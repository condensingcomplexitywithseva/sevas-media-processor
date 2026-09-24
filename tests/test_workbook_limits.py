# Copyright 2026 Vsevolod Belonogov
# SPDX-License-Identifier: Apache-2.0


import json
import logging
import re
import sqlite3
import sys
import time
import tracemalloc
from pathlib import Path
from typing import Any

import pytest

SRC = Path(__file__).resolve().parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

import openpyxl
from openpyxl.cell.cell import Cell

import data_exporter
from db_controller import SQLiteDatabaseController
from schemas import OverallResult, Status

from test_results_report import (
    REQUEST_STATUSES,
    RESULTS_TITLE,
    add_file,
    answer_row,
    cell_text,
    export_workbook,
    make_exporter,
    new_name,
    read_csv_rows,
    record_configuration,
    request_outcome,
    results_sheet,
    seed_acceptance_example,
    seed_one_file,
    sheet_rows,
    store_requests,
)

SUMMARY = "Summary"
DIVIDER = "Details>>>"
AI_ORDER = [SUMMARY, RESULTS_TITLE, DIVIDER, "Master Registry", "Page Log", "LLM Requests", "LLM Answers"]
AI_OFF_ORDER = [SUMMARY, RESULTS_TITLE, DIVIDER, "Master Registry", "Page Log"]
EXCEL_TITLE_FORBIDDEN = set(":\\/?*[]")
ROWS_CUT = re.compile(r"^(?P<sheet>.+?): (?P<cut>\d+) rows omitted; (?P<kept>\d+) retained\.$")
COLUMNS_CUT = re.compile(r"^(?P<sheet>.+?): (?P<cut>\d+) columns omitted:")
CELLS_CUT = re.compile(r"^(?P<sheet>.+?): Text shortened: .+ \(column \d+\), (?P<cut>\d+) cells?\.$")




def summary_table(workbook) -> list[tuple[str, str, Any]]:
    headings = {"Files": "file records", "File attempts": "file records", "Pages": "pages",
                "AI requests": "requests", "Answer problems": "answers",
                "Complete CSV reports, saved beside this workbook": "export", "Reports not saved": "failed",
                "Workbook omissions": "limits", "Text replacements": "text"}
    labels = {
        "file records": {"Fully processed": "ok", "Partly processed": "partial_fail",
                         "Failed": "fail", "Skipped": "skipped", "Distinct files": "distinct files"},
        "pages": {"Converted": "ok", "Failed": "failure", "Skipped": "skipped"},
        "requests": {"Completed": "ok", "Stopped by user": "aborted_by_user",
                     "Images could not be prepared": "encoding_failure",
                     "Reply exceeded response limit": "token_limit_exceeded", "Request failed": "network_failure",
                     "Unreadable provider reply": "provider_reply_parse_error",
                     "Invalid answer format": "invalid_json_answer", "Not sent": "not_attempted",
                     "Interrupted; execution uncertain": "interrupted"},
        "answers": {"Unverified page claim": "hallucinated_page_number",
                    "Additional answer for the same page": "duplicate_page_number",
                    "No answer returned for this page": "no_row_returned"},
    }
    result = []
    group = None
    for row in sheet_rows(workbook, SUMMARY):
        first, second = [*row, None, None][:2]
        if first in headings and second is None:
            group = headings[first]
        elif first == data_exporter.WORKBOOK_COMPLETE_LINE:
            result.append(("limits", first, None))
        elif group and second is not None and first not in ("Item", "Report", "Excel sheet"):
            if group == "limits":
                result.append((group, f"{first}: {second}", None))
            else:
                result.append((group, labels.get(group, {}).get(first, first), second))
    return result


def section(table, name) -> dict[str, Any]:
    return {item: count for sect, item, count in table if sect == name}


def limits_lines(table) -> list[str]:
    return [item for sect, item, _ in table if sect == "limits"]


def limits_cells(workbook):
    active = False
    for row in workbook[SUMMARY]:
        if row[0].value == data_exporter.WORKBOOK_COMPLETE_LINE:
            yield row[0]
        if row[0].value == "Workbook omissions":
            active = True
        elif active and row[1].value is None:
            active = False
        elif active and row[0].value != "Excel sheet":
            yield row[1]


def export_dir(db_path, out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    exporter = make_exporter(db_path)
    exporter.export_all_formats(out_dir)
    return out_dir


def workbook_in(out_dir):
    return openpyxl.load_workbook(next(out_dir.glob("database_export_*.xlsx")))


def csv_named(out_dir, stem) -> Path:
    matches = list(out_dir.glob(f"{stem}_*.csv"))
    assert len(matches) == 1, (stem, matches)
    return matches[0]


def data_rows(sheet) -> list[list[Any]]:
    return [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)]


def marker_only(value) -> bool:
    return isinstance(value, str) and value.strip() in (
        data_exporter.ROW_OMISSION_MARKER,
        data_exporter.COLUMN_OMISSION_MARKER,
    )




def test_sheet_order_on_an_ai_run(tmp_path):
    seed_acceptance_example(tmp_path / "state.db")
    workbook = export_workbook(tmp_path / "state.db", tmp_path / "out")
    assert workbook.sheetnames == AI_ORDER


@pytest.mark.parametrize("ai_on", [False, True])
def test_summary_explains_run_and_counts_before_the_numbers(tmp_path, ai_on):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    record_configuration(c, ai_on)
    c.close()
    book = export_workbook(db, tmp_path / "out")
    rows = sheet_rows(book, "Summary")
    heading = next(i for i, row in enumerate(rows) if row[0] == "Files")
    text = " ".join(str(row[0] or "") for row in rows[:heading])
    assert rows[0][0] == "Your processing report"
    assert rows[1][0] == ("0 files - AI enabled" if ai_on else "0 files - AI off")
    assert ("AI detail sheets" in text) == ai_on
    assert "If you need more detail, you can use file_id" in text
    assert "attempt" not in text and "database records" not in text
    assert book[SUMMARY].max_column == 2
    assert ["Item", "Count"] in rows
    assert ["Complete CSV reports, saved beside this workbook", None] in rows
    assert ["Saved beside this workbook.", None] not in rows
    for row in book[SUMMARY]:
        if row[0].value in ("Files", "Pages", "AI requests", "Complete CSV reports, saved beside this workbook"):
            assert row[0].font.bold and row[0].font.sz == 14
            assert row[0].row is not None
            assert book[SUMMARY].row_dimensions[row[0].row].height >= 23
    assert book[SUMMARY]["A1"].font.sz == 16
    assert ["section", "item", "count"] not in rows
    assert not any("Clean" in str(row) or "clean" in str(row) or "formula bar" in str(row) for row in rows)
    assert "Answer problems" not in [row[0] for row in rows]
    assert "Record 1" not in str(rows), "CSV locator guidance belongs only to shortened text"


def test_omitted_columns_are_named_and_distinct_from_omitted_rows(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_COLUMNS", 5)
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_ROWS", 3)
    seed_ten_results_rows(tmp_path / "state.db")
    out = export_dir(tmp_path / "state.db", tmp_path / "out")
    book = workbook_in(out)
    lines = limits_lines(summary_table(book))
    assert ["Excel sheet", "Omitted"] in sheet_rows(book, SUMMARY)
    column = next(line for line in lines if line.startswith("Results: 2 columns omitted:"))
    assert "llm_answer, notes (columns 5-6)" in column
    assert section(summary_table(book), "export")["Results"] == csv_named(out, "results").name
    assert any(line.startswith("Results: 8 rows omitted; 2 retained.") for line in lines)
    assert book["Results"].cell(1, 5).value == data_exporter.COLUMN_OMISSION_MARKER
    assert book["Results"].cell(4, 1).value == data_exporter.ROW_OMISSION_MARKER


def test_shortened_text_names_the_complete_csv_and_record(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_CELL_CHARS", 200)
    text = "x" * 300
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None)],
        requests=[request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": text}),))],
    )
    out = export_dir(db, tmp_path / "out")
    book = workbook_in(out)
    _, rows = results_sheet(book)
    cell = rows[0]["llm_answer"]
    assert cell.startswith(data_exporter.TRUNCATION_MARKER)
    assert csv_named(out, "results").name in cell and "record 1, column 5" in cell
    assert "Summary for full data" not in cell
    lines = limits_lines(summary_table(book))
    assert any("Results: Text shortened: llm_answer (column 5), 1 cell." in line for line in lines)
    _, complete = read_csv_rows(csv_named(out, "results"))
    assert complete[0]["llm_answer"] == text


def test_sheet_order_on_an_ai_off_run(tmp_path):
    seed_one_file(tmp_path / "state.db", [(1, "1_p1.jpg", "ok", "", None)])
    workbook = export_workbook(tmp_path / "state.db", tmp_path / "out")
    assert workbook.sheetnames == AI_OFF_ORDER


def test_every_title_obeys_excels_own_rules(tmp_path):
    seed_acceptance_example(tmp_path / "state.db")
    workbook = export_workbook(tmp_path / "state.db", tmp_path / "out")
    for title in workbook.sheetnames:
        assert 1 <= len(title) <= 31, title
        assert not (set(title) & EXCEL_TITLE_FORBIDDEN), title
        assert "- Part" not in title, title
    assert "Page Log" in workbook.sheetnames and "Page Log - Part 1" not in workbook.sheetnames


def test_the_divider_sheet_is_empty(tmp_path):
    seed_acceptance_example(tmp_path / "state.db")
    workbook = export_workbook(tmp_path / "state.db", tmp_path / "out")
    assert list(workbook[DIVIDER].iter_rows(values_only=True)) == []


def test_the_real_constants_sit_under_excels_limits():
    assert data_exporter.EXCEL_MAX_ROWS == 1_000_000 < 1_048_576
    assert data_exporter.EXCEL_MAX_COLUMNS == 16_000 < 16_384
    assert data_exporter.EXCEL_MAX_CELL_CHARS == 32_000 < 32_767
    assert len(data_exporter.TRUNCATION_MARKER) < data_exporter.EXCEL_MAX_CELL_CHARS


def test_every_sheet_fits_excel_under_the_real_constants(tmp_path):
    seed_acceptance_example(tmp_path / "state.db")
    workbook = export_workbook(tmp_path / "state.db", tmp_path / "out")
    for sheet in workbook.worksheets:
        assert sheet.max_row <= 1_048_576 and sheet.max_column <= 16_384, sheet.title




def test_summary_counts_the_acceptance_example_with_zeros_and_distinct_files(tmp_path):
    seed_acceptance_example(tmp_path / "state.db")
    table = summary_table(export_workbook(tmp_path / "state.db", tmp_path / "out"))
    files = section(table, "file records")
    assert {m.value: files[m.value] for m in OverallResult} == {"ok": 1, "partial_fail": 2, "fail": 2, "skipped": 0}
    assert "distinct files" not in files
    pages = section(table, "pages")
    assert {m.value: pages[m.value] for m in Status} == {"ok": 5, "failure": 3, "skipped": 0}
    requests = section(table, "requests")
    assert requests == {
        "ok": 2,
        "network_failure": 1,

    }
    answers = section(table, "answers")
    assert answers == {}, "No answer-problems block when no such problems were recorded"
    assert section(table, "ai") == {}


def test_summary_distinct_files_counts_paths_not_records(tmp_path):
    c = SQLiteDatabaseController(tmp_path / "state.db")
    add_file(
        c,
        1,
        "photo.png",
        [(1, "", "failure", "x", None)],
        total=0,
        page_range="",
        range_status="failure",
        comment="first",
    )
    add_file(c, 2, "photo.png", [(1, "2_p1.jpg", "ok", "", None)], comment="second")
    record_configuration(c, ai_on=False)
    c.close()
    files = section(summary_table(export_workbook(tmp_path / "state.db", tmp_path / "out")), "file records")
    assert files["fail"] == 1 and files["ok"] == 1 and files["distinct files"] == 1
    book = export_workbook(tmp_path / "state.db", tmp_path / "second")
    assert book[SUMMARY]["A2"].value == "2 attempts for 1 file - AI off"
    assert "File attempts" in str(sheet_rows(book, SUMMARY))


def test_summary_says_ai_was_off_and_has_no_ai_blocks(tmp_path):
    seed_one_file(tmp_path / "state.db", [(1, "1_p1.jpg", "ok", "", None)])
    book = export_workbook(tmp_path / "state.db", tmp_path / "out")
    table = summary_table(book)
    assert section(table, "requests") == {} and section(table, "answers") == {}
    assert book[SUMMARY]["A2"].value == "1 file - AI off"


def test_summary_counts_equal_the_sheets_below_it(tmp_path):
    seed_acceptance_example(tmp_path / "state.db")
    workbook = export_workbook(tmp_path / "state.db", tmp_path / "out")
    table = summary_table(workbook)

    def column(title, name):
        rows = sheet_rows(workbook, title)
        index = [cell_text(c) for c in rows[0]].index(name)
        return [cell_text(r[index]) for r in rows[1:]]

    overall = column("Master Registry", "overall_result")
    assert {k: v for k, v in section(table, "file records").items() if k != "distinct files"} == {
        m.value: overall.count(m.value) for m in OverallResult
    }
    page_statuses = column("Page Log", "page_to_jpeg_status")
    assert section(table, "pages") == {m.value: page_statuses.count(m.value) for m in Status}
    request_statuses = column("LLM Requests", "request_status")
    assert section(table, "requests") == {s: request_statuses.count(s) for s in REQUEST_STATUSES
                                          if s == "ok" or request_statuses.count(s)}
    errors = column("LLM Answers", "llm_error")
    assert section(table, "answers") == {error: errors.count(error) for error in set(errors) if error}


def test_summary_limits_block_reads_complete_when_nothing_was_cut(tmp_path):
    seed_acceptance_example(tmp_path / "state.db")
    workbook = export_workbook(tmp_path / "state.db", tmp_path / "out")
    complete = new_name(data_exporter, "WORKBOOK_COMPLETE_LINE", "the Summary's complete line")
    assert limits_lines(summary_table(workbook)) == [complete]
    cell = next(limits_cells(workbook))
    assert not cell.font.bold




def seed_long_comments(db_path, length):
    seed_one_file(
        db_path,
        [(1, "", "failure", "p" * length, None)],
        total=0,
        page_range="",
        range_status="failure",
        comment="f" * length,
    )
    return db_path


def test_cells_at_the_cap_pass_and_over_it_are_cut_within_the_cap(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_CELL_CHARS", 60)
    seed_one_file(
        tmp_path / "state.db",
        [(1, "", "failure", "p" * 61, None)],
        total=0,
        page_range="",
        range_status="failure",
        comment="f" * 60,
    )
    workbook = workbook_in(export_dir(tmp_path / "state.db", tmp_path / "out"))
    registry = sheet_rows(workbook, "Master Registry")
    assert registry[1][registry[0].index("file_to_jpegs_comment")] == "f" * 60
    pages = sheet_rows(workbook, "Page Log")
    cut = pages[1][pages[0].index("page_to_jpeg_comment")]
    assert cut.startswith(data_exporter.TRUNCATION_MARKER) and len(cut) == 60


def test_cut_cells_are_counted_across_the_whole_workbook(tmp_path, monkeypatch, caplog):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_CELL_CHARS", 40)
    seed_long_comments(tmp_path / "state.db", 100)
    with caplog.at_level(logging.WARNING):
        out = export_dir(tmp_path / "state.db", tmp_path / "out")
    workbook = workbook_in(out)
    lines = limits_lines(summary_table(workbook))
    cells = [CELLS_CUT.match(line) for line in lines]
    cells = [m for m in cells if m]
    assert sum(int(m["cut"]) for m in cells) == 3, lines
    assert new_name(data_exporter, "WORKBOOK_COMPLETE_LINE", "the Summary's complete line") not in lines
    assert sum("cells contain shortened text across" in r.getMessage() for r in caplog.records) == 1
    for cell in limits_cells(workbook):
        assert cell.font.bold and cell.font.color is not None
        assert str(cell.font.color.rgb).upper().endswith("FF0000")




def seed_all_columns(db_path):
    return seed_one_file(
        db_path,
        [(1, "1_p1.jpg", "ok", "", 1.0)],
        path="clip.mp4",
        requests=[request_outcome(pages=(1,), error="note", answer_rows=(answer_row(1, "1", "", {"answer": "a"}),))],
    )


def test_columns_beyond_the_cap_become_one_marker_column(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_COLUMNS", 4)
    seed_all_columns(tmp_path / "state.db")
    out = export_dir(tmp_path / "state.db", tmp_path / "out")
    workbook = workbook_in(out)
    sheet = workbook[RESULTS_TITLE]
    assert sheet.max_column == 4
    header = [c.value for c in sheet[1]]
    assert header[:3] == ["file_id", "file_path", "pages"] and marker_only(header[3])
    assert all(marker_only(row[3]) for row in data_rows(sheet))
    hits = [
        m
        for m in (COLUMNS_CUT.match(line) for line in limits_lines(summary_table(workbook)))
        if m and m["sheet"] == RESULTS_TITLE
    ]
    assert len(hits) == 1 and int(hits[0]["cut"]) == 3
    headers, _ = read_csv_rows(csv_named(out, "results"))
    assert headers == ["file_id", "file_path", "pages", "file_result", "llm_answer", "notes"]


def test_two_over_wide_sheets_get_one_line_each(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_COLUMNS", 5)
    seed_all_columns(tmp_path / "state.db")
    workbook = workbook_in(export_dir(tmp_path / "state.db", tmp_path / "out"))
    sheets_named = sorted(
        m["sheet"] for m in (COLUMNS_CUT.match(line) for line in limits_lines(summary_table(workbook))) if m
    )
    over_wide = sorted(
        s.title
        for s in workbook.worksheets
        if s.title != SUMMARY and s.max_column == 5 and marker_only(s.cell(1, 5).value)
    )
    assert sheets_named == over_wide and len(over_wide) >= 2, (sheets_named, over_wide)


def test_seventeen_thousand_columns_survive_with_a_complete_csv(tmp_path):
    keys = [f"k{i}" for i in range(17_000)]
    seed_one_file(
        tmp_path / "state.db",
        [(1, "1_p1.jpg", "ok", "", None)],
        requests=[
            request_outcome(pages=(1,), answer_rows=(answer_row(1, "1", "", {k: f"v{i}" for i, k in enumerate(keys)}),))
        ],
        declared=keys,
    )
    out = export_dir(tmp_path / "state.db", tmp_path / "out")
    csv_headers, csv_rows = read_csv_rows(csv_named(out, "results"))
    assert [h for h in csv_headers if h.startswith("llm_")] == [f"llm_{k}" for k in keys]
    assert csv_rows[0][f"llm_{keys[-1]}"] == "v16999"
    workbook = workbook_in(out)
    for sheet in workbook.worksheets:
        assert sheet.max_column <= 16_384, sheet.title
    assert workbook[RESULTS_TITLE].max_column == data_exporter.EXCEL_MAX_COLUMNS
    assert any(COLUMNS_CUT.match(line) for line in limits_lines(summary_table(workbook)))




def seed_ten_results_rows(db_path):
    return seed_one_file(
        db_path,
        [(n, f"{n}.jpg", "ok", "", None) for n in range(1, 11)],
        requests=[
            request_outcome(request_number=n, pages=(n,), answer_rows=(answer_row(values={"answer": str(n)}),))
            for n in range(1, 11)
        ],
    )


def test_rows_beyond_the_cap_become_one_marker_row(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_ROWS", 6)
    seed_ten_results_rows(tmp_path / "state.db")
    out = export_dir(tmp_path / "state.db", tmp_path / "out")
    workbook = workbook_in(out)
    rows = data_rows(workbook[RESULTS_TITLE])
    assert len(rows) == 6 and all(marker_only(c) for c in rows[-1])
    assert [row[2] for row in rows[:5]] == [1, 2, 3, 4, 5]
    hits = [
        m
        for m in (ROWS_CUT.match(line) for line in limits_lines(summary_table(workbook)))
        if m and m["sheet"] == RESULTS_TITLE
    ]
    assert len(hits) == 1 and (int(hits[0]["cut"]), int(hits[0]["kept"])) == (5, 5)
    assert section(summary_table(workbook), "export")["Results"] == csv_named(out, "results").name
    _, complete = read_csv_rows(csv_named(out, "results"))
    assert len(complete) == 10


def test_the_page_log_truncates_instead_of_splitting(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_ROWS", 5)
    seed_one_file(tmp_path / "state.db", [(n, f"1_p{n}.jpg", "ok", "", None) for n in range(1, 13)])
    out = export_dir(tmp_path / "state.db", tmp_path / "out")
    workbook = workbook_in(out)
    assert [t for t in workbook.sheetnames if t.startswith("Page Log")] == ["Page Log"]
    rows = data_rows(workbook["Page Log"])
    assert len(rows) == 5 and all(marker_only(c) for c in rows[-1] if c is not None)
    assert [r[2] for r in rows[:4]] == [1, 2, 3, 4]
    hits = [m for m in (ROWS_CUT.match(line) for line in limits_lines(summary_table(workbook))) if m]
    assert [(m["sheet"], int(m["cut"]), int(m["kept"])) for m in hits] == [("Page Log", 8, 4)]
    assert section(summary_table(workbook), "export")["Page Log"] == csv_named(out, "page_log").name
    _, csv_rows = read_csv_rows(csv_named(out, "page_log"))
    assert len(csv_rows) == 12


@pytest.mark.parametrize("title, stem", [("LLM Requests", "llm_requests"), ("LLM Answers", "llm_answers")])
def test_the_ai_sheets_truncate_the_same_way(tmp_path, monkeypatch, title, stem):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_ROWS", 5)
    seed_one_file(
        tmp_path / "state.db",
        [(n, f"1_p{n}.jpg", "ok", "", None) for n in range(1, 13)],
        requests=[
            request_outcome(request_number=n, pages=(n,), answer_rows=(answer_row(n, str(n), "", {"answer": "a"}),))
            for n in range(1, 13)
        ],
    )
    out = export_dir(tmp_path / "state.db", tmp_path / "out")
    workbook = workbook_in(out)
    rows = data_rows(workbook[title])
    assert len(rows) == 5 and all(marker_only(c) for c in rows[-1] if c is not None)
    hits = [m for m in (ROWS_CUT.match(line) for line in limits_lines(summary_table(workbook))) if m]
    assert any(m["sheet"] == title for m in hits)
    assert section(summary_table(workbook), "export")[title] == csv_named(out, stem).name


def test_rows_and_columns_cut_on_one_sheet_are_both_announced(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_ROWS", 4)
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_COLUMNS", 4)
    seed_ten_results_rows(tmp_path / "state.db")
    workbook = workbook_in(export_dir(tmp_path / "state.db", tmp_path / "out"))
    sheet = workbook[RESULTS_TITLE]
    assert sheet.max_column == 4 and len(data_rows(sheet)) == 4
    corner = sheet.cell(row=sheet.max_row, column=4).value
    assert marker_only(corner), "one marker in the corner, never two"
    lines = limits_lines(summary_table(workbook))
    assert any(m and m["sheet"] == RESULTS_TITLE for m in map(ROWS_CUT.match, lines))
    assert any(m and m["sheet"] == RESULTS_TITLE for m in map(COLUMNS_CUT.match, lines))


def test_a_cut_on_the_last_sheet_written_is_still_announced(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_CELL_CHARS", 40)
    seed_one_file(
        tmp_path / "state.db",
        [(1, "1_p1.jpg", "ok", "", None)],
        requests=[request_outcome(pages=(1,), answer_rows=(answer_row(1, "1", "", {"answer": "a" * 100}),))],
    )
    workbook = workbook_in(export_dir(tmp_path / "state.db", tmp_path / "out"))
    assert any(CELLS_CUT.match(line) for line in limits_lines(summary_table(workbook)))




def test_the_csvs_are_complete_when_every_cut_fires(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_ROWS", 4)
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_COLUMNS", 4)
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_CELL_CHARS", 40)
    seed_long_comments(tmp_path / "state.db", 100)
    c = SQLiteDatabaseController(tmp_path / "state.db")
    add_file(c, 2, "clip.mp4", [(n, f"2_p{n}.jpg", "ok", "", float(n)) for n in range(1, 8)])
    store_requests(
        c,
        2,
        [
            request_outcome(
                pages=tuple(range(1, 8)),
                answer_rows=tuple(answer_row(n, str(n), "", {"answer": "a" * 100}) for n in range(1, 8)),
            )
        ],
    )
    record_configuration(c, ai_on=True, declared=("answer",))
    c.close()
    out = export_dir(tmp_path / "state.db", tmp_path / "out")
    headers, rows = read_csv_rows(csv_named(out, "results"))
    assert len(headers) == 6 and len(rows) == 8
    assert "f" * 100 in rows[0]["notes"]
    assert all(r["llm_answer"] == "a" * 100 for r in rows[1:])
    results_text = csv_named(out, "results").read_text(encoding="utf-8")
    assert data_exporter.TRUNCATION_MARKER.strip() not in results_text


def test_page_range_exports_verbatim(tmp_path):
    fragments = ", ".join(str(n) for n in range(1, 100, 2))
    seed_one_file(tmp_path / "state.db", [(1, "1_p1.jpg", "ok", "", None)], total=99, page_range=fragments)
    out = export_dir(tmp_path / "state.db", tmp_path / "out")
    _, rows = read_csv_rows(csv_named(out, "file_registry"))
    assert rows[0]["page_range"] == fragments
    registry = sheet_rows(workbook_in(out), "Master Registry")
    assert registry[1][registry[0].index("page_range")] == fragments


def test_one_export_shares_one_timestamp_across_every_file(tmp_path):
    seed_acceptance_example(tmp_path / "state.db")
    out = export_dir(tmp_path / "state.db", tmp_path / "out")
    stamps = {re.sub(r"^[a-z_]+_", "", p.stem) for p in out.glob("*.csv")}
    stamps |= {re.sub(r"^database_export_", "", p.stem) for p in out.glob("*.xlsx")}
    assert len(stamps) == 1, stamps
    assert {re.sub(r"_\d{8}_\d{6}$", "", p.stem) for p in out.glob("*.csv")} >= {
        "results",
        "file_registry",
        "page_log",
        "llm_requests",
        "llm_answers",
    }




def test_large_csv_report_preserves_every_record(tmp_path):
    db = tmp_path / "state.db"
    spec = resource_database(db, "ai", 2000)
    outcome = make_exporter(db).export_csv(tmp_path / "out")
    assert_resource_output(outcome, spec, workbook=False)




def test_summary_counts_are_the_runs_truth_when_rows_were_cut(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_ROWS", 6)
    seed_ten_results_rows(tmp_path / "state.db")
    workbook = workbook_in(export_dir(tmp_path / "state.db", tmp_path / "out"))
    table = summary_table(workbook)
    assert section(table, "pages")["ok"] == 10
    assert section(table, "file records")["ok"] == 1
    assert "distinct files" not in section(table, "file records")
    assert len(data_rows(workbook["Page Log"])) == 6, "the sheet itself was cut"


GRIN = "\U0001f600"


def test_the_cell_cap_counts_what_excel_counts(tmp_path):
    seed_one_file(
        tmp_path / "state.db",
        [(1, "1_p1.jpg", "ok", "", None)],
        requests=[request_outcome(pages=(1,), answer_rows=(answer_row(1, "1", "", {"answer": GRIN * 16_384}),))],
    )
    out = export_dir(tmp_path / "state.db", tmp_path / "out")
    workbook = workbook_in(out)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str):
                    units = len(cell.encode("utf-16-le")) // 2
                    assert units <= 32_767, (sheet.title, units)
                    assert units <= data_exporter.EXCEL_MAX_CELL_CHARS, (sheet.title, units)
    assert any(CELLS_CUT.match(line) for line in limits_lines(summary_table(workbook)))
    _, rows = read_csv_rows(csv_named(out, "llm_answers"))
    assert rows[0]["llm_answer"] == GRIN * 16_384


def test_the_marker_passes_through_the_same_funnel_as_every_cell(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "TRUNCATION_MARKER", "[CUT\x07] ")
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_CELL_CHARS", 60)
    seed_long_comments(tmp_path / "state.db", 100)
    out = export_dir(tmp_path / "state.db", tmp_path / "out")
    _, rows = results_sheet(workbook_in(out))
    assert rows[0]["notes"].startswith("[CUT�] ")


def test_giant_file_csv_preserves_every_page(tmp_path):
    db = tmp_path / "state.db"
    spec = resource_database(db, "giant", 200000)
    outcome = make_exporter(db).export_csv(tmp_path / "out")
    assert_resource_output(outcome, spec, workbook=False)


def test_bound_export_lookups_use_indexes_including_aliased_tables(tmp_path, monkeypatch):
    db_path = tmp_path / "state.db"
    seed_acceptance_example(db_path)
    statements = []
    real_connect = sqlite3.connect

    class RecordingConnection(sqlite3.Connection):
        def execute(self, sql, parameters=(), /):
            if parameters:
                statements.append((sql, parameters))
            return super().execute(sql, parameters)

    def recording_connect(*args, **kwargs):
        kwargs["factory"] = RecordingConnection
        return real_connect(*args, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(data_exporter.sqlite3, "connect", recording_connect)
        make_exporter(db_path).export_csv(tmp_path / "out", timestamp="t")
    assert len(statements) >= 3

    def scans(connection, sql, parameters):
        return [
            row[3] for row in connection.execute("EXPLAIN QUERY PLAN " + sql, parameters) if row[3].startswith("SCAN")
        ]

    with real_connect(db_path) as connection:
        for sql, parameters in statements:
            assert scans(connection, sql, parameters) == [], sql
        control = "SELECT p.page_number FROM page_log p WHERE p.file_id=?"
        assert scans(connection, control, (1,)) == []
        connection.execute("DROP INDEX ix_page_log_file_id")
        connection.execute("DROP INDEX ix_page_log_file_page")
    with real_connect(db_path) as connection:
        assert scans(connection, control, (1,)), "an unindexed aliased lookup must be detected"


def test_summary_and_the_divider_are_never_cut_or_counted(tmp_path, monkeypatch):
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_ROWS", 4)
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_COLUMNS", 4)
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_CELL_CHARS", 60)
    seed_long_comments(tmp_path / "state.db", 100)
    workbook = workbook_in(export_dir(tmp_path / "state.db", tmp_path / "out"))
    for title in (SUMMARY, DIVIDER):
        for row in workbook[title].iter_rows(values_only=True):
            assert not any(marker_only(cell) for cell in row), title
    assert len(list(workbook[SUMMARY].iter_rows())) > 4, "Summary is never row-cut"
    assert workbook[SUMMARY].max_column == 2
    lines = limits_lines(summary_table(workbook))
    assert lines and not any(SUMMARY in line or DIVIDER in line for line in lines)


@pytest.mark.parametrize("error, label", [
    ("hallucinated_page_number", "Unverified page claim"),
    ("duplicate_page_number", "Additional answer for the same page"),
    ("no_row_returned", "No answer returned for this page"),
])
def test_summary_displays_only_recorded_answer_problems(tmp_path, error, label):
    values = None if error == "no_row_returned" else {"answer": "invoice"}
    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)],
                      requests=[request_outcome(pages=(1,), answer_rows=(answer_row(1, "1", error, values),))])
    book = export_workbook(db, tmp_path / "out")
    assert section(summary_table(book), "answers") == {error: 1}
    rows = sheet_rows(book, SUMMARY)
    assert [label, 1] in rows
    assert ["Answer problems", None] in rows
    assert "clean" not in str(rows).lower()


@pytest.mark.parametrize("all_fail", [False, True])
def test_summary_separates_saved_csvs_from_failures(tmp_path, monkeypatch, all_fail):
    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)])
    original = data_exporter.SQLiteDataExporter._write_csv

    def fail(self, path, headers, rows):
        if all_fail or "page_id" in headers:
            raise PermissionError("Destination is not writable")
        return original(self, path, headers, rows)

    monkeypatch.setattr(data_exporter.SQLiteDataExporter, "_write_csv", fail)
    with pytest.raises(data_exporter.ExportError):
        export_dir(db, tmp_path / "out")
    book = workbook_in(tmp_path / "out")
    table = summary_table(book)
    saved = section(table, "export")
    failures = section(table, "failed")
    assert set(saved) == (set() if all_fail else {"Results", "Master Registry"})
    assert set(failures) == ({"Results", "Master Registry", "Page Log"} if all_fail else {"Page Log"})
    assert not set(saved) & set(failures)
    for report, filename in saved.items():
        assert (tmp_path / "out" / filename).is_file(), report
    rows = sheet_rows(book, SUMMARY)
    assert ("Complete CSV reports" in str(rows)) == (not all_fail)
    assert "Retry export or save reports elsewhere" in str(rows)
    for row in book[SUMMARY]:
        if row[0].value in failures:
            assert row[1].font.color.rgb == "FFFF0000"


@pytest.mark.parametrize("cut", [False, True])
def test_adversarial_summary_census_is_independent_of_answer_multiplicity_and_cuts(tmp_path, monkeypatch, cut):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    try:
        add_file(c, 90, "repeat.pdf", [(1, "90.jpg", "ok", "", None)])
        add_file(c, 10, "repeat.pdf", [(1, "10.jpg", "ok", "", None), (2, "", "failure", "bad page", None)])
        add_file(c, 50, "other/repeat.pdf", [], range_status="skipped")
        store_requests(c, 10, [request_outcome(pages=(1,), answer_rows=(
            answer_row(1, values={"answer": "first"}),
            answer_row(1, "1", "duplicate_page_number", {"answer": "second"}),
            answer_row(None, "99", "hallucinated_page_number", {"answer": "third"}),
        ))])
        store_requests(c, 90, [request_outcome(pages=(1,), answer_rows=(answer_row(1, "1", "no_row_returned"),)),
                               request_outcome(request_number=2, pages=(1,), status="network_failure", error="offline"),
                               request_outcome(request_number=3, pages=(1,), status="not_attempted")])
        record_configuration(c, True, output_mode="table_per_page")
    finally:
        c.close()
    if cut:
        monkeypatch.setattr(data_exporter, "EXCEL_MAX_ROWS", 2)
        monkeypatch.setattr(data_exporter, "EXCEL_MAX_COLUMNS", 3)
    book = export_workbook(db, tmp_path / "out")
    try:
        table = summary_table(book)
        assert section(table, "file records") == {
            "ok": 0, "partial_fail": 2, "fail": 0, "skipped": 1, "distinct files": 2,
        }
        assert section(table, "pages") == {"ok": 2, "failure": 1, "skipped": 0}
        assert section(table, "requests") == {"ok": 2, "network_failure": 1, "not_attempted": 1}
        assert section(table, "answers") == {
            "hallucinated_page_number": 1, "duplicate_page_number": 1, "no_row_returned": 1,
        }
        assert book[SUMMARY]["A2"].value == "3 attempts for 2 files - AI enabled"
        headings = {"Your processing report", "File attempts", "Pages", "AI requests", "Answer problems",
                    "Complete CSV reports, saved beside this workbook"}
        found = set()
        for row in book[SUMMARY]:
            if row[0].value in headings:
                found.add(row[0].value)
                assert row[0].font.bold
                assert row[0].font.sz == (16 if row[0].value == "Your processing report" else 14)
                assert row[0].row is not None
                assert book[SUMMARY].row_dimensions[row[0].row].height >= 23
        assert found == headings
        assert "clean" not in str(sheet_rows(book, SUMMARY)).lower()
        assert len(read_csv_rows(csv_named(tmp_path / "out", "results"))[1]) == 7
    finally:
        book.close()


@pytest.mark.parametrize("failed_report", ["Results", "Master Registry", "Page Log", "LLM Requests", "LLM Answers"])
def test_adversarial_reordered_csv_locators_remain_exact_with_each_missing_sibling(
    tmp_path, monkeypatch, failed_report,
):
    db = tmp_path / "state.db"
    long = 'line one, "quoted"\r\n_x000D_' + '\U0001f600' * 400
    c = SQLiteDatabaseController(db)
    try:
        add_file(c, 9007199254740993, "input.pdf", [(7, long, "ok", long, None)], comment=long)
        store_requests(c, 9007199254740993, [request_outcome(pages=(7,), raw_answer=long,
            answer_rows=(answer_row(7, "7", "", {"answer": long}),))])
        record_configuration(c, True, output_mode="table_per_page")
    finally:
        c.close()
    headers_to_title = {
        ("file_id", "file_path", "pages", "file_result", "llm_answer", "notes"): "Results",
        ("file_id", "file_path", "file_ext", "total_pages", "page_range", "range_status", "file_to_jpegs_comment",
         "file_to_jpegs_status", "file_to_llm_status", "overall_result"): "Master Registry",
        ("file_id", "page_id", "page_number", "output_file", "page_to_jpeg_status", "page_to_jpeg_comment",
         "video_frame_timestamp"): "Page Log",
        ("file_id", "request_id", "request_number", "pages", "request_status",
         "raw_llm_answer", "llm_network_error"): "LLM Requests",
        ("file_id", "page_id", "request_id", "llm_answer_id",
         "raw_model_page_number", "llm_error", "llm_answer"): "LLM Answers",
    }
    original = data_exporter.SQLiteDataExporter._write_csv
    reached = []

    def fail_one(self, path, headers, rows):
        title = headers_to_title[tuple(headers)]
        reached.append(title)
        if title == failed_report:
            raise PermissionError("one report unavailable")
        return original(self, path, headers, rows)

    monkeypatch.setattr(data_exporter.SQLiteDataExporter, "_write_csv", fail_one)
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_CELL_CHARS", 240)
    with pytest.raises(data_exporter.ExportError) as caught:
        make_exporter(db).export_all_formats(tmp_path / "out")
    outcome = caught.value.outcome
    assert reached == list(headers_to_title.values())
    assert [(r["report"], r["format"]) for r in outcome.failed] == [(failed_report, "csv")]
    assert {r["report"] for r in outcome.saved} == set(headers_to_title.values()) - {failed_report} | {"Workbook"}
    assert {Path(r["path"]).name for r in outcome.saved} == {p.name for p in (tmp_path / "out").iterdir()}
    book = openpyxl.load_workbook(next(r["path"] for r in outcome.saved if r["format"] == "xlsx"))
    try:
        for headers, title in headers_to_title.items():
            rows = sheet_rows(book, title)
            assert rows[0] == list(headers)
            assert str(rows[1][0]) == "9007199254740993"
            marked = [(column, value) for column, value in enumerate(rows[1], 1)
                      if isinstance(value, str) and value.startswith("[TEXT SHORTENED]")]
            assert marked, title
            if title == failed_report:
                assert all("CSV unavailable" in value for _, value in marked)
            else:
                path = Path(next(r["path"] for r in outcome.saved if r["report"] == title))
                csv_headers, csv_rows = read_csv_rows(path)
                assert csv_headers == list(headers)
                for column, value in marked:
                    assert f"{path.name}; record 1, column {column}." in value
                    assert long in csv_rows[0][csv_headers[column - 1]]
                    assert len(value.encode("utf-16-le")) // 2 <= 240
        saved = section(summary_table(book), "export")
        assert set(saved) == set(headers_to_title.values()) - {failed_report}
        assert section(summary_table(book), "failed") == {failed_report: "one report unavailable"}
    finally:
        book.close()


@pytest.mark.parametrize("size", [0, 1, 2, 3, 4])
def test_adversarial_row_budget_exact_boundary_preserves_prefix_and_csv(tmp_path, monkeypatch, size):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    try:
        record_configuration(c, False)
        for fid in range(1, size + 1):
            add_file(c, fid, f"{fid}.png", [(1, f"{fid}.jpg", "ok", "", None)])
    finally:
        c.close()
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_ROWS", 3)
    book = export_workbook(db, tmp_path / "out")
    try:
        actual = sheet_rows(book, "Results")[1:]
        assert [row[0] for row in actual] == (list(range(1, size + 1)) if size <= 3 else [1, 2, "[ROWS OMITTED]"])
        assert len(read_csv_rows(csv_named(tmp_path / "out", "results"))[1]) == size
        lines = limits_lines(summary_table(book))
        if size <= 3:
            assert lines == ["No rows, columns or cell text shortened"]
        else:
            assert "Results: 2 rows omitted; 2 retained." in lines
        assert section(summary_table(book), "file records") == {"ok": size, "partial_fail": 0, "fail": 0, "skipped": 0}
    finally:
        book.close()


def resource_database(db, kind, size, payload=0):
    ai = kind in ("ai", "wrapped", "capped", "recovery")
    files, pages = (1, size) if kind == "giant" else (size, 10 if kind == "ai" else 1)
    note = ("line of report text\n" * ((payload + 19) // 20))[:payload]
    controller = SQLiteDatabaseController(db)
    record_configuration(controller, ai)
    controller.close()
    with sqlite3.connect(db) as con:
        con.executemany(
            "INSERT INTO file_registry (file_id,file_path,file_ext,total_pages,file_to_jpegs_comment,"
            "page_range,range_status,processing_complete,ai_enabled,processing_error,"
            "source_size,source_mtime_ns,source_sha256) "
            "VALUES (?, ?, '.pdf', ?, ?, ?, 'ok', 1, 0, '', 1000, 0, ?)",
            ((i, f"file-{i}.pdf", pages, note or "done", f"1-{pages}", f"{i:064x}")
             for i in range(1, files + 1)),
        )
        con.executemany(
            "INSERT INTO page_log VALUES (?, ?, ?, ?, 'ok', '', '')",
            (((i - 1) * pages + p, i, p, f"{i}_page_{p}.jpg")
             for i in range(1, files + 1) for p in range(1, pages + 1)),
        )
        if ai:
            con.executemany(
                "INSERT INTO llm_requests VALUES (?, ?, 1, ?, 'ok', 'raw', '')",
                ((i, i, f"1-{pages}") for i in range(1, files + 1)),
            )
            con.executemany(
                "INSERT INTO llm_answers VALUES (?, ?, ?, NULL, '', '', ?)",
                (((i - 1) * 2 + k, i, i, json.dumps({"answer": f"answer-{i}-{k}"}))
                 for i in range(1, files + 1) for k in (1, 2)),
            )
        if kind == "recovery":
            con.execute("DROP VIEW overall_result")
    return files, pages, ai, note


def resource_rows(spec, title):
    files, pages, ai, note = spec
    registry_headers = ["file_id", "processing_complete", "ai_enabled", "processing_error",
                        "file_path", "source_size", "source_mtime_ns", "source_sha256",
                        "file_ext", "total_pages",
                        "file_to_jpegs_comment", "page_range", "range_status"]
    page_headers = ["page_id", "file_id", "page_number", "output_file",
                    "page_to_jpeg_status", "page_to_jpeg_comment", "video_frame_timestamp"]
    request_headers = ["request_id", "file_id", "request_number", "pages",
                       "request_status", "raw_llm_answer", "llm_network_error"]
    answer_headers = ["llm_answer_id", "request_id", "file_id", "page_id",
                      "raw_model_page_number", "llm_error", "values_json"]
    if title in ("Results", "Master Registry", "Raw file_registry"):
        if title == "Results":
            yield ["file_id", "file_path", *(["pages"] if ai else []), "file_result",
                   *(["llm_answer"] if ai else []), "notes"]
        elif title == "Master Registry":
            yield ["file_id", "file_path", "file_ext", "total_pages", "page_range", "range_status",
                   "file_to_jpegs_comment", "file_to_jpegs_status", "file_to_llm_status", "overall_result"]
        else:
            yield registry_headers
        for i in range(1, files + 1):
            if title == "Results":
                if ai:
                    for k in (1, 2):
                        yield [i, f"file-{i}.pdf", f"1-{pages}", "ok", f"answer-{i}-{k}", note]
                else:
                    yield [i, f"file-{i}.pdf", "ok", note]
            elif title == "Master Registry":
                yield [i, f"file-{i}.pdf", ".pdf", pages, f"1-{pages}", "ok", note or "done",
                       "ok", "ok" if ai else "skipped", "ok"]
            else:
                yield [i, 1, 0, "", f"file-{i}.pdf", 1000, 0, f"{i:064x}", ".pdf", pages,
                       note or "done", f"1-{pages}", "ok"]
    elif title in ("Page Log", "Raw page_log"):
        yield (["file_id", "page_id", *page_headers[2:]] if title == "Page Log" else page_headers)
        for i in range(1, files + 1):
            for p in range(1, pages + 1):
                identity = [i, (i - 1) * pages + p]
                yield [*(identity if title == "Page Log" else identity[::-1]), p, f"{i}_page_{p}.jpg", "ok", "", ""]
    elif title in ("LLM Requests", "Raw llm_requests"):
        yield (["file_id", "request_id", *request_headers[2:]] if title == "LLM Requests" else request_headers)
        if ai:
            for i in range(1, files + 1):
                yield [i, i, 1, f"1-{pages}", "ok", "raw", ""]
    elif title in ("LLM Answers", "Raw llm_answers"):
        yield (["file_id", "page_id", "request_id", "llm_answer_id", "raw_model_page_number",
                "llm_error", "llm_answer"] if title == "LLM Answers" else answer_headers)
        if ai:
            for i in range(1, files + 1):
                for k in (1, 2):
                    answer = f"answer-{i}-{k}"
                    identity = (i - 1) * 2 + k
                    if title == "LLM Answers":
                        yield [i, None, i, identity, "", "", answer]
                    else:
                        yield [identity, i, i, None, "", "", json.dumps({"answer": answer})]
    elif title == "Raw run_configuration":
        yield ["configuration_id", "ai_enabled", "source_root", "processing_digest", "output_mode", "declared_columns"]
        yield [1, int(ai), "", "", "table_per_file" if ai else "", '["answer"]' if ai else "[]"]
    else:
        raise AssertionError(f"Unknown expected report: {title}")


def assert_resource_output(outcome, spec, *, workbook=True, recovery=False, row_cap=None):
    import csv
    from itertools import zip_longest

    expected = ["Results", "Master Registry", "Page Log"]
    if spec[2]:
        expected += ["LLM Requests", "LLM Answers"]
    if recovery:
        expected = [name for name in expected if name not in ("Results", "Master Registry")]
        assert {(f["report"], f["format"]) for f in outcome.failed} == {
            ("Results", "csv"), ("Master Registry", "csv"), ("Workbook", "xlsx")}
    else:
        assert not outcome.failed, outcome.failed
    assert outcome.recovery == recovery
    wanted_inventory = [(title, "csv") for title in expected]
    if workbook:
        wanted_inventory.append(("Recovery workbook" if recovery else "Workbook", "xlsx"))
    assert [(p["report"], p["format"]) for p in outcome.saved] == wanted_inventory, "report inventory"
    assert len({p["path"] for p in outcome.saved}) == len(outcome.saved)

    def compare(actual, wanted, title):
        for number, (got, want) in enumerate(zip_longest(actual, wanted), 1):
            assert got is not None and want is not None, (title, number, "row count")
            assert [cell_text(v) for v in got] == [cell_text(v) for v in want], (title, number, got, want)

    for item in outcome.saved:
        assert Path(item["path"]).is_file(), item
        if item["format"] == "csv":
            with Path(item["path"]).open(encoding="utf-8", newline="") as handle:
                compare(csv.reader(handle), resource_rows(spec, item["report"]), item["report"])
    if not workbook:
        return
    book = openpyxl.load_workbook(outcome.saved[-1]["path"], read_only=True)
    try:
        titles = (["Raw file_registry", "Raw page_log", "Raw llm_requests", "Raw llm_answers",
                   "Raw run_configuration"] if recovery else expected)
        expected_sheets = ["Summary", *titles] if recovery else ["Summary", titles[0], "Details>>>", *titles[1:]]
        assert book.sheetnames == expected_sheets
        if not recovery:
            assert list(book["Details>>>"].values) == []
            table = summary_table(book)
            assert section(table, "file records")["ok"] == spec[0]
            assert section(table, "pages")["ok"] == spec[0] * spec[1]
            if spec[2]:
                assert section(table, "requests")["ok"] == spec[0]
            assert section(table, "export") == {
                p["report"]: Path(p["path"]).name for p in outcome.saved if p["format"] == "csv"}
        else:
            assert "Recovery report" in str(list(book["Summary"].values))
        for title in titles:
            rows = resource_rows(spec, title)
            if row_cap is not None:
                header, *records = rows
                rows = iter([header, *records[:row_cap - 1], ["[ROWS OMITTED]"] * len(header)])
            compare(book[title].values, rows, title)
    finally:
        book.close()


def resource_export(exporter, directory, kind):
    try:
        return (exporter.export_csv if kind == "giant" else exporter.export_all_formats)(directory)
    except data_exporter.ExportError as exc:
        if kind != "recovery":
            raise
        return exc.outcome


def assert_resource_growth(tmp_path, kind, mutation="healthy"):
    import gc
    from openpyxl.worksheet._write_only import WriteOnlyWorksheet

    sizes = (2048, 16384) if kind == "giant" else (64, 512)
    peaks = []
    real_csv = data_exporter.SQLiteDataExporter._write_csv
    real_book = data_exporter.SQLiteDataExporter._workbook
    real_append = WriteOnlyWorksheet.append
    for size in sizes:
        folder = tmp_path / str(size)
        folder.mkdir()
        db = folder / "state.db"
        spec = resource_database(db, kind, size, payload=0 if kind == "giant" else 4096)
        retained = []

        def csv_writer(self, target, headers, rows):
            if mutation == "retain-rows":
                rows = iter(list(rows))
            return real_csv(self, target, headers, rows)

        def book_writer(self, target, reports, summary, outcome, **kwargs):
            if mutation == "retain-rows":
                copies = []
                for report in reports:
                    stored = list(report.rows())
                    copies.append(data_exporter.Report(report.title, report.stem, report.headers,
                                  lambda stored=stored: iter(stored)))
                reports = copies
            return real_book(self, target, reports, summary, outcome, **kwargs)

        def append(sheet, row, retained=retained):
            if mutation == "retain-cells":
                retained.append(tuple(row))
            return real_append(sheet, row)

        with pytest.MonkeyPatch.context() as patch:
            patch.setattr(data_exporter.SQLiteDataExporter, "_write_csv", csv_writer)
            patch.setattr(data_exporter.SQLiteDataExporter, "_workbook", book_writer)
            patch.setattr(WriteOnlyWorksheet, "append", append)
            if kind == "capped":
                patch.setattr(data_exporter, "EXCEL_MAX_ROWS", 8)
            gc.collect()
            tracemalloc.start()
            try:
                outcome = resource_export(make_exporter(db), folder / "out", kind)
                peaks.append(tracemalloc.get_traced_memory()[1])
            finally:
                tracemalloc.stop()
        assert_resource_output(outcome, spec, workbook=kind != "giant", recovery=kind == "recovery",
                               row_cap=8 if kind == "capped" else None)
    growth = peaks[1] - peaks[0]
    assert growth < 512 * 1024, f"retained export allocation grew: {kind}, sizes={sizes}, peaks={peaks}"


@pytest.mark.parametrize("kind", ["giant", "ai", "wrapped", "capped", "recovery"])
def test_report_allocation_does_not_grow_with_retained_rows(tmp_path, kind):
    assert_resource_growth(tmp_path, kind)


@pytest.mark.parametrize("kind", ["files", "giant", "ai", "wrapped", "capped", "recovery"])
def test_allocation_guard_rejects_whole_report_materialization(tmp_path, kind):
    with pytest.raises(AssertionError, match="retained export allocation grew"):
        assert_resource_growth(tmp_path, kind, "retain-rows")


@pytest.mark.parametrize("kind", ["files", "wrapped", "recovery"])
def test_allocation_guard_rejects_retained_workbook_cells(tmp_path, kind):
    with pytest.raises(AssertionError, match="retained export allocation grew"):
        assert_resource_growth(tmp_path, kind, "retain-cells")


def resource_work(tmp_path, size, mutation="healthy"):
    db = tmp_path / "state.db"
    spec = resource_database(db, "ai", size)
    steps = 0
    parses = 0
    real_connect = sqlite3.connect
    real_values = data_exporter.SQLiteDataExporter._values

    class CountingConnection(sqlite3.Connection):
        def execute(self, sql, parameters=(), /):
            if mutation == "scan" and parameters and "FROM page_log WHERE file_id=?" in sql:
                sql = sql.replace("FROM page_log WHERE", "FROM page_log NOT INDEXED WHERE")
            return super().execute(sql, parameters)

    def connect(*args, **kwargs):
        con = real_connect(*args, **kwargs, factory=CountingConnection)

        def progress():
            nonlocal steps
            steps += 1000
            return 0

        con.set_progress_handler(progress, 1000)
        return con

    def values(raw, columns, error):
        nonlocal parses
        parses += 1
        return real_values(raw, columns, error)

    original_rows = data_exporter.SQLiteDataExporter._results_rows

    def repeated(self, con, config):
        for row in original_rows(self, con, config):
            if mutation in ("reparse", "reparse-once"):
                for _ in range(size if mutation == "reparse" else 1):
                    self._values('{"answer":"unchanged"}', ["answer"], "")
            yield row

    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(data_exporter.sqlite3, "connect", connect)
        patch.setattr(data_exporter.SQLiteDataExporter, "_values", staticmethod(values))
        patch.setattr(data_exporter.SQLiteDataExporter, "_results_rows", repeated)
        outcome = make_exporter(db).export_csv(tmp_path / "out")
    assert_resource_output(outcome, spec, workbook=False)
    return steps, parses


def assert_resource_work_growth(tmp_path, mutation):
    counts = []
    for size in (100, 400):
        folder = tmp_path / str(size)
        folder.mkdir()
        counts.append(resource_work(folder, size, mutation))
        assert counts[-1][1] <= 4 * size, ("answer parsing", size, counts[-1])
    for index, label in enumerate(("database instructions", "answer parsing")):
        assert 0 < counts[0][index] <= counts[1][index] < 8 * counts[0][index], (label, counts)


@pytest.mark.parametrize("mutation", ["healthy", "scan", "reparse", "reparse-once"])
def test_report_work_scales_without_repeated_scans_or_parsing(tmp_path, mutation):
    if mutation == "healthy":
        assert_resource_work_growth(tmp_path, mutation)
    else:
        with pytest.raises(AssertionError, match=r"database instructions|answer parsing"):
            assert_resource_work_growth(tmp_path, mutation)


@pytest.mark.parametrize("damage", ["healthy", "missing-report", "missing-row", "changed-cell", "reordered-rows"])
def test_resource_oracle_checks_inventory_and_every_detail(tmp_path, damage):
    import csv

    db = tmp_path / "state.db"
    spec = resource_database(db, "ai", 3)
    outcome = make_exporter(db).export_all_formats(tmp_path / "out")
    if damage == "missing-report":
        outcome.saved = [p for p in outcome.saved if p["report"] != "Page Log"]
    elif damage != "healthy":
        path = Path(next(p["path"] for p in outcome.saved if p["report"] == "Page Log"))
        with path.open(encoding="utf-8", newline="") as f:
            rows = list(csv.reader(f))
        if damage == "missing-row":
            rows.pop()
        elif damage == "changed-cell":
            rows[2][3] = "wrong.jpg"
        else:
            rows[1], rows[2] = rows[2], rows[1]
        with path.open("w", encoding="utf-8", newline="") as f:
            csv.writer(f).writerows(rows)
    if damage == "healthy":
        assert_resource_output(outcome, spec)
    else:
        with pytest.raises(AssertionError, match=r"inventory|Page Log"):
            assert_resource_output(outcome, spec)


@pytest.mark.parametrize("title", ["Results", "Master Registry", "Page Log", "LLM Requests", "LLM Answers",
                                    "Summary", "Details>>>", "Raw file_registry", "Raw page_log",
                                    "Raw llm_requests", "Raw llm_answers", "Raw run_configuration"])
def test_resource_oracle_rejects_damage_in_every_worksheet(tmp_path, title):
    recovery = title.startswith("Raw ")
    kind = "recovery" if recovery else "ai"
    db = tmp_path / "state.db"
    spec = resource_database(db, kind, 3)
    outcome = resource_export(make_exporter(db), tmp_path / "out", kind)
    assert_resource_output(outcome, spec, recovery=recovery)
    path = outcome.saved[-1]["path"]
    book = openpyxl.load_workbook(path)
    try:
        sheet = book[title]
        if title == "Summary":
            for row in sheet:
                if row[0].value == "Fully processed":
                    count_cell = row[1]
                    assert isinstance(count_cell, Cell)
                    count_cell.value = 999
                    break
            else:
                raise AssertionError("missing file count")
        elif title == "Details>>>":
            sheet["A1"] = "unexpected record"
        else:
            sheet["B2"] = "corrupt record"
        book.save(path)
    finally:
        book.close()
    with pytest.raises(AssertionError):
        assert_resource_output(outcome, spec, recovery=recovery)


def benchmark_sample(folder, kind, size, slow=False):
    db = folder / "state.db"
    spec = resource_database(db, kind, size)
    phases = {"csv": 0.0, "xlsx": 0.0}

    class MeasuredExporter(data_exporter.SQLiteDataExporter):
        def _write_csv(self, target_path, headers, rows):
            start = time.perf_counter()
            try:
                return super()._write_csv(target_path, headers, rows)
            finally:
                phases["csv"] += time.perf_counter() - start

        def _workbook(self, *args, **kwargs):
            start = time.perf_counter()
            try:
                return super()._workbook(*args, **kwargs)
            finally:
                phases["xlsx"] += time.perf_counter() - start

        def _results_rows(self, con, config):
            previous = []
            for row in super()._results_rows(con, config):
                if slow:
                    previous.append(row)
                    for old in previous:
                        json.dumps(old, ensure_ascii=False)
                yield row

    assert not tracemalloc.is_tracing(), "speed measurements must not trace allocations"
    start, cpu = time.perf_counter(), time.process_time()
    exporter = MeasuredExporter(db)
    outcome = (exporter.export_all_formats if kind == "files" else exporter.export_csv)(folder / "out")
    result = {"kind": kind, "size": size, "seconds": time.perf_counter() - start,
              "cpu": time.process_time() - cpu, **phases}
    assert_resource_output(outcome, spec, workbook=kind == "files")
    return result


def assert_benchmark_scaling(samples):
    from statistics import median

    sizes = sorted({s["size"] for s in samples})
    assert len(sizes) >= 2
    assert sizes[-1] == 4 * sizes[0]
    times = [median(s["seconds"] for s in samples if s["size"] == size) for size in sizes]
    assert times[0] > 0
    for index in range(1, len(sizes)):
        factor = sizes[index] / sizes[index - 1]
        assert times[index] < 2 * factor * times[index - 1], (
            f"export time scaling: sizes={sizes}, medians={times}")
    assert times[-1] < 8 * times[0], f"export time scaling: sizes={sizes}, medians={times}"


def run_benchmark_child(command, scratch, timeout=120):
    import os
    import subprocess

    env = dict(os.environ, TEMP=str(scratch), TMP=str(scratch))
    return subprocess.run(command, env=env, capture_output=True, text=True, check=True, timeout=timeout)


@pytest.mark.parametrize("bad", [False, True])
def test_benchmark_verdict_separates_proportional_and_quadratic_work(bad):
    samples = [{"size": n, "seconds": (n / 1000) ** (2 if bad else 1) * offset}
               for n in (1000, 2000, 4000) for offset in (0.9, 1.0, 1.1)]
    if bad:
        with pytest.raises(AssertionError, match="export time scaling"):
            assert_benchmark_scaling(samples)
    else:
        assert_benchmark_scaling(samples)


def test_benchmark_watchdog_reaps_child_and_owns_its_temporary_files(tmp_path):
    import subprocess
    import tempfile

    target = tmp_path / "owned"
    with tempfile.TemporaryDirectory(dir=tmp_path) as scratch:
        script = ("import os,tempfile,time; from pathlib import Path; "
                  f"p=Path({str(target)!r}); p.write_text(tempfile.gettempdir()); "
                  "f=open(Path(tempfile.gettempdir()) / 'owned.xml','w'); "
                  "f.write('owned'); f.flush(); time.sleep(60)")
        with pytest.raises(subprocess.TimeoutExpired):
            run_benchmark_child([sys.executable, "-c", script], scratch, timeout=3)
        assert target.read_text() == scratch
        files = list(Path(scratch).iterdir())
        assert files
        for path in files:
            path.unlink()
    assert not Path(scratch).exists()


def report_benchmark():
    import importlib.metadata
    import os
    import platform
    import tempfile

    if len(sys.argv) > 1 and sys.argv[1] == "--sample":
        folder, kind, size, slow = sys.argv[2:]
        print(json.dumps(benchmark_sample(Path(folder), kind, int(size), slow == "slow")))
        return
    assert sys.argv[1:] == ["--benchmark"], "use --benchmark for the required sequential report check"
    print(json.dumps({"python": sys.version, "os": platform.platform(), "cpus": os.cpu_count(),
                      "sqlite": sqlite3.sqlite_version, "openpyxl": importlib.metadata.version("openpyxl"),
                      "samples_per_size": 3, "tracing": False}), flush=True)
    with tempfile.TemporaryDirectory(prefix="report-benchmark-") as root:
        for kind, sizes, slow in [("files", (1000, 2000, 4000), False),
                                  ("ai", (1000, 2000, 4000), False),
                                  ("giant", (50000, 100000, 200000), False),
                                  ("ai", (500, 2000), True)]:
            samples = []
            for size in sizes:
                for repeat in range(3):
                    folder = Path(root) / f"{kind}-{size}-{repeat}-{slow}"
                    folder.mkdir()
                    result = run_benchmark_child(
                        [sys.executable, str(Path(__file__).resolve()), "--sample", str(folder),
                         kind, str(size), "slow" if slow else "healthy"], folder)
                    sample = json.loads(result.stdout)
                    samples.append(sample)
                    print(json.dumps({**sample, "repeat": repeat + 1, "control": slow}), flush=True)
            if slow:
                with pytest.raises(AssertionError, match="export time scaling"):
                    assert_benchmark_scaling(samples)
                print("Repeated-prefix control rejected", flush=True)
            else:
                assert_benchmark_scaling(samples)
    print("Report benchmark passed: healthy scaling and broken control verified", flush=True)


if __name__ == "__main__":
    report_benchmark()
