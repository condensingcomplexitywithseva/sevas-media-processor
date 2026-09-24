# Copyright 2026 Vsevolod Belonogov
# SPDX-License-Identifier: Apache-2.0

import csv
import json
import sqlite3
import sys
from pathlib import Path
from typing import Any

SRC = Path(__file__).resolve().parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from data_exporter import (LLM_ANSWERS_REPORT_COLUMNS,
                           SQLiteDataExporter)
from db_controller import SQLiteDatabaseController
from schemas import AnswerRow, RequestOutcome, FileSummary, PageResult
from test_results_report import excel_value

DECLARED = ["genre", "answer"]
SPREAD_HEADERS = [*LLM_ANSWERS_REPORT_COLUMNS, "llm_genre", "llm_answer"]
RESULTS_SHEET_TITLE = "Results"
RESULTS_HEADERS = ["file_id", "file_path", "pages", "file_result", "llm_genre", "llm_answer", "notes"]


def seed_db(tmp_path, answer_values, page_count=2, declared=DECLARED):
    db_path = tmp_path / "state.db"
    controller = SQLiteDatabaseController(db_path)
    controller.handle_file_started(1, "video.mp4", ".mp4", "Stub")
    for number in range(1, page_count + 1):
        controller.handle_frame_saved(1, PageResult(
            number, f"1_page_{number}.jpg", "ok", "", capture_seconds=None))
    controller.handle_file_completed(
        1, FileSummary(page_count, f"1-{page_count}", "ok", "done"))
    controller.finalize_file(1)
    controller.handle_llm_requests(1, [
        RequestOutcome(1, tuple(range(1, page_count + 1)), "ok", "raw reply", "")])
    controller.record_run_configuration(True, "table_per_page", declared)
    controller.close()

    connection = sqlite3.connect(db_path)
    try:
        request_id = connection.execute(
            "SELECT request_id FROM llm_requests").fetchone()[0]
        pages = dict(connection.execute(
            "SELECT page_number, page_id FROM page_log").fetchall())
        for page_number, llm_error, values_json in answer_values:
            connection.execute(
                "INSERT INTO llm_answers (request_id, file_id, page_id,"
                " raw_model_page_number, llm_error, values_json)"
                " VALUES (?, 1, ?, ?, ?, ?)",
                (request_id,
                 pages.get(page_number),
                 "" if page_number is None else str(page_number),
                 llm_error, values_json))
        connection.commit()
    finally:
        connection.close()
    return db_path


def read_csv(path):
    with open(path, encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def csv_headers(path):
    with open(path, encoding="utf-8", newline="") as handle:
        return next(csv.reader(handle))


def export_csvs(db_path, out_dir):
    exporter = SQLiteDataExporter(db_path)
    exporter.export_csv(out_dir, timestamp="t")
    return out_dir / "llm_answers_t.csv"


def export_xlsx(db_path, out_dir):
    import openpyxl

    exporter = SQLiteDataExporter(db_path)
    exporter.export_all_formats(out_dir)
    workbook_path = next(out_dir.glob("database_export_*.xlsx"))
    return openpyxl.load_workbook(workbook_path, read_only=False)


def sheet_rows(workbook, title) -> list[list[Any]]:
    sheet = workbook[title]
    return [[excel_value(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]



def test_llm_answers_csv_spreads_declared_columns(tmp_path):
    db_path = seed_db(tmp_path, [
        (1, "", json.dumps({"genre": "war", "answer": "long text"})),
        (2, "no_row_returned", ""),
    ])
    answers_csv = export_csvs(db_path, tmp_path / "out")

    assert csv_headers(answers_csv) == SPREAD_HEADERS
    rows = read_csv(answers_csv)
    assert len(rows) == 2
    assert rows[0]["llm_genre"] == "war"
    assert rows[0]["llm_answer"] == "long text"
    assert rows[1]["llm_genre"] == "" and rows[1]["llm_answer"] == ""
    assert "values_json" not in rows[0]


def test_non_scalar_values_are_json_encoded_into_cells(tmp_path):
    db_path = seed_db(tmp_path, [
        (1, "", json.dumps({"genre": ["war", "drama"],
                            "answer": {"nested": True}})),
    ])
    answers_csv = export_csvs(db_path, tmp_path / "out")

    row = read_csv(answers_csv)[0]
    assert json.loads(row["llm_genre"]) == ["war", "drama"]
    assert json.loads(row["llm_answer"]) == {"nested": True}


def test_formula_shaped_values_round_trip_the_csv_verbatim(tmp_path):
    hostile = "=HYPERLINK(\"http://evil.example\",\"click\")"
    db_path = seed_db(tmp_path, [
        (1, "", json.dumps({"genre": hostile, "answer": "+1-2"})),
    ])
    answers_csv = export_csvs(db_path, tmp_path / "out")

    row = read_csv(answers_csv)[0]
    assert row["llm_genre"] == hostile
    assert row["llm_answer"] == "+1-2"


def test_llm_answers_sheet_spreads_the_same_columns(tmp_path):
    db_path = seed_db(tmp_path, [
        (1, "", json.dumps({"genre": "war", "answer": "text"})),
    ])
    workbook = export_xlsx(db_path, tmp_path / "out")
    rows = sheet_rows(workbook, "LLM Answers")
    assert rows[0] == SPREAD_HEADERS
    header_index = dict(zip(rows[0], range(len(rows[0])), strict=True))
    assert rows[1][header_index["llm_genre"]] == "war"



def test_reader_sheet_is_one_row_per_answer_row(tmp_path):
    db = seed_db(tmp_path, [(1, "", json.dumps({"genre":"war", "answer":"a1"})),
        (2, "no_row_returned", ""),
        (None, "hallucinated_page_number", json.dumps({"genre":"g9", "answer":"a9"}))])
    rows = sheet_rows(export_xlsx(db, tmp_path / "out"), "Results")
    assert rows[0] == ["file_id", "file_path", "pages", "file_result", "llm_genre", "llm_answer", "notes"]
    assert len(rows) == 4
    assert [(r[2], r[4]) for r in rows[1:]] == [(1,"war"), (2,None), (None,"g9")]
    assert "No answer returned" in rows[2][-1] and "Unverified page claim" in rows[3][-1]


def test_results_per_file_rows_list_the_pages_actually_sent(tmp_path):
    db_path = tmp_path / "state.db"
    controller = SQLiteDatabaseController(db_path)
    controller.handle_file_started(1, "doc.pdf", ".pdf", "Stub")
    for number in range(1, 8):
        if number == 6:
            controller.handle_frame_saved(1, PageResult(6, "", "failure", "decoder choked"))
        else:
            controller.handle_frame_saved(1, PageResult(
                number, f"1_page_{number}.jpg", "ok", ""))
    controller.handle_file_completed(1, FileSummary(7, "1-7", "ok", "done"))
    controller.finalize_file(1)
    controller.handle_llm_requests(1, [
        RequestOutcome(1, (1, 2, 3), "ok", "r1", "",
                     answer_rows=(AnswerRow(None, "", "", json.dumps({"genre": "g", "answer": "a"})),)),
        RequestOutcome(2, (4, 5, 7), "ok", "r2", "",
                     answer_rows=(AnswerRow(None, "", "", json.dumps({"genre": "h", "answer": "b"})),)),
    ])
    controller.record_run_configuration(True, 'table_per_file', DECLARED)
    controller.close()

    workbook = export_xlsx(db_path, tmp_path / "out")
    rows = sheet_rows(workbook, RESULTS_SHEET_TITLE)
    assert [(r[2], r[4]) for r in rows[1:]] == [("1-3", "g"), ("4-5, 7", "h")]
    assert all("failed conversion" in r[-1] for r in rows[1:])


def test_stored_page_cell_only_converts_canonical_single_page():
    for value, expected in [("", ""), ("4", 4), ("1-3", "1-3"), ("1, 3", "1, 3"), ("003", "003")]:
        assert SQLiteDataExporter._page_cell(value) == expected


def test_results_has_one_row_per_file_when_ai_is_off(tmp_path):
    db_path = tmp_path / "state.db"
    controller = SQLiteDatabaseController(db_path)
    controller.handle_file_started(1, "img.png", ".png", "Stub")
    for number in (1, 2):
        controller.handle_frame_saved(1, PageResult(
            number, f"1_page_{number}.jpg", "ok", ""))
    controller.handle_file_completed(1, FileSummary(2, "1-2", "ok", "done"))
    controller.finalize_file(1)
    controller.record_run_configuration(False, "", ())
    controller.close()

    workbook = export_xlsx(db_path, tmp_path / "out")
    rows = sheet_rows(workbook, RESULTS_SHEET_TITLE)
    assert rows[0] == ["file_id", "file_path", "file_result", "notes"]
    assert len(rows) == 2 and rows[1][1] == "img.png"


def test_results_sheet_is_always_present(tmp_path):
    db_path = seed_db(tmp_path, [
        (1, "", json.dumps({"genre": "war", "answer": "a"})),
    ])
    workbook = export_xlsx(db_path, tmp_path / "out")
    assert RESULTS_SHEET_TITLE in workbook.sheetnames
    assert "Joined Answers" not in workbook.sheetnames


def test_results_csv_is_always_present_without_a_legacy_joined_report(tmp_path):
    db_path = seed_db(tmp_path, [
        (1, "", json.dumps({"genre": "war", "answer": "a"})),
    ])
    out_dir = tmp_path / "out"
    export_csvs(db_path, out_dir)
    names = {p.name for p in out_dir.iterdir()}
    assert "results_t.csv" in names
    assert not any("joined" in name.lower() for name in names), names



def test_control_characters_in_spread_cells_do_not_kill_the_workbook(tmp_path):
    hostile = "bell \x07 and null \x00 chars"
    db_path = seed_db(tmp_path, [
        (1, "", json.dumps({"genre": hostile, "answer": "fine"})),
    ])
    workbook = export_xlsx(db_path, tmp_path / "out")
    for title in ("LLM Answers", RESULTS_SHEET_TITLE):
        rows = sheet_rows(workbook, title)
        index = dict(zip(rows[0], range(len(rows[0])), strict=True))
        cell = rows[1][index["llm_genre"]]
        assert "\x07" not in cell and "\x00" not in cell
        assert "�" in cell

    answers_csv = export_csvs(db_path, tmp_path / "out2")
    assert hostile in answers_csv.read_text(encoding="utf-8")


def test_oversized_spread_cells_are_cut_with_the_visible_marker(tmp_path):
    huge = "x" * 40_000
    db_path = seed_db(tmp_path, [
        (1, "", json.dumps({"genre": "g", "answer": huge})),
    ])
    workbook = export_xlsx(db_path, tmp_path / "out")
    for title in ("LLM Answers", RESULTS_SHEET_TITLE):
        rows = sheet_rows(workbook, title)
        index = dict(zip(rows[0], range(len(rows[0])), strict=True))
        cell = rows[1][index["llm_answer"]]
        assert len(cell) <= 32_767
        assert cell.startswith("[TEXT SHORTENED] ")



def test_full_run_lands_answers_in_db_views_and_exports(tmp_path):
    import logging
    from types import SimpleNamespace

    from PIL import Image

    from batch_orchestrator import BatchOrchestrator
    from media_classifier import MediaClassifier
    from schemas import Status
    from fake_llm import make_server
    from fake_llm.generic import openai_reply
    from fake_llm.harness import build_client

    input_dir = tmp_path / "input"
    input_dir.mkdir()
    (input_dir / "clip.mp4").write_bytes(b"x")
    output_dir = tmp_path / "output"
    output_dir.mkdir()

    class JpegRouter:
        output_folder = output_dir
        relative_or_orphan = staticmethod(MediaClassifier.relative_or_orphan)

        def evaluate_and_route(self, file_id, path, root):
            rel, orphaned = self.relative_or_orphan(path, root)

            def gen():
                for number in (1, 2, 3):
                    name = f"{file_id}_page_{number}.jpg"
                    Image.new("RGB", (8, 8), (number * 20, 0, 0)).save(
                        output_dir / name, "JPEG")
                    yield PageResult(number, name, Status.OK.value, "")
                return FileSummary(3, "1-3", Status.OK.value, "done")

            return rel, path.suffix, "Stub", gen(), orphaned

    class StubLogger:
        app_logger = logging.getLogger("test-llm-answer-e2e")

        def log_file_started(self, *args):
            pass

        def log_frame_saved(self, *args):
            pass

        def log_file_completed(self, *args):
            pass

        def log_llm_completed(self, *args):
            pass

        def log_critical_error(self, *args):
            pass

    srv = make_server("openai").start()
    try:
        srv.queue(json=openai_reply("sorry, here it is in prose"))
        srv.queue(json=openai_reply(json.dumps([
            {"page": 1, "genre": "war", "answer": "a1"},
            {"page": 3, "genre": "peace", "answer": "a3"},
        ])))
        client = build_client(
            "openai", srv.base_url,
            LLM_OUTPUT_MODE="table_per_page",
            LLM_OUTPUT_COLUMNS="genre,answer",
            LLM_OUTPUT_COLUMN_INSTRUCTIONS="",
            LLM_JSON_MAX_ATTEMPTS=3,
            LLM_ABORT_ON_MALFORMED_JSON=False,
        )
        settings = SimpleNamespace(
            INPUT_FOLDER_PATH=input_dir,
            NO_RETRY_STATUSES=[Status.OK],
            JPEG_QUALITY=90,
            MAX_DIMENSION=4096,
            ENABLE_LLM_INFERENCE=True,
            MAX_CONSECUTIVE_LLM_FAILURES=3,
            MAX_JPEGS_PER_INFERENCE=10,
        )
        db_path = tmp_path / "state.db"
        controller = SQLiteDatabaseController(db_path)
        BatchOrchestrator(
            settings, controller, JpegRouter(), StubLogger(), client
        ).execute_batch_processing_loop()

        statuses = controller.get_file_statuses(1)
        controller.record_run_configuration(True, "table_per_page", DECLARED)
        controller.close()
        assert len(srv.requests) == 2
    finally:
        srv.stop()

    assert statuses["file_to_jpegs_status"] == "ok"
    assert statuses["file_to_llm_status"] == "partial_llm_failure"
    assert statuses["overall_result"] == "partial_fail"

    connection = sqlite3.connect(db_path)
    try:
        stored = connection.execute(
            "SELECT llm_error, values_json FROM llm_answers"
            " ORDER BY llm_answer_id").fetchall()
        request_row = connection.execute(
            "SELECT request_status, llm_network_error FROM llm_requests"
        ).fetchone()
    finally:
        connection.close()
    assert request_row[0] == "ok"
    assert "attempt 1" in request_row[1]
    errors = sorted(row[0] for row in stored)
    assert errors == ["", "", "no_row_returned"]

    answers_csv = export_csvs(db_path, tmp_path / "reports")
    rows = read_csv(answers_csv)
    assert len(rows) == 3
    genres = {row["llm_genre"] for row in rows}
    assert genres == {"war", "peace", ""}

RESULTS_REPORT_COLUMNS = ["file_id", "file_path", "file_result", "notes"]
