# Copyright 2026 Vsevolod Belonogov
# SPDX-License-Identifier: Apache-2.0


import csv
import itertools
import json
import re
import logging
import threading
import sqlite3
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import Cell
import pytest

SRC = Path(__file__).resolve().parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))
import data_exporter
import schemas
from fake_llm import make_server
from db_controller import SQLiteDatabaseController
from schemas import FileSummary, PageResult

RESULTS_TITLE = "Results"
RESULTS_STEM = "results"
FIXED_COLUMNS = ["file_id", "file_path", "pages", "file_result", "notes"]
REQUEST_STATUSES = tuple(s.value for s in schemas.RequestStatus)


def new_name(module, name: str, ruling: str):
    value = getattr(module, name, None)
    assert value is not None, f"{module.__name__}.{name} must exist ({ruling})"
    return value


def request_outcome(**fields):
    outcome_type = new_name(schemas, "RequestOutcome", "the request vocabulary with recorded pages")
    values: dict[str, Any] = {
        "request_number": 1,
        "pages": (),
        "status": "ok",
        "raw_answer": "",
        "error": "",
        "answer_rows": (),
    }
    values.update(fields)
    return outcome_type(**values)


def answer_row(page_number=None, raw_claim="", llm_error="", values=None):
    return schemas.AnswerRow(
        page_number, raw_claim, llm_error, "" if values is None else json.dumps(values, ensure_ascii=False)
    )


def store_requests(controller, file_id: int, outcomes) -> None:
    handler = getattr(controller, "handle_llm_requests", None)
    assert handler is not None, "SQLiteDatabaseController.handle_llm_requests must exist (the request vocabulary)"
    handler(file_id, outcomes)


def record_configuration(controller, ai_on: bool, declared=("answer",), output_mode="table_per_file") -> None:
    if ai_on:
        controller.record_run_configuration(True, output_mode, tuple(declared))
    else:
        controller.record_run_configuration(False, "", ())


def add_file(
    controller, file_id, path, pages, total=None, page_range=None, range_status="ok", comment="done", extension=None
):
    controller.handle_file_started(file_id, path, extension or Path(path).suffix, "Stub")
    for number, output_file, status, page_comment, capture in pages:
        controller.handle_frame_saved(
            file_id, PageResult(number, output_file, status, page_comment, capture_seconds=capture)
        )
    if total is None:
        total = len(pages)
    if page_range is None:
        page_range = f"1-{total}" if total > 1 else ("1" if total == 1 else "")
    controller.handle_file_completed(file_id, FileSummary(total, page_range, range_status, comment))
    controller.finalize_file(file_id)


def seed_acceptance_example(db_path: Path) -> None:
    c = SQLiteDatabaseController(db_path)
    add_file(
        c,
        1,
        "notes.txt",
        [(1, "", "failure", "Unsupported file extension: .txt", None)],
        total=0,
        page_range="",
        range_status="failure",
        comment="Unsupported file extension: .txt",
    )
    add_file(
        c,
        2,
        "scan.pdf",
        [
            (1, "2_scan_page_1.jpg", "ok", "", None),
            (2, "", "failure", "pdfium: render failed (corrupt object stream)", None),
            (3, "2_scan_page_3.jpg", "ok", "", None),
        ],
        total=3,
        page_range="1-3",
        comment="Saved 2 of 3 | Details: pdfium: render failed",
    )
    store_requests(
        c,
        2,
        [
            request_outcome(
                request_number=1,
                pages=(1, 3),
                status="ok",
                raw_answer='{"answer": "invoice"}',
                answer_rows=(answer_row(values={"answer": "invoice"}),),
            )
        ],
    )
    add_file(
        c,
        3,
        "broken.heic",
        [(1, "", "failure", "Catastrophic Image error: cannot identify image file", None)],
        total=0,
        page_range="",
        range_status="failure",
        comment="Catastrophic Image error: cannot identify image file",
    )
    add_file(
        c,
        4,
        "receipt.png",
        [(1, "4_receipt_page_1.jpg", "ok", "", None)],
        total=1,
        page_range="1",
        comment="Successfully saved all 1 requested frames",
    )
    store_requests(
        c,
        4,
        [
            request_outcome(
                request_number=1,
                pages=(1,),
                status="network_failure",
                error="FATAL AUTHENTICATION ERROR: HTTP 401 from provider",
            )
        ],
    )
    add_file(
        c,
        5,
        "clip.mp4",
        [(1, "5_clip_page_1_t00_00_00_00.jpg", "ok", "", 0.0), (2, "5_clip_page_2_t00_00_05_00.jpg", "ok", "", 5.0)],
        total=300,
        page_range="00:00:00-00:00:10",
        comment="Successfully saved all 2 requested frames",
    )
    store_requests(
        c,
        5,
        [
            request_outcome(
                request_number=1,
                pages=(1,),
                status="ok",
                raw_answer='{"answer": "a cat"}',
                error="Base64 Encoding Failure [5_clip_page_2_t00_00_05_00.jpg]: [Errno 2] No such file",
                answer_rows=(answer_row(values={"answer": "a cat"}),),
            )
        ],
    )
    record_configuration(c, ai_on=True, declared=("answer",))
    c.close()


def make_exporter(db_path):
    return data_exporter.SQLiteDataExporter(db_path)


def export_results_csv(db_path, out_dir) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    exporter = make_exporter(db_path)
    exporter.export_csv(out_dir, timestamp="t")
    path = out_dir / f"{RESULTS_STEM}_t.csv"
    assert path.exists(), f"the export wrote no {path.name} (Results is always written)"
    return path


def export_workbook(db_path, out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    exporter = make_exporter(db_path)
    exporter.export_all_formats(out_dir)
    return openpyxl.load_workbook(next(out_dir.glob("database_export_*.xlsx")))


def read_csv_rows(path) -> tuple[list[str], list[dict[str, str]]]:
    old_limit = csv.field_size_limit(1 << 30)
    try:
        with open(path, encoding="utf-8", newline="") as handle:
            headers = next(csv.reader(handle))
        with open(path, encoding="utf-8", newline="") as handle:
            rows = list(csv.DictReader(handle))
    finally:
        csv.field_size_limit(old_limit)
    return headers, rows


def excel_value(value):
    if not isinstance(value, str):
        return value
    return re.sub(r"_x([0-9a-fA-F]{4})_", lambda match: chr(int(match[1], 16)), value)


def sheet_rows(workbook, title) -> list[list[Any]]:
    assert title in workbook.sheetnames, (title, workbook.sheetnames)
    return [[excel_value(cell) for cell in row] for row in workbook[title].iter_rows(values_only=True)]


def results_sheet(workbook) -> tuple[list[str], list[dict[str, Any]]]:
    rows = sheet_rows(workbook, RESULTS_TITLE)
    headers = [str(cell) for cell in rows[0]]
    return headers, [dict(zip(headers, row, strict=True)) for row in rows[1:]]


def sql(db_path, statement, *params):
    connection = sqlite3.connect(db_path)
    try:
        return connection.execute(statement, params).fetchall()
    finally:
        connection.close()


def cell_text(value) -> str:
    return "" if value is None else str(value)


def seed_one_file(db_path, pages, requests=(), declared=("answer",), **file_kwargs):
    c = SQLiteDatabaseController(db_path)
    add_file(c, 1, file_kwargs.pop("path", "doc.pdf"), pages, **file_kwargs)
    if requests:
        store_requests(c, 1, list(requests))
    record_configuration(c, ai_on=bool(requests), declared=declared)
    c.close()
    return db_path


def run_through_the_real_app(tmp_path, monkeypatch, files, server, abort_flag=None, **overrides):
    from test_llm_run_cruelty import point_tokens_at_the_fake, run_core, settings_for
    import central_logger

    monkeypatch.setattr(central_logger, "setup_logging", lambda *a, **k: None)
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    for name, make in files:
        make(input_dir / name)
    point_tokens_at_the_fake(monkeypatch, "openai")
    settings = settings_for(
        (input_dir, tmp_path / "output"), "openai", server.base_url, LLM_OUTPUT_COLUMNS="answer", **overrides
    )
    return run_core(settings, abort_flag=abort_flag)


def test_acceptance_example_is_five_complete_reader_rows(tmp_path):
    db = tmp_path / "state.db"
    seed_acceptance_example(db)
    headers, rows = read_csv_rows(export_results_csv(db, tmp_path / "csv"))
    assert headers == ["file_id", "file_path", "pages", "file_result", "llm_answer", "notes"]
    assert [(r["file_path"], r["pages"], r["file_result"], r["llm_answer"]) for r in rows] == [
        ("notes.txt", "", "fail", ""),
        ("scan.pdf", "1, 3", "partial_fail", "invoice"),
        ("broken.heic", "", "fail", ""),
        ("receipt.png", "1", "partial_fail", ""),
        ("clip.mp4", "1", "ok", "a cat"),
    ]
    assert "Unsupported file extension" in rows[0]["notes"]
    assert "1 page(s) failed conversion" in rows[1]["notes"]
    assert "cannot identify image file" in rows[2]["notes"]
    assert "FATAL AUTHENTICATION ERROR" in rows[3]["notes"]
    assert "Base64 Encoding Failure" in rows[4]["notes"]
    workbook = export_workbook(db, tmp_path / "xlsx")
    sheet_headers, sheet_data = results_sheet(workbook)
    assert sheet_headers == headers
    assert [[cell_text(row[h]) for h in headers] for row in sheet_data] == [[r[h] for h in headers] for r in rows]
    assert isinstance(sheet_data[3]["pages"], int)
    assert isinstance(sheet_data[1]["pages"], str)
    sheet = workbook["Results"]
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:F6"
    assert sheet.column_dimensions["A"].width == 10 and sheet.column_dimensions["B"].width == 32
    assert all(cell.font.bold for cell in sheet[1])


PAIRS = [
    (a, b)
    for a, b in itertools.product(REQUEST_STATUSES, repeat=2)
    if a not in ("aborted_by_user", "not_attempted") or b == "not_attempted"
]


@pytest.mark.parametrize("mode", ["table_per_file", "table_per_page"])
@pytest.mark.parametrize("failed_page", [False, True])
@pytest.mark.parametrize("pair", PAIRS)
def test_mixed_request_outcomes_preserve_answers_and_gaps(tmp_path, mode, failed_page, pair):
    db = tmp_path / "state.db"
    controller = SQLiteDatabaseController(db)
    try:
        pages = [(n, f"{n}.jpg", "ok", "", None) for n in range(1, 5)]
        if failed_page:
            pages.append((5, "", "failure", "render failure", None))
        add_file(controller, 1, "sample.pdf", pages)
        record_configuration(controller, True, output_mode=mode)
        expected = []
        requests = []
        for number, status in enumerate(pair, 1):
            carried = (1, 2) if number == 1 else (3, 4)
            answers = []
            if status == "ok":
                if mode == "table_per_file":
                    answers = [answer_row(values={"answer": f"response-{number}"})]
                    expected.append((f"{carried[0]}-{carried[-1]}", f"response-{number}"))
                else:
                    for page in carried:
                        answers.append(answer_row(page, str(page), values={"answer": f"response-{page}"}))
                        expected.append((str(page), f"response-{page}"))
            else:
                expected.append((f"{carried[0]}-{carried[-1]}", ""))
            requests.append(
                request_outcome(
                    request_number=number,
                    pages=carried,
                    status=status,
                    error="" if status == "ok" else f"reason-{number}-{status}",
                    answer_rows=tuple(answers),
                )
            )
        store_requests(controller, 1, list(reversed(requests)))
        overall = controller.get_file_statuses(1)["overall_result"]
    finally:
        controller.close()
    _, rows = read_csv_rows(export_results_csv(db, tmp_path / "out"))
    assert [(r["pages"], r["llm_answer"]) for r in rows] == expected
    assert all(r["file_path"] == "sample.pdf" and r["file_result"] == overall for r in rows)
    assert all(("failed conversion" in r["notes"]) == failed_page for r in rows)
    for number, status in enumerate(pair, 1):
        if status != "ok":
            assert sum(f"reason-{number}-{status}" in row["notes"] for row in rows) == 1


def test_unverified_duplicate_and_missing_answers_are_not_lost(tmp_path):
    db = tmp_path / "state.db"
    seed_one_file(
        db,
        [(1, "one.jpg", "ok", "", None), (2, "two.jpg", "ok", "", None)],
        requests=[
            request_outcome(
                pages=(1, 2),
                status="invalid_json_answer",
                error="Page claims mismatch",
                answer_rows=(
                    answer_row(1, "1", values={"answer": "first"}),
                    answer_row(1, "1", "duplicate_page_number", {"answer": "second"}),
                    answer_row(None, "99", "hallucinated_page_number", {"answer": "unverified"}),
                    answer_row(2, "", "no_row_returned"),
                ),
            )
        ],
    )
    c = SQLiteDatabaseController(db)
    record_configuration(c, True, output_mode="table_per_page")
    c.close()
    _, rows = read_csv_rows(export_results_csv(db, tmp_path / "out"))
    assert [(r["pages"], r["llm_answer"]) for r in rows] == [
        ("1", "first"),
        ("1", "second"),
        ("", "unverified"),
        ("2", ""),
    ]
    assert "Additional answer" in rows[1]["notes"]
    assert "Unverified page claim: 99" in rows[2]["notes"]
    assert "No answer returned" in rows[3]["notes"]
    assert all("Page claims mismatch" in r["notes"] for r in rows)


@pytest.mark.parametrize(
    "value,expected",
    [
        (None, ""),
        ("", ""),
        (0, "0"),
        (False, "false"),
        (True, "true"),
        (123456789012345678901, "123456789012345678901"),
        ([1, "two"], '[1, "two"]'),
        ({"a": 1}, '{"a": 1}'),
    ],
)
def test_values_keep_their_meaning(tmp_path, value, expected):
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None)],
        requests=[request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": value}),))],
    )
    _, rows = read_csv_rows(export_results_csv(db, tmp_path / "out"))
    assert len(rows) == 1 and rows[0]["llm_answer"] == expected
    assert ("AI returned empty values" in rows[0]["notes"]) == (value is None or value == "")


@pytest.mark.parametrize("range_status", ["ok", "failure", "skipped", "truncated", "partial_skip"])
def test_ai_off_is_one_row_per_file_even_with_page_failures(tmp_path, range_status):
    db = seed_one_file(
        tmp_path / "state.db",
        [(n, "", "failure", "failed", None) for n in range(1, 10)],
        range_status=range_status,
        comment="Recorded range comment",
    )
    headers, rows = read_csv_rows(export_results_csv(db, tmp_path / "out"))
    assert headers == ["file_id", "file_path", "file_result", "notes"] and len(rows) == 1
    assert "9 page(s) failed conversion" in rows[0]["notes"]
    assert "Recorded range comment" in rows[0]["notes"]
    assert ("Fewer pages or frames than requested" in rows[0]["notes"]) == (range_status == "truncated")
    assert ("outside the source" in rows[0]["notes"]) == (range_status == "partial_skip")


def test_empty_run_and_ai_on_without_requests_keep_recorded_headers(tmp_path):
    for ai_on in (False, True):
        db = tmp_path / f"{ai_on}.db"
        c = SQLiteDatabaseController(db)
        record_configuration(c, ai_on)
        c.close()
        headers, rows = read_csv_rows(export_results_csv(db, tmp_path / str(ai_on)))
        assert rows == []
        assert headers == (
            ["file_id", "file_path", "pages", "file_result", "llm_answer", "notes"]
            if ai_on
            else ["file_id", "file_path", "file_result", "notes"]
        )


def test_missing_ai_evidence_and_torn_files_are_explicit(tmp_path):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    add_file(c, 1, "complete-conversion.png", [(1, "one.jpg", "ok", "", None)])
    c.handle_file_started(2, "torn.pdf", ".pdf", "Stub")
    c.handle_frame_saved(2, PageResult(1, "two.jpg", "ok", ""))
    record_configuration(c, True)
    c.close()
    _, rows = read_csv_rows(export_results_csv(db, tmp_path / "out"))
    assert len(rows) == 2
    assert rows[0]["file_result"] == "ok"
    assert all("AI completion cannot be confirmed" in r["notes"] for r in rows)
    assert "File completion was not recorded" in rows[1]["notes"]


def test_attempts_are_preserved_and_distinguished(tmp_path):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    for ident in (1, 2):
        add_file(c, ident, "same.png", [(1, "one.jpg", "ok", "", None)])
        store_requests(
            c, ident, [request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": str(ident)}),))]
        )
    record_configuration(c, True)
    c.close()
    _, rows = read_csv_rows(export_results_csv(db, tmp_path / "out"))
    assert [(r["llm_answer"], r["notes"]) for r in rows] == [("1", "File attempt 1"), ("2", "File attempt 2")]


@pytest.mark.parametrize("stored", ["2, 1", "1, 1", "nonsense", "003", "\u00b3", ""])
def test_stored_page_text_is_exported_without_reinterpretation(tmp_path, stored):
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None)],
        requests=[request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": "a"}),))],
    )
    with sqlite3.connect(db) as con:
        con.execute("UPDATE llm_requests SET pages=?", (stored,))
    _, rows = read_csv_rows(export_results_csv(db, tmp_path / "out"))
    assert rows[0]["pages"] == stored


@pytest.mark.parametrize("raw", ["", "{", "[]", "{}", '{"answer":1,"extra":2}', '{"answer":NaN}', '{"answer":1e999}'])
def test_inconsistent_answers_refuse_normal_reports_and_preserve_raw_recovery(tmp_path, raw):
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None)],
        requests=[request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": "original"}),))],
    )
    with sqlite3.connect(db) as con:
        con.execute("UPDATE llm_answers SET values_json=?", (raw,))
    exporter = make_exporter(db)
    with pytest.raises(data_exporter.ExportError) as refused:
        exporter.export_all_formats(tmp_path / "out")
    outcome = refused.value.outcome
    assert outcome.recovery
    assert {r["report"] for r in outcome.failed} >= {"Results", "LLM Answers", "Workbook"}
    assert not list((tmp_path / "out").glob("results_*.csv"))
    book = openpyxl.load_workbook(next((tmp_path / "out").glob("recovery_*.xlsx")))
    assert "Results" not in book.sheetnames
    rows = sheet_rows(book, "Raw llm_answers")
    assert cell_text(rows[1][rows[0].index("values_json")]) == raw


@pytest.mark.parametrize(
    "text",
    [
        "=1+1",
        "+100",
        "-2",
        "@name",
        "00001",
        "2026-09-14",
        "00:01:33.37",
        "1234567890123456789012345",
        "#N/A",
        'commas, quotes " and\nnewlines',
        "_x000D_",
    ],
)
def test_every_model_string_is_explicit_text_in_saved_workbook(tmp_path, text):
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None)],
        path="=cmd().png",
        requests=[
            request_outcome(pages=(1,), raw_answer=text, error=text, answer_rows=(answer_row(values={"answer": text}),))
        ],
    )
    workbook = export_workbook(db, tmp_path / "out")
    for sheet in workbook.worksheets:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    assert cell.data_type in ("s", "inlineStr")
    _, rows = results_sheet(workbook)
    assert rows[0]["llm_answer"] == text and rows[0]["file_path"] == "=cmd().png"


def test_recovered_retry_and_encoding_skip_remain_visible(tmp_path):
    note = "HTTP 503 recovered on retry; image 2 could not be encoded"
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None), (2, "two.jpg", "ok", "", None)],
        requests=[request_outcome(pages=(1,), error=note, answer_rows=(answer_row(values={"answer": "a"}),))],
    )
    _, rows = read_csv_rows(export_results_csv(db, tmp_path / "out"))
    assert len(rows) == 1 and rows[0]["pages"] == "1" and rows[0]["notes"] == note


def test_missing_openpyxl_is_failure_with_csv_inventory(tmp_path, monkeypatch):
    import builtins

    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)])
    original = builtins.__import__

    def missing(name, *args, **kwargs):
        if name == "openpyxl":
            raise ImportError("missing Excel library")
        return original(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", missing)
    with pytest.raises(data_exporter.ExportError) as failure:
        make_exporter(db).export_all_formats(tmp_path / "out")
    assert len(failure.value.outcome.saved) == 3
    assert failure.value.outcome.failed[0]["category"] == "installation"
    assert not list((tmp_path / "out").glob("*.xlsx"))


def test_failed_workbook_never_leaves_a_partial_final_or_temp_file(tmp_path, monkeypatch):
    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)])

    from openpyxl.writer.excel import ExcelWriter

    def fail(self):
        self._archive.writestr("partial.txt", b"partial")
        raise RuntimeError("save interrupted")

    monkeypatch.setattr(ExcelWriter, "write_data", fail)
    with pytest.raises(data_exporter.ExportError) as failure:
        make_exporter(db).export_all_formats(tmp_path / "out")
    assert len(failure.value.outcome.saved) == 3
    assert all(p.suffix == ".csv" for p in (tmp_path / "out").iterdir())


def test_csv_failure_does_not_prevent_workbook_and_summary_does_not_promise_missing_csv(tmp_path, monkeypatch):
    db = seed_one_file(tmp_path / "state.db", [(n, "one.jpg", "ok", "", None) for n in range(1, 8)])
    original = data_exporter.SQLiteDataExporter._write_csv

    def fail(self, path, headers, rows):
        if "page_id" in headers:
            raise PermissionError("page log CSV locked")
        return original(self, path, headers, rows)

    monkeypatch.setattr(data_exporter.SQLiteDataExporter, "_write_csv", fail)
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_ROWS", 3)
    with pytest.raises(data_exporter.ExportError) as failure:
        make_exporter(db).export_all_formats(tmp_path / "out")
    outcome = failure.value.outcome
    assert any(r["format"] == "xlsx" for r in outcome.saved)
    book = openpyxl.load_workbook(next((tmp_path / "out").glob("*.xlsx")))
    text = str(sheet_rows(book, "Summary"))
    assert "CSV unavailable; retained database" in text and "page log CSV locked" in text


def test_same_timestamp_preserves_previous_exports(tmp_path):
    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)])
    exporter = make_exporter(db)
    first = exporter.export_csv(tmp_path / "out", "t")
    first_bytes = {r["path"]: Path(r["path"]).read_bytes() for r in first.saved}
    second = exporter.export_csv(tmp_path / "out", "t")
    assert not set(first_bytes) & {r["path"] for r in second.saved}
    assert all(Path(path).read_bytes() == content for path, content in first_bytes.items())


@pytest.mark.parametrize(
    "stored,expected",
    [
        ("_x005F_x000D_", "_x000D_"),
        ("_x000D_\n", "\r\n"),
        ("_x005F_x0041_x005F_x0042_", "_x0041_x0042_"),
        ("_x005F_x005F_", "_x005F_"),
    ],
)
def test_office_string_reference_is_not_recursive(stored, expected):
    assert excel_value(stored) == expected


@pytest.mark.parametrize("text", ["_x000D_", "_x0041_", "_x005F_", "_x005F_x000D_", "_x0041_x0042_", "_X000D_"])
def test_literal_excel_escape_spellings_survive_every_text_surface(tmp_path, text):
    key = "_x0041_"
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    add_file(c, 1, "literal_x000D_.png", [(1, "one_x0041_.jpg", "ok", "", None)])
    record_configuration(c, True, declared=(key,), output_mode="table_per_page")
    raw = json.dumps([{"page": "_x0042_", key: text}])
    store_requests(
        c,
        1,
        [
            request_outcome(
                pages=(1,),
                status="invalid_json_answer",
                raw_answer=raw,
                error="failure_x0041_",
                answer_rows=(answer_row(None, "_x0042_", "hallucinated_page_number", {key: text}),),
            )
        ],
    )
    c.close()
    out = tmp_path / "out"
    make_exporter(db).export_all_formats(out)
    book = openpyxl.load_workbook(next(out.glob("*.xlsx")))
    headers, rows = results_sheet(book)
    assert "llm_" + key in headers and rows[0]["llm_" + key] == text
    assert rows[0]["file_path"] == "literal_x000D_.png"
    assert "failure_x0041_" in rows[0]["notes"] and "_x0042_" in rows[0]["notes"]
    requests = sheet_rows(book, "LLM Requests")
    assert requests[1][requests[0].index("raw_llm_answer")] == raw
    answers = sheet_rows(book, "LLM Answers")
    assert answers[1][answers[0].index("raw_model_page_number")] == "_x0042_"
    _, csv_rows = read_csv_rows(next(out.glob("results_*.csv")))
    assert csv_rows[0]["llm_" + key] == text


def test_escape_encoding_expansion_does_not_trigger_library_truncation(tmp_path):
    import zipfile
    import xml.etree.ElementTree as ET

    text = "_x000D_" * 4000
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None)],
        requests=[request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": text}),))],
    )
    make_exporter(db).export_all_formats(tmp_path / "out")
    with zipfile.ZipFile(next((tmp_path / "out").glob("*.xlsx"))) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))  # noqa: S314
        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        node = root.find(".//x:c[@r='E2']/x:is/x:t", ns)
        assert node is not None and node.text is not None
        assert len(node.text) > 32767
        assert excel_value(node.text) == text
    _, rows = read_csv_rows(next((tmp_path / "out").glob("results_*.csv")))
    assert rows[0]["llm_answer"] == text


@pytest.mark.parametrize("line_break", ["\n", "\r\n"])
def test_real_line_breaks_are_preserved_and_have_visible_row_height(tmp_path, line_break):
    text = "first" + line_break + "second"
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None)],
        requests=[request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": text}),))],
    )
    book = export_workbook(db, tmp_path / "out")
    for title in ("Results", "LLM Answers"):
        rows = sheet_rows(book, title)
        column = rows[0].index("llm_answer") + 1
        assert rows[1][column - 1] == text
        assert book[title].cell(2, column).alignment.wrap_text
        assert book[title].row_dimensions[2].height >= 34


def test_large_file_ids_remain_exact_in_all_linked_reports(tmp_path):
    ident = 9007199254740993
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    add_file(c, ident, "large-id.png", [(1, "one.jpg", "ok", "", None)])
    record_configuration(c, True)
    store_requests(c, ident, [request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": "a"}),))])
    c.close()
    book = export_workbook(db, tmp_path / "out")
    for title in ("Results", "Master Registry", "Page Log", "LLM Requests", "LLM Answers"):
        rows = sheet_rows(book, title)
        column = rows[0].index("file_id") + 1
        assert rows[1][column - 1] == str(ident)
        assert book[title].cell(2, column).data_type == "s"


def test_specific_request_cause_precedes_file_wide_warning(tmp_path):
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None), (2, "two.jpg", "ok", "", None)],
        requests=[
            request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": "a"}),)),
            request_outcome(request_number=2, pages=(2,), status="network_failure", error="Provider connection failed"),
        ],
    )
    book = export_workbook(db, tmp_path / "out")
    _, rows = results_sheet(book)
    assert [row["file_id"] for row in rows] == [1, 1]
    assert rows[1]["notes"].startswith("Request 2: Provider connection failed")
    assert book["Results"].row_dimensions[3].height > 15


@pytest.mark.parametrize("field", ["request_id", "file_id", "page_id"])
def test_broken_answer_links_are_refused_instead_of_losing_a_reader_row(tmp_path, field):
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None)],
        requests=[request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": "preserve me"}),))],
    )
    with sqlite3.connect(db) as con:
        statement = {
            "request_id": "UPDATE llm_answers SET request_id=999",
            "file_id": "UPDATE llm_answers SET file_id=999",
            "page_id": "UPDATE llm_answers SET page_id=999",
        }[field]
        con.execute(statement)
    with pytest.raises(data_exporter.ExportError) as refused:
        make_exporter(db).export_all_formats(tmp_path / "out")
    assert refused.value.outcome.recovery
    assert "inconsistent record references" in str(refused.value)
    book = openpyxl.load_workbook(next((tmp_path / "out").glob("recovery_*.xlsx")))
    assert "preserve me" in str(sheet_rows(book, "Raw llm_answers"))


def test_the_joined_sheet_and_its_toggle_are_gone():
    import inspect

    from config_validator import Settings, SettingsAIDormant

    for model in (Settings, SettingsAIDormant):
        assert "EXPORT_JOINED_SHEET" not in model.model_fields
        assert "EXPORT_JOINED_SHEET" not in model().model_dump()
    assert "joined_sheet" not in inspect.signature(data_exporter.SQLiteDataExporter).parameters
    example = json.loads((SRC.parent / "settings.example.json").read_text(encoding="utf-8"))
    assert "EXPORT_JOINED_SHEET" not in example
    for template in (SRC / "templates").rglob("*.html"):
        assert "EXPORT_JOINED_SHEET" not in template.read_text(encoding="utf-8")
    for locale in sorted((SRC / "locales").glob("*.json")):
        keys = json.loads(locale.read_text(encoding="utf-8"))
        assert "lbl_joined_sheet" not in keys and "hint_joined_sheet" not in keys


@pytest.mark.parametrize("code", [0x0, 0x7, 0xB, 0x1F, 0xFFFE, 0xFFFF])
def test_all_xml_forbidden_characters_are_replaced_and_csv_retains_them(tmp_path, code):
    text = "before" + chr(code) + "after"
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None)],
        requests=[
            request_outcome(pages=(1,), raw_answer=text, error=text, answer_rows=(answer_row(values={"answer": text}),))
        ],
    )
    outcome = make_exporter(db).export_all_formats(tmp_path / "out")
    book = openpyxl.load_workbook(next((tmp_path / "out").glob("*.xlsx")))
    _, rows = results_sheet(book)
    assert rows[0]["llm_answer"] == "before\ufffdafter"
    _, csv_rows = read_csv_rows(next((tmp_path / "out").glob("results_*.csv")))
    assert csv_rows[0]["llm_answer"] == text
    assert any("characters replaced" in notice for notice in outcome.notices)


@pytest.mark.parametrize("text", ["\t\n\r", "\r\n", "\r\r\n", "\n", "\u0020\ud7ff\ue000\ufffd", "\U00010000\U0010ffff"])
def test_valid_xml_boundary_characters_survive(tmp_path, text):
    db = seed_one_file(
        tmp_path / "state.db",
        [(1, "one.jpg", "ok", "", None)],
        requests=[request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": text}),))],
    )
    outcome = make_exporter(db).export_all_formats(tmp_path / "out")
    book = openpyxl.load_workbook(next((tmp_path / "out").glob("*.xlsx")))
    _, rows = results_sheet(book)
    assert rows[0]["llm_answer"] == text
    assert not any("characters replaced" in notice for notice in outcome.notices)


@pytest.mark.parametrize("kind", ["disk_full", "access_denied", "locked"])
def test_promotion_failures_have_bounded_appropriate_retries(tmp_path, monkeypatch, kind):
    import errno

    calls, waits = [], []
    exporter = make_exporter(tmp_path / "unused.db")

    def reject(*args):
        calls.append(args)
        if kind == "disk_full":
            raise OSError(errno.ENOSPC, "disk full")
        if kind == "access_denied":
            raise PermissionError("denied")
        raise OSError(errno.EBUSY, "sharing violation")

    monkeypatch.setattr(data_exporter.os, "rename", reject)
    monkeypatch.setattr(data_exporter.time, "sleep", waits.append)
    with pytest.raises(OSError):

        def write(path):
            path.write_bytes(b"complete")

        exporter._atomic_write(tmp_path / "report.csv", write)
    assert len(calls) == (4 if kind == "locked" else 1)
    assert sum(waits) == (0.75 if kind == "locked" else 0)
    assert list(tmp_path.iterdir()) == []


def test_transient_promotion_lock_preserves_existing_file_and_serializes_once(tmp_path, monkeypatch):
    import errno

    target = tmp_path / "report.csv"
    target.write_bytes(b"previous")
    calls, writes = [], []
    original = data_exporter.os.rename

    def rename(source, destination):
        calls.append(destination)
        if len(calls) == 1:
            raise OSError(errno.EBUSY, "sharing violation")
        return original(source, destination)

    def write(path):
        writes.append(path)
        path.write_bytes(b"new")

    monkeypatch.setattr(data_exporter.os, "rename", rename)
    monkeypatch.setattr(data_exporter.time, "sleep", lambda delay: None)
    final = make_exporter(tmp_path / "unused.db")._atomic_write(target, write)
    assert target.read_bytes() == b"previous" and final.read_bytes() == b"new"
    assert len(writes) == 1 and len(calls) == 2


def test_cleanup_failure_does_not_mask_write_failure(tmp_path, monkeypatch, caplog):
    original = Path.unlink

    def deny_temp(path, *args, **kwargs):
        if path.name.startswith(".report-"):
            raise PermissionError("cleanup denied")
        return original(path, *args, **kwargs)

    def fail(path):
        raise RuntimeError("original save failure")

    with monkeypatch.context() as patch:
        patch.setattr(Path, "unlink", deny_temp)
        with pytest.raises(RuntimeError, match="original save failure"):
            make_exporter(tmp_path / "unused.db")._atomic_write(tmp_path / "report.csv", fail)
    assert "Could not remove temporary report" in caplog.text
    for path in tmp_path.iterdir():
        path.unlink()


def test_workbook_save_does_not_need_to_keep_all_rows_in_memory(tmp_path):
    from test_workbook_limits import assert_resource_growth, assert_resource_output, resource_database

    assert_resource_growth(tmp_path, "files")
    db = tmp_path / "original-size.db"
    spec = resource_database(db, "files", 2000)
    outcome = make_exporter(db).export_all_formats(tmp_path / "original-size")
    assert_resource_output(outcome, spec)


def test_a_fatal_credential_rejection_reaches_results_through_the_real_app(tmp_path, monkeypatch):
    from test_llm_run_cruelty import make_png
    from test_workbook_limits import section, summary_table

    server = make_server("openai").start()
    try:
        server.queue(status=401, json={"error": {"message": "Incorrect API key provided"}})
        events, db_path = run_through_the_real_app(
            tmp_path, monkeypatch, [(n, make_png) for n in ("a.png", "b.png", "c.png")], server
        )
        assert len(server.requests) == 1
    finally:
        server.stop()
    assert events[-1]["type"] != "done", events[-3:]
    _, rows = read_csv_rows(export_results_csv(db_path, tmp_path / "reports"))
    assert len(rows) == 1 and rows[0]["file_result"] == "partial_fail"
    assert "FATAL AUTHENTICATION ERROR" in rows[0]["notes"]
    table = summary_table(export_workbook(db_path, tmp_path / "xlsx"))
    assert sum(section(table, "file records").values()) == 1


def test_stop_mid_run_reaches_results_as_not_attempted_rows(tmp_path, monkeypatch):
    from test_llm_run_cruelty import make_tiff
    from fake_llm.generic import openai_reply

    flag = threading.Event()
    server = make_server("openai").start()
    try:
        server.queue(on_request=lambda rec: flag.set(), json=openai_reply('{"answer": "a"}'))
        events, db_path = run_through_the_real_app(
            tmp_path,
            monkeypatch,
            [("doc.tif", lambda p: make_tiff(p, page_count=3))],
            server,
            abort_flag=flag,
            MAX_JPEGS_PER_INFERENCE=1,
        )
    finally:
        server.stop()
    assert events[-1]["type"] != "failed", events[-3:]
    _, rows = read_csv_rows(export_results_csv(db_path, tmp_path / "reports"))
    assert rows and all(r["file_result"] == "partial_fail" for r in rows)
    assert any(r["pages"] == "3" and "Request not sent" in r["notes"] for r in rows)
    assert all(r["pages"] for r in rows)


@pytest.mark.skipif(sys.platform != "win32", reason="NTFS file names are UTF-16 and may carry a lone surrogate")
def test_a_file_named_with_a_lone_surrogate_is_reported_not_lost(tmp_path, monkeypatch, caplog):
    from test_llm_run_cruelty import exported, make_png, run_core, settings_for
    import central_logger

    monkeypatch.setattr(central_logger, "setup_logging", lambda *a, **k: None)
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    make_png(input_dir / "photo_\ud800.png")
    make_png(input_dir / "plain.png")
    settings = settings_for(
        (input_dir, tmp_path / "output"), "openai", "http://127.0.0.1:9", ENABLE_LLM_INFERENCE=False
    )
    with caplog.at_level(logging.CRITICAL + 10):
        events, db_path = run_core(settings)
    assert events[-1]["type"] == "done"
    registered = [row[0] for row in sql(db_path, "SELECT file_path FROM file_registry ORDER BY file_id")]
    assert len(registered) == 2, f"the file must be registered, not lost: {registered}"
    assert any("photo_" in name and "�" in name for name in registered)
    assert exported(settings.TECH_FOLDER_PATH, "results") is not None, "the run's own export landed"
    _, rows = read_csv_rows(export_results_csv(db_path, tmp_path / "reports"))
    assert sorted(r["file_path"] for r in rows) == sorted(registered)


@pytest.mark.parametrize("mode", ["table_per_file", "table_per_page"])
def test_adversarial_row_conservation_uses_independent_ordered_facts(tmp_path, mode):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    try:
        for fid, path in [(40, "same.pdf"), (30, "empty.pdf"), (20, "other/same.pdf"), (10, "same.pdf")]:
            add_file(c, fid, path, [(1, f"{fid}.jpg", "ok", "", None)])
        store_requests(c, 10, [
            request_outcome(request_number=2, pages=(1,), status="network_failure", error="offline")])
        store_requests(c, 40, [request_outcome(pages=(1,), answer_rows=(answer_row(1, values={"answer": False}),))])
        store_requests(c, 20, [request_outcome(pages=(1,), status="not_attempted")])
        store_requests(c, 10, [request_outcome(pages=(1,), answer_rows=(
            answer_row(1, values={"answer": 0}), answer_row(1, "1", "duplicate_page_number", {"answer": 0}),
        ))])
        record_configuration(c, True, output_mode=mode)
    finally:
        c.close()
    before = db.read_bytes()
    outcome = make_exporter(db).export_all_formats(tmp_path / "out")
    csv_path = next(Path(r["path"]) for r in outcome.saved if r["report"] == "Results")
    headers, rows = read_csv_rows(csv_path)
    assert headers == ["file_id", "file_path", "pages", "file_result", "llm_answer", "notes"]
    assert [(r["file_id"], r["file_path"], r["pages"], r["file_result"], r["llm_answer"]) for r in rows] == [
        ("10", "same.pdf", "1", "partial_fail", "0"),
        ("10", "same.pdf", "1", "partial_fail", "0"),
        ("10", "same.pdf", "1", "partial_fail", ""),
        ("20", "other/same.pdf", "1", "partial_fail", ""),
        ("30", "empty.pdf", "", "ok", ""),
        ("40", "same.pdf", "1", "ok", "false"),
    ]
    assert all("File attempt 1" in r["notes"] for r in rows[:3])
    assert "File attempt 2" in rows[-1]["notes"]
    assert "File attempt" not in rows[3]["notes"]
    assert "Additional answer" in rows[1]["notes"]
    assert "offline" in rows[2]["notes"]
    assert "not sent" in rows[3]["notes"]
    assert "completion cannot be confirmed" in rows[4]["notes"]
    book = openpyxl.load_workbook(next(r["path"] for r in outcome.saved if r["format"] == "xlsx"))
    try:
        _, actual = results_sheet(book)
        assert [{k: cell_text(v) for k, v in row.items()} for row in actual] == rows
        assert len(sheet_rows(book, "LLM Requests")) == 5
        assert len(sheet_rows(book, "LLM Answers")) == 4
    finally:
        book.close()
    assert db.read_bytes() == before


@pytest.mark.parametrize("broken", ["orphan_page", "orphan_request", "cross_file_page", "cross_file_request"])
def test_adversarial_broken_record_families_refuse_and_preserve_every_raw_row(tmp_path, broken):
    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)],
                       requests=[request_outcome(pages=(1,), answer_rows=(answer_row(1, values={"answer": "keep"}),))])
    c = SQLiteDatabaseController(db)
    try:
        add_file(c, 2, "second.pdf", [(1, "two.jpg", "ok", "", None)])
        store_requests(c, 2, [request_outcome(pages=(1,), answer_rows=(answer_row(1, values={"answer": "second"}),))])
    finally:
        c.close()
    queries = {
        "file_registry": "SELECT * FROM file_registry ORDER BY 1",
        "page_log": "SELECT * FROM page_log ORDER BY 1",
        "llm_requests": "SELECT * FROM llm_requests ORDER BY 1",
        "llm_answers": "SELECT * FROM llm_answers ORDER BY 1",
        "run_configuration": "SELECT * FROM run_configuration ORDER BY 1",
    }
    with sqlite3.connect(db) as con:
        con.execute({
            "orphan_page": "UPDATE page_log SET file_id=999 WHERE page_id=2",
            "orphan_request": "UPDATE llm_requests SET file_id=999 WHERE request_id=2",
            "cross_file_page": "UPDATE llm_answers SET page_id=2 WHERE llm_answer_id=1",
            "cross_file_request": "UPDATE llm_answers SET request_id=2 WHERE llm_answer_id=1",
        }[broken])
        expected = {table: con.execute(query).fetchall() for table, query in queries.items()}
    with pytest.raises(data_exporter.ExportError) as refused:
        make_exporter(db).export_all_formats(tmp_path / "out")
    outcome = refused.value.outcome
    assert outcome.recovery
    assert not any(r["report"] in ("Results", "Workbook") for r in outcome.saved)
    book = openpyxl.load_workbook(next(r["path"] for r in outcome.saved if r["report"] == "Recovery workbook"))
    try:
        for table, records in expected.items():
            assert [[cell_text(v) for v in row] for row in sheet_rows(book, "Raw " + table)[1:]] == [
                [cell_text(v) for v in row] for row in records
            ]
    finally:
        book.close()


def test_adversarial_unreferenced_orphan_page_cannot_be_a_complete_report(tmp_path):
    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)])
    with sqlite3.connect(db) as con:
        con.execute("UPDATE page_log SET file_id=999")
    with pytest.raises(data_exporter.ExportError) as refused:
        make_exporter(db).export_all_formats(tmp_path / "out")
    assert refused.value.outcome.recovery
    assert not any(r["report"] == "Workbook" for r in refused.value.outcome.saved)


@pytest.mark.parametrize("view", ["overall_result", "file_to_jpegs_status", "file_to_llm_status"])
def test_adversarial_broken_summary_view_still_recovers_readable_raw_tables(tmp_path, view):
    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)])
    with sqlite3.connect(db) as con:
        con.execute({
            "overall_result": "DROP VIEW overall_result",
            "file_to_jpegs_status": "DROP VIEW file_to_jpegs_status",
            "file_to_llm_status": "DROP VIEW file_to_llm_status",
        }[view])
    with pytest.raises(data_exporter.ExportError) as refused:
        make_exporter(db).export_all_formats(tmp_path / "out")
    outcome = refused.value.outcome
    assert any(r["report"] == "Page Log" for r in outcome.saved)
    assert outcome.recovery, "Readable source tables require raw recovery even when Summary cannot be counted"
    book = openpyxl.load_workbook(next(r["path"] for r in outcome.saved if r["report"] == "Recovery workbook"))
    try:
        assert len(sheet_rows(book, "Raw file_registry")) == 2
        assert len(sheet_rows(book, "Raw page_log")) == 2
        assert view in str(sheet_rows(book, "Summary"))
    finally:
        book.close()


@pytest.mark.parametrize("stage", ["create", "open", "write", "close", "serialize"])
@pytest.mark.parametrize("kind", ["success", "transient_lock", "persistent_lock", "disk_full", "access_denied"])
def test_adversarial_save_stage_matrix(tmp_path, monkeypatch, stage, kind):
    import builtins
    import errno

    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)])
    out = tmp_path / "out"
    out.mkdir()
    previous = out / "earlier.csv"
    previous.write_bytes(b"earlier completed export")
    calls, waits = [], []
    original_open = builtins.open
    original_create = data_exporter.tempfile.mkstemp
    from openpyxl.writer.excel import ExcelWriter
    original_save = ExcelWriter.write_data

    def fault():
        calls.append(stage)
        if kind == "success" or (kind == "transient_lock" and len(calls) > 1):
            return
        if kind in ("transient_lock", "persistent_lock"):
            raise OSError(errno.EBUSY, "injected sharing violation")
        if kind == "disk_full":
            raise OSError(errno.ENOSPC, "injected disk full")
        raise PermissionError(errno.EACCES, "injected access denial")

    class Stream:
        def __init__(self, wrapped):
            self.wrapped = wrapped
            self.writes = 0

        def __enter__(self):
            self.wrapped.__enter__()
            return self

        def write(self, value):
            self.writes += 1
            if stage == "write" and self.writes == 2:
                fault()
            return self.wrapped.write(value)

        def __exit__(self, *args):
            result = self.wrapped.__exit__(*args)
            if stage == "close":
                fault()
            return result

    def open_report(*args, **kwargs):
        if stage == "open":
            fault()
        return Stream(original_open(*args, **kwargs))

    def create(*args, **kwargs):
        fault()
        return original_create(*args, **kwargs)

    def save(writer):
        fault()
        return original_save(writer)

    exporter = make_exporter(db)
    original_reports = exporter._reports
    monkeypatch.setattr(exporter, "_reports", lambda con, config: original_reports(con, config)[:1])
    monkeypatch.setattr(data_exporter.time, "sleep", waits.append)
    if stage == "create":
        monkeypatch.setattr(data_exporter.tempfile, "mkstemp", create)
    elif stage == "serialize":
        monkeypatch.setattr(ExcelWriter, "write_data", save)
    else:
        monkeypatch.setattr(data_exporter, "open", open_report, raising=False)
    try:
        outcome = (exporter.export_all_formats(out) if stage == "serialize" else exporter.export_csv(out, "fault"))
    except data_exporter.ExportError as exc:
        outcome = exc.outcome
    expected_calls = {"success": 1, "transient_lock": 2, "persistent_lock": 4,
                      "disk_full": 1, "access_denied": 1}[kind]
    assert calls, "The intended fault boundary was never exercised"
    assert previous.read_bytes() == b"earlier completed export"
    assert not list(out.glob(".report-*.tmp"))
    assert len(calls) == expected_calls, (stage, kind, outcome.as_dict())
    assert len(waits) == (expected_calls - 1 if "lock" in kind else 0)
    assert sum(waits) <= 1.0
    if kind in ("success", "transient_lock"):
        assert not outcome.failed, outcome.as_dict()
        result = next(Path(r["path"]) for r in outcome.saved if r["report"] == "Results")
        assert len(read_csv_rows(result)[1]) == 1
        if stage == "serialize":
            book = openpyxl.load_workbook(next(r["path"] for r in outcome.saved if r["format"] == "xlsx"))
            try:
                assert len(results_sheet(book)[1]) == 1
            finally:
                book.close()
    else:
        expected_category = {
            "persistent_lock": "locked", "disk_full": "disk_full", "access_denied": "access_denied",
        }[kind]
        assert [r["category"] for r in outcome.failed] == [expected_category]
        assert not any(r["format"] == "xlsx" for r in outcome.saved)
    assert {Path(r["path"]).name for r in outcome.saved} == {p.name for p in out.iterdir()} - {previous.name}


def test_adversarial_export_snapshot_survives_a_commit_between_formats(tmp_path, monkeypatch):
    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)],
                       requests=[request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": "before"}),))])
    with sqlite3.connect(db) as con:
        assert con.execute("PRAGMA journal_mode=WAL").fetchone()[0] == "wal"
    exporter = make_exporter(db)
    original = exporter._write_csv
    mutations = []

    def write_then_change(path, headers, rows):
        original(path, headers, rows)
        if not mutations:
            with sqlite3.connect(db, timeout=2) as con:
                con.execute("UPDATE llm_answers SET values_json=?", ('{"answer":"after"}',))
            mutations.append(True)

    monkeypatch.setattr(exporter, "_write_csv", write_then_change)
    outcome = exporter.export_all_formats(tmp_path / "out")
    assert mutations == [True]
    for report in ("Results", "LLM Answers"):
        path = next(r["path"] for r in outcome.saved if r["report"] == report)
        assert read_csv_rows(path)[1][0]["llm_answer"] == "before"
    book = openpyxl.load_workbook(next(r["path"] for r in outcome.saved if r["format"] == "xlsx"))
    try:
        assert results_sheet(book)[1][0]["llm_answer"] == "before"
        assert sheet_rows(book, "LLM Answers")[1][-1] == "before"
    finally:
        book.close()
    assert sql(db, "SELECT values_json FROM llm_answers") == [('{"answer":"after"}',)]



@pytest.mark.parametrize("phase", ["construct", "serialize", "archive_close"])
def test_workbook_retry_rebuilds_consumed_sheets_and_closes_every_archive(tmp_path, monkeypatch, phase):
    import errno
    from openpyxl.writer.excel import ExcelWriter

    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)])
    workbooks, archives, waits = [], [], []
    faulted = False
    original_book = openpyxl.Workbook
    original_write = ExcelWriter.write_data
    original_close = data_exporter.ZipFile.close

    def fail_once(where):
        nonlocal faulted
        if where == phase and not faulted:
            faulted = True
            raise OSError(errno.EBUSY, "temporary workbook lock")

    def construct(*args, **kwargs):
        fail_once("construct")
        book = original_book(*args, **kwargs)
        workbooks.append(book)
        return book

    def write(writer):
        archives.append(writer._archive)
        original_write(writer)
        fail_once("serialize")

    def close(archive):
        was_open = archive.fp is not None
        original_close(archive)
        if was_open:
            fail_once("archive_close")

    with monkeypatch.context() as patch:
        patch.setattr(openpyxl, "Workbook", construct)
        patch.setattr(ExcelWriter, "write_data", write)
        patch.setattr(data_exporter.ZipFile, "close", close)
        patch.setattr(data_exporter.time, "sleep", waits.append)
        outcome = make_exporter(db).export_all_formats(tmp_path / "out")
    assert faulted and waits == [0.25]
    assert len(workbooks) == (1 if phase == "construct" else 2)
    assert len({id(book) for book in workbooks}) == len(workbooks)
    assert all(archive.fp is None for archive in archives)
    assert not list((tmp_path / "out").glob(".report-*.tmp"))
    book = openpyxl.load_workbook(next(r["path"] for r in outcome.saved if r["format"] == "xlsx"))
    try:
        assert len(results_sheet(book)[1]) == 1
        assert len(sheet_rows(book, "Page Log")) == 2
    finally:
        book.close()


@pytest.mark.parametrize("locked", [False, True])
def test_reservation_close_failure_releases_descriptor_and_cleans_partial_file(tmp_path, monkeypatch, locked):
    import errno

    handles, waits = [], []
    failed = False
    fdopen = data_exporter.os.fdopen

    class Reservation:
        def __init__(self, handle):
            self.handle = handle

        @property
        def closed(self):
            return self.handle.closed

        def close(self):
            nonlocal failed
            if not failed:
                failed = True
                raise OSError(errno.EBUSY if locked else errno.EACCES, "reservation close failed")
            self.handle.close()

    def reserve(*args, **kwargs):
        handle = fdopen(*args, **kwargs)
        handles.append(handle)
        return Reservation(handle)

    monkeypatch.setattr(data_exporter.os, "fdopen", reserve)
    monkeypatch.setattr(data_exporter.time, "sleep", waits.append)
    def write(path):
        path.write_bytes(b"complete")

    exporter = make_exporter(tmp_path / "unused.db")
    target = tmp_path / "report.csv"
    if locked:
        assert exporter._atomic_write(target, write) == target
        assert target.read_bytes() == b"complete"
        assert waits == [0.25] and len(handles) == 2
    else:
        with pytest.raises(PermissionError, match="reservation close failed"):
            exporter._atomic_write(target, write)
        assert not target.exists() and waits == [] and len(handles) == 1
    assert all(handle.closed for handle in handles)
    assert not list(tmp_path.glob(".report-*.tmp"))


def test_reservation_wrapper_failure_closes_the_unclaimed_descriptor(tmp_path, monkeypatch):
    import errno

    fdopen, close = data_exporter.os.fdopen, data_exporter.os.close
    unclaimed, closed = [], []

    def wrap(fd, *args, **kwargs):
        if not unclaimed:
            unclaimed.append(fd)
            raise OSError(errno.EBUSY, "wrapper unavailable")
        return fdopen(fd, *args, **kwargs)

    def release(fd):
        closed.append(fd)
        return close(fd)

    monkeypatch.setattr(data_exporter.os, "fdopen", wrap)
    monkeypatch.setattr(data_exporter.os, "close", release)
    monkeypatch.setattr(data_exporter.time, "sleep", lambda seconds: None)
    def write(path):
        path.write_bytes(b"complete")

    path = make_exporter(tmp_path / "unused.db")._atomic_write(tmp_path / "report.csv", write)
    assert unclaimed and closed == unclaimed
    assert path.read_bytes() == b"complete" and not list(tmp_path.glob(".report-*.tmp"))


def test_writing_and_promotion_share_one_retry_budget(tmp_path, monkeypatch):
    import errno

    writes, renames, waits = [], [], []

    def write(path):
        writes.append(path)
        path.write_bytes(b"partial" if len(writes) == 1 else b"complete")
        if len(writes) == 1:
            raise OSError(errno.EBUSY, "write locked")

    def promote(source, destination):
        renames.append(source)
        assert Path(source).read_bytes() == b"complete"
        raise OSError(errno.EBUSY, "rename locked")

    monkeypatch.setattr(data_exporter.os, "rename", promote)
    monkeypatch.setattr(data_exporter.time, "sleep", waits.append)
    with pytest.raises(OSError, match="rename locked"):
        make_exporter(tmp_path / "unused.db")._atomic_write(tmp_path / "report.csv", write)
    assert len(writes) == 2 and len(renames) == 3 and waits == [0.25] * 3
    assert len(set(renames)) == 1 and list(tmp_path.iterdir()) == []


def test_archive_cleanup_failure_preserves_the_original_save_error(tmp_path, monkeypatch, caplog):
    import errno
    from openpyxl.writer.excel import ExcelWriter

    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)])
    original_close = data_exporter.ZipFile.close

    def fail(writer):
        writer._archive.writestr("partial.txt", b"partial")
        raise OSError(errno.ENOSPC, "original disk full")

    def close(archive):
        was_open = archive.fp is not None
        original_close(archive)
        if was_open:
            raise PermissionError("secondary close failure")

    monkeypatch.setattr(ExcelWriter, "write_data", fail)
    monkeypatch.setattr(data_exporter.ZipFile, "close", close)
    with pytest.raises(data_exporter.ExportError) as failure:
        make_exporter(db).export_all_formats(tmp_path / "out")
    assert [(r["category"], r["error"]) for r in failure.value.outcome.failed] == [
        ("disk_full", "[Errno 28] original disk full")]
    assert "secondary close failure" in caplog.text
    assert not list((tmp_path / "out").glob(".report-*.tmp"))



@pytest.mark.parametrize("ai_on", [False, True])
@pytest.mark.parametrize("empty_registry", [False, True])
def test_orphan_page_recovery_is_independent_of_ai_and_registry_size(tmp_path, ai_on, empty_registry):
    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)])
    controller = SQLiteDatabaseController(db)
    record_configuration(controller, ai_on)
    controller.close()
    with sqlite3.connect(db) as con:
        if empty_registry:
            con.execute("DELETE FROM file_registry")
        else:
            con.execute("UPDATE page_log SET file_id=999")
        expected = con.execute("SELECT * FROM page_log ORDER BY page_id").fetchall()
    before = db.read_bytes()
    with pytest.raises(data_exporter.ExportError) as failure:
        make_exporter(db).export_all_formats(tmp_path / "out")
    outcome = failure.value.outcome
    assert outcome.recovery and any("Stored page 1 has no file record" in r["error"] for r in outcome.failed)
    assert not any(r["report"] in ("Results", "Workbook") for r in outcome.saved)
    book = openpyxl.load_workbook(next(r["path"] for r in outcome.saved if r["report"] == "Recovery workbook"))
    try:
        assert [[cell_text(v) for v in row] for row in sheet_rows(book, "Raw page_log")[1:]] == [
            [cell_text(v) for v in row] for row in expected]
    finally:
        book.close()
    assert db.read_bytes() == before



@pytest.mark.parametrize("recovery", [False, True])
def test_full_text_guidance_does_not_offer_an_unavailable_csv(tmp_path, monkeypatch, recovery):
    db = seed_one_file(tmp_path / "state.db", [(1, "one.jpg", "ok", "", None)],
                       requests=[request_outcome(pages=(1,), answer_rows=(answer_row(values={"answer": "x" * 500}),))])
    monkeypatch.setattr(data_exporter, "EXCEL_MAX_CELL_CHARS", 240)
    if recovery:
        with sqlite3.connect(db) as con:
            con.execute("DROP VIEW overall_result")
    else:
        original = data_exporter.SQLiteDataExporter._write_csv

        def fail_results(self, target, headers, rows):
            if "file_result" in headers:
                raise PermissionError("Results CSV unavailable")
            return original(self, target, headers, rows)

        monkeypatch.setattr(data_exporter.SQLiteDataExporter, "_write_csv", fail_results)
    with pytest.raises(data_exporter.ExportError) as failure:
        make_exporter(db).export_all_formats(tmp_path / "out")
    outcome = failure.value.outcome
    assert outcome.recovery == recovery
    assert any(r["format"] == "csv" for r in outcome.saved)
    book = openpyxl.load_workbook(next(r["path"] for r in outcome.saved if r["format"] == "xlsx"))
    try:
        summary = str(sheet_rows(book, "Summary"))
        assert "Full text for fields without a CSV remains in the run database." in summary
        assert "Use the listed CSVs for full text." not in summary
    finally:
        book.close()


@pytest.mark.parametrize("key, minimum_width, minimum_height", [
    ("WWWMMMMWWWMMMM", 35, 19),
    ("W" * 70, 60, 49),
    ("漢字" * 22, 60, 34),
    ("Широкий_заголовок_результата", 60, 34),
    ("first_line\nsecond_line", 32, 34),
])
def test_filtered_headers_allow_wide_glyphs_and_wrapping(tmp_path, key, minimum_width, minimum_height):
    db = seed_one_file(tmp_path / "run.db", [(1, "1.jpg", "ok", "", None)], [request_outcome(
        pages=(1,), answer_rows=(answer_row(values={key: "value"}),),
    )], declared=(key,))
    book = export_workbook(db, tmp_path / "out")
    for title in ("Results", "LLM Answers"):
        sheet = book[title]
        header = next(c for c in sheet[1] if excel_value(c.value) == "llm_" + key)
        assert isinstance(header, Cell)
        assert minimum_width <= sheet.column_dimensions[header.column_letter].width <= 60
        assert minimum_height <= sheet.row_dimensions[1].height <= 409
        assert header.alignment.wrap_text and header.font.bold
        assert sheet.auto_filter.ref and sheet.freeze_panes == "A2"
        assert sheet.cell(2, header.column).value == "value"
    book.close()


def test_fixed_and_recovery_headers_leave_filter_button_space(tmp_path):
    db = seed_one_file(tmp_path / "run.db", [(1, "1.jpg", "ok", "", None)])
    for recovery in (False, True):
        if recovery:
            with sqlite3.connect(db) as con:
                con.execute("DROP VIEW overall_result")
        exporter = make_exporter(db)
        directory = tmp_path / str(recovery)
        try:
            outcome = exporter.export_all_formats(directory)
        except data_exporter.ExportError as exc:
            assert recovery
            outcome = exc.outcome
        assert outcome.recovery is recovery
        workbook_path = next(item["path"] for item in outcome.saved if item["format"] == "xlsx")
        book = openpyxl.load_workbook(workbook_path)
        sheet = book["Raw page_log" if recovery else "Page Log"]
        for name in ("page_to_jpeg_comment", "video_frame_timestamp"):
            cell = next(c for c in sheet[1] if c.value == name)
            assert isinstance(cell, Cell)
            assert sheet.column_dimensions[cell.column_letter].width >= 27
        assert sheet.row_dimensions[1].height >= 19
        book.close()


def test_source_change_notes_compare_consecutive_attempts_of_one_path(tmp_path):
    db_path = tmp_path / "state.db"
    c = SQLiteDatabaseController(db_path)
    record_configuration(c, ai_on=False)
    attempts = [("a.png", "A"), ("a.png", "A"), ("b.png", "B"), ("a.png", "B"), ("a.png", ""), ("a.png", "C")]
    for file_id, (path, digest) in enumerate(attempts, start=1):
        c.handle_file_started(file_id, path, ".png", "Stub", source=schemas.SourceFingerprint(10, file_id, digest))
        c.handle_frame_saved(file_id, PageResult(1, f"{file_id}.jpg", "ok", ""))
        c.handle_file_completed(file_id, FileSummary(1, "1", "ok", "done"))
        c.finalize_file(file_id)
    c.close()
    expected = [
        ("a.png", "File attempt 1"),
        ("a.png", "File attempt 2 | Source file changed after this attempt"),
        ("b.png", ""),
        ("a.png", "File attempt 3 | Source file changed since attempt 2"),
        ("a.png", "File attempt 4"),
        ("a.png", "File attempt 5"),
    ]
    _, rows = read_csv_rows(export_results_csv(db_path, tmp_path / "csv"))
    assert [(row["file_path"], row["notes"]) for row in rows] == expected
    _, sheet = results_sheet(export_workbook(db_path, tmp_path / "xlsx"))
    assert [(row["file_path"], row["notes"] or "") for row in sheet] == expected
