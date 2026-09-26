# Copyright 2026 Vsevolod Belonogov
# SPDX-License-Identifier: Apache-2.0

import csv
import re
import sqlite3
import sys
from pathlib import Path

import openpyxl
import pytest

SRC = Path(__file__).resolve().parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from db_controller import SQLiteDatabaseController
from test_results_report import (
    add_file,
    answer_row,
    export_results_csv,
    export_workbook,
    record_configuration,
    request_outcome,
    store_requests,
)

COLUMN = "video_frame_timestamp"
MODES = ["table_per_page", "table_per_file"]
HMS = re.compile(r"^\d{2,}:\d{2}:\d{2}\.\d{2}$")



def pages_named(cell: str) -> list[int]:
    named: list[int] = []
    for part in filter(None, (p.strip() for p in cell.split(","))):
        first, _, last = part.partition("-")
        named.extend(range(int(first), int(last or first) + 1))
    return named


def expected_frame_time(page_log: list[tuple], file_id: str, pages_cell: str) -> str:
    named = pages_named(pages_cell)
    if len(named) != 1:
        return ""
    saved = [stamp for fid, number, status, stamp in page_log
             if str(fid) == file_id and number == named[0] and status == "ok"]
    return saved[0] if len(saved) == 1 else ""


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with open(path, encoding="utf-8", newline="") as handle:
        rows = list(csv.reader(handle))
    return rows[0], [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


def page_log_of(db_path: Path) -> list[tuple]:
    with sqlite3.connect(db_path) as con:
        return con.execute("SELECT file_id, page_number, page_to_jpeg_status, video_frame_timestamp "
                           "FROM page_log").fetchall()


def assert_results_follow_page_log(db_path: Path, csv_path: Path, workbook) -> list[dict[str, str]]:
    headers, rows = read_csv(csv_path)
    assert headers.count(COLUMN) == 1, headers
    assert headers.index(COLUMN) == headers.index("pages") + 1
    assert headers[headers.index(COLUMN) + 1] == "file_result"
    page_log = page_log_of(db_path)
    mismatches = [
        (r["file_id"], r["pages"], r[COLUMN], expected_frame_time(page_log, r["file_id"], r["pages"]))
        for r in rows if r[COLUMN] != expected_frame_time(page_log, r["file_id"], r["pages"])]
    assert mismatches == [], "(file_id, pages, shown, expected): " + repr(mismatches)
    sheet = workbook["Results"]
    sheet_headers = [cell.value for cell in sheet[1]]
    assert sheet_headers == headers
    column = sheet_headers.index(COLUMN) + 1
    cells = [sheet.cell(row, column) for row in range(2, sheet.max_row + 1)]
    assert [cell.value or "" for cell in cells] == [r[COLUMN] for r in rows]
    for cell in cells:
        if cell.value:
            assert isinstance(cell.value, str) and cell.data_type == "s", (cell.value, cell.data_type)
            assert HMS.match(cell.value), cell.value
    return rows


def export_both(db_path: Path, out: Path):
    return export_results_csv(db_path, out / "csv"), export_workbook(db_path, out / "xlsx")



def video_pages(times: list[float], name: str = "clip") -> list[tuple]:
    return [(n, f"{name}_page_{n}.jpg", "ok", "", t) for n, t in enumerate(times, 1)]


@pytest.mark.parametrize("mode", MODES)
def test_each_row_takes_its_own_attempts_frame_and_page_numbers_never_cross_match(tmp_path, mode):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    first = [float(n) for n in range(12)]
    second = [100.0 + n for n in range(12)]
    other = [50.25 + n for n in range(12)]
    for file_id, path, times in ((1, "clip.mp4", first), (2, "clip.mp4", second), (3, "other.mp4", other)):
        add_file(c, file_id, path, video_pages(times, f"{file_id}_clip"))
        if mode == "table_per_page":
            store_requests(c, file_id, [request_outcome(request_number=1, pages=(1, 10, 11), answer_rows=tuple(
                answer_row(p, str(p), "", {"answer": f"{file_id}-{p}"}) for p in (1, 10, 11)))])
        else:
            store_requests(c, file_id, [
                request_outcome(request_number=n, pages=(p,), answer_rows=(
                    answer_row(None, "", "", {"answer": f"{file_id}-{p}"}),))
                for n, p in enumerate((1, 10, 11), 1)])
    record_configuration(c, True, output_mode=mode)
    c.close()

    csv_path, workbook = export_both(db, tmp_path)
    rows = assert_results_follow_page_log(db, csv_path, workbook)
    assert [(r["file_id"], r["pages"], r[COLUMN]) for r in rows] == [
        ("1", "1", "00:00:00.00"), ("1", "10", "00:00:09.00"), ("1", "11", "00:00:10.00"),
        ("2", "1", "00:01:40.00"), ("2", "10", "00:01:49.00"), ("2", "11", "00:01:50.00"),
        ("3", "1", "00:00:50.25"), ("3", "10", "00:00:59.25"), ("3", "11", "00:01:00.25"),
    ]


def test_every_per_page_answer_kind_shows_the_time_of_the_page_it_names(tmp_path):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    add_file(c, 1, "clip.mp4", video_pages([0.0, 1.5, 3.0, 4.5]))
    store_requests(c, 1, [
        request_outcome(request_number=1, pages=(1, 2), answer_rows=(
            answer_row(1, "1", "", {"answer": "a"}),
            answer_row(1, "1", "duplicate_page_number", {"answer": "again"}),
            answer_row(None, "3", "hallucinated_page_number", {"answer": "claimed"}),
            answer_row(2, "", "no_row_returned", None),
        ), status="invalid_json_answer"),
        request_outcome(request_number=2, pages=(3, 4), answer_rows=(
            answer_row(3, "3", "", {"answer": "c"}), answer_row(4, "4", "", {"answer": "d"}))),
    ])
    record_configuration(c, True, output_mode="table_per_page")
    c.close()

    csv_path, workbook = export_both(db, tmp_path)
    rows = assert_results_follow_page_log(db, csv_path, workbook)
    assert [(r["pages"], r[COLUMN], r["llm_answer"]) for r in rows] == [
        ("1", "00:00:00.00", "a"),
        ("1", "00:00:00.00", "again"),
        ("", "", "claimed"),
        ("2", "00:00:01.50", ""),
        ("3", "00:00:03.00", "c"),
        ("4", "00:00:04.50", "d"),
    ]


ANSWERLESS_STATUSES = ["ok", "aborted_by_user", "encoding_failure", "token_limit_exceeded",
                       "network_failure", "provider_reply_parse_error", "invalid_json_answer",
                       "not_attempted"]


@pytest.mark.parametrize("mode", MODES)
@pytest.mark.parametrize("status", ANSWERLESS_STATUSES)
def test_an_answerless_single_frame_request_shows_its_frames_time(tmp_path, mode, status):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    add_file(c, 1, "clip.mp4", video_pages([0.0, 2.0, 4.0]))
    store_requests(c, 1, [
        request_outcome(request_number=1, pages=(1, 2), answer_rows=(
            answer_row(*((1, "1") if mode == "table_per_page" else (None, "")), "", {"answer": "a"}),
            *((answer_row(2, "2", "", {"answer": "b"}),) if mode == "table_per_page" else ()))),
        request_outcome(request_number=2, pages=(3,), status=status,
                        error="" if status == "ok" else "detail"),
    ])
    record_configuration(c, True, output_mode=mode)
    c.close()

    csv_path, workbook = export_both(db, tmp_path)
    rows = assert_results_follow_page_log(db, csv_path, workbook)
    assert (rows[-1]["pages"], rows[-1][COLUMN]) == ("3", "00:00:04.00")


@pytest.mark.parametrize("mode", MODES)
def test_both_answer_modes_agree_when_a_saved_frame_also_has_a_failure_row(tmp_path, mode):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    add_file(c, 1, "clip.mp4", [*video_pages([2.5, 5.0]), (1, "", "failure", "PyAV critical failure", None)],
             range_status="failure")
    if mode == "table_per_page":
        answers = (answer_row(1, "1", "", {"answer": "a"}), answer_row(2, "2", "", {"answer": "b"}))
        store_requests(c, 1, [request_outcome(request_number=1, pages=(1, 2), answer_rows=answers)])
    else:
        store_requests(c, 1, [
            request_outcome(request_number=n, pages=(p,), answer_rows=(answer_row(None, "", "", {"answer": a}),))
            for n, (p, a) in enumerate(((1, "a"), (2, "b")), 1)])
    record_configuration(c, True, output_mode=mode)
    c.close()

    csv_path, workbook = export_both(db, tmp_path)
    rows = assert_results_follow_page_log(db, csv_path, workbook)
    assert [(r["pages"], r[COLUMN]) for r in rows] == [("1", "00:00:02.50"), ("2", "00:00:05.00")]
    with sqlite3.connect(db) as con:
        linked = con.execute("SELECT p.page_to_jpeg_status FROM llm_answers a "
                             "JOIN page_log p ON p.page_id = a.page_id").fetchall()
    assert all(status == "ok" for (status,) in linked), linked



@pytest.fixture
def real_run(tmp_path, monkeypatch):
    import central_logger
    from fake_llm import make_server
    from test_llm_run_cruelty import (
        HonestModel, make_png, make_tiff, make_video, point_tokens_at_the_fake, run_core, settings_for)

    monkeypatch.setattr(central_logger, "setup_logging", lambda *a, **k: None)
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    make_video(input_dir / "c_clip.mp4")
    make_png(input_dir / "b_photo.png")
    make_tiff(input_dir / "a_doc.tif", page_count=3)

    def run(mode):
        if mode == "off":
            from config_validator import Settings
            settings = Settings(INPUT_FOLDER_PATH=input_dir, OUTPUT_FOLDER_PATH=tmp_path / "output",
                                ENABLE_LLM_INFERENCE=False, VIDEO_MODE="SUMMARY",
                                VIDEO_SUMMARY_TARGET_TOTAL_FRAMES=5, VIDEO_SUMMARY_SCENE_SENSITIVITY=0.0,
                                START_OVER=False)
            events, db_path = run_core(settings)
            return events, db_path, settings.TECH_FOLDER_PATH, 0
        point_tokens_at_the_fake(monkeypatch, "openai")
        server = make_server("openai").start()
        try:
            server.content = HonestModel(mode)
            settings = settings_for((input_dir, tmp_path / "output"), "openai", server.base_url,
                                    LLM_OUTPUT_MODE=mode, LLM_OUTPUT_COLUMNS="genre,n",
                                    MAX_JPEGS_PER_INFERENCE=4)
            events, db_path = run_core(settings)
            return events, db_path, settings.TECH_FOLDER_PATH, len(server.requests)
        finally:
            server.stop()

    return run


def saved_exports(tech_folder: Path):
    csv_path = sorted(tech_folder.glob("results_*.csv"))[-1]
    workbook = openpyxl.load_workbook(sorted(tech_folder.glob("database_export_*.xlsx"))[-1])
    return csv_path, workbook


@pytest.mark.parametrize("mode", MODES)
def test_a_real_run_shows_page_log_times_for_every_row_naming_one_frame(real_run, mode):
    events, db_path, tech, request_count = real_run(mode)
    assert events[-1] == {"type": "done"}, events[-3:]
    assert request_count == 1 + 1 + 2
    csv_path, workbook = saved_exports(tech)
    rows = assert_results_follow_page_log(db_path, csv_path, workbook)

    stamps = sorted(stamp for _, _, status, stamp in page_log_of(db_path) if stamp and status == "ok")
    assert len(stamps) == 5 and stamps[0] == "00:00:00.00", stamps
    video = [r for r in rows if r["file_path"] == "c_clip.mp4"]
    others = [r for r in rows if r["file_path"] != "c_clip.mp4"]
    assert all(r[COLUMN] == "" for r in others), others
    if mode == "table_per_page":
        assert sorted(r[COLUMN] for r in video) == stamps
    else:
        assert [(r["pages"], r[COLUMN]) for r in video] == [("1-4", ""), ("5", stamps[-1])]


def test_an_ai_off_run_keeps_results_without_the_column_and_page_log_with_it(real_run):
    events, _, tech, _ = real_run("off")
    assert events[-1] == {"type": "done"}, events[-3:]
    csv_path, workbook = saved_exports(tech)
    headers, _ = read_csv(csv_path)
    assert headers == ["file_id", "file_path", "file_result", "notes"]
    assert [cell.value for cell in workbook["Results"][1]] == headers
    page_log = workbook["Page Log"]
    stamp_column = [cell.value for cell in page_log[1]].index(COLUMN)
    stamps = [row[stamp_column] for row in page_log.iter_rows(min_row=2, values_only=True)]
    assert sum(1 for stamp in stamps if stamp) == 5
