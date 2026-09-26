# Copyright 2026 Vsevolod Belonogov
# SPDX-License-Identifier: Apache-2.0

import csv
import json
import re
import sys
from pathlib import Path

SRC = Path(__file__).resolve().parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

import openpyxl
import pytest
from openpyxl.utils import get_column_letter
from PIL import Image

from test_results_report import (
    add_file, answer_row, export_results_csv, export_workbook, record_configuration,
    request_outcome, store_requests)
from test_llm_run_cruelty import HonestModel, make_png, make_tiff, point_tokens_at_the_fake, run_core, settings_for
from db_controller import SQLiteDatabaseController
import central_logger
from fake_llm import make_server

av = pytest.importorskip("av")
np = pytest.importorskip("numpy")

COLUMN = "video_frame_timestamp"
FPS = 7
FRAMES = 30
LABEL = re.compile(r"page (\d+), extracted at (\d{2}:\d{2}:\d{2}\.\d{2})")


def grey(index: int) -> int:
    return 20 + 7 * index


def true_time(index: int) -> str:
    hundredths = round(index * 100 / FPS)
    return f"00:00:{hundredths // 100:02d}.{hundredths % 100:02d}"


def make_numbered_video(path: Path) -> None:
    with av.open(str(path), mode="w") as container:
        stream = container.add_stream("mpeg4", rate=FPS)
        stream.width, stream.height = 64, 64
        stream.pix_fmt = "yuv420p"
        for i in range(FRAMES):
            pixels = np.full((64, 64, 3), grey(i), dtype=np.uint8)
            for packet in stream.encode(av.VideoFrame.from_ndarray(pixels, format="rgb24")):
                container.mux(packet)
        for packet in stream.encode():
            container.mux(packet)


def decoded_levels(video: Path) -> list[float]:
    with av.open(str(video)) as container:
        levels = [float(np.asarray(frame.to_image().convert("L"), dtype=np.float64).mean())
                  for frame in container.decode(video=0)]
    assert len(levels) == FRAMES
    return levels


def frame_index_shown(jpeg: Path, levels: list[float]) -> int:
    with Image.open(jpeg) as image:
        mean = float(np.asarray(image.convert("L"), dtype=np.float64).mean())
    index = min(range(len(levels)), key=lambda i: abs(levels[i] - mean))
    assert abs(levels[index] - mean) < 2, (jpeg.name, mean, levels[index])
    return index


def make_gif(path: Path) -> None:
    frames = [Image.new("RGB", (32, 32), (40 * i, 200 - 40 * i, 90)) for i in range(3)]
    frames[0].save(path, save_all=True, append_images=frames[1:], duration=100)


def make_pdf(path: Path) -> None:
    pages = [Image.new("RGB", (64, 64), (230, 230 - 60 * i, 30)) for i in range(2)]
    pages[0].save(path, "PDF", save_all=True, append_images=pages[1:])


def csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with open(path, encoding="utf-8", newline="") as handle:
        rows = list(csv.reader(handle))
    return rows[0], [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


def newest(folder: Path, prefix: str) -> Path:
    matches = sorted(folder.glob(f"{prefix}_*"))
    assert matches, f"no {prefix} report was written"
    return matches[-1]


def workbook_results(workbook) -> tuple[list[str], list[list]]:
    sheet = workbook["Results"]
    headers = [cell.value for cell in sheet[1]]
    return headers, [list(row) for row in sheet.iter_rows(min_row=2)]



@pytest.fixture
def sandbox(tmp_path, monkeypatch):
    monkeypatch.setattr(central_logger, "setup_logging", lambda *a, **k: None)
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    return input_dir, tmp_path / "output"


REAL_CASES = [
    ("table_per_page", 4, "SUMMARY"),
    ("table_per_page", 4, "SAMPLING"),
    ("table_per_file", 1, "SUMMARY"),
    ("table_per_file", 4, "SUMMARY"),
]


@pytest.mark.parametrize(("mode", "per_request", "video_mode"), REAL_CASES)
def test_results_shows_the_true_time_of_the_one_frame_a_row_names(
        sandbox, monkeypatch, mode, per_request, video_mode):
    input_dir, output_dir = sandbox
    make_numbered_video(input_dir / "e_clip.mp4")
    make_png(input_dir / "b_photo.png")
    make_tiff(input_dir / "a_scan.tif", page_count=2)
    make_gif(input_dir / "c_anim.gif")
    make_pdf(input_dir / "d_doc.pdf")
    point_tokens_at_the_fake(monkeypatch, "openai")
    server = make_server("openai").start()
    try:
        server.content = HonestModel(mode)
        settings = settings_for(
            sandbox, "openai", server.base_url, LLM_OUTPUT_MODE=mode, LLM_OUTPUT_COLUMNS="genre,n",
            MAX_JPEGS_PER_INFERENCE=per_request, VIDEO_MODE=video_mode,
            VIDEO_SUMMARY_TARGET_TOTAL_FRAMES=6, VIDEO_SAMPLING_CAPTURE_RATE_FPS=2,
            VIDEO_SAMPLING_MAX_FRAMES_BUDGET=6, VIDEO_SAMPLING_SCENE_SENSITIVITY=0.0,
            ANIMATION_SCENE_SENSITIVITY=0.0)
        events, _ = run_core(settings)
        told: dict[int, set[str]] = {}
        for request in server.requests:
            for page, stamp in LABEL.findall(json.dumps(request.json)):
                told.setdefault(int(page), set()).add(stamp)
    finally:
        server.stop()
    assert events[-1] == {"type": "done"}, events[-3:]

    tech = settings.TECH_FOLDER_PATH
    headers, rows = csv_rows(newest(tech, "results"))
    assert headers == ["file_id", "file_path", "pages", COLUMN, "file_result", "llm_genre", "llm_n", "notes"]

    _, page_log = csv_rows(newest(tech, "page_log"))
    video_id = next(r["file_id"] for r in rows if r["file_path"] == "e_clip.mp4")
    jpegs = {p.name: p for p in output_dir.rglob("*.jpg")}
    levels = decoded_levels(input_dir / "e_clip.mp4")
    shown = {int(r["page_number"]): frame_index_shown(jpegs[r["output_file"]], levels)
             for r in page_log if r["file_id"] == video_id and r["page_to_jpeg_status"] == "ok"}
    assert len(shown) >= 5 and 0 in shown.values(), shown
    assert told.keys() == shown.keys(), (told, shown)

    video = [r for r in rows if r["file_id"] == video_id]
    single = [r for r in video if re.fullmatch(r"\d+", r["pages"])]
    for r in video:
        if r in single:
            page = int(r["pages"])
            assert r[COLUMN] == true_time(shown[page]), (r["pages"], r[COLUMN], shown[page])
            assert told[page] == {r[COLUMN]}, (page, told[page], r[COLUMN])
        else:
            assert r[COLUMN] == "", r
    assert all(r[COLUMN] == "" for r in rows if r["file_id"] != video_id), rows
    if per_request == 1 or mode == "table_per_page":
        assert sorted(int(r["pages"]) for r in single) == sorted(shown), video
        assert "00:00:00.00" in {r[COLUMN] for r in single}, video
    else:
        assert len(single) == (1 if len(shown) % per_request == 1 else 0), video

    workbook = openpyxl.load_workbook(newest(tech, "database_export"))
    sheet_headers, sheet_rows = workbook_results(workbook)
    assert sheet_headers == headers
    column = headers.index(COLUMN)
    assert [row[column].value or "" for row in sheet_rows] == [r[COLUMN] for r in rows]
    assert all(row[column].data_type == "s" for row in sheet_rows if row[column].value)



def video_pages(times):
    return [(n, f"clip_page_{n}.jpg", "ok", "", t) for n, t in enumerate(times, 1)]


@pytest.mark.parametrize("mode", ["table_per_page", "table_per_file"])
def test_rows_naming_several_pages_or_none_stay_empty_whatever_the_spelling(tmp_path, mode):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    add_file(c, 1, "clip.mp4", video_pages([0.5 * n for n in range(11)]))
    per_page = mode == "table_per_page"
    outcomes = [
        request_outcome(request_number=1, pages=(1, 2), answer_rows=(
            (answer_row(1, "1", "", {"answer": "x"}), answer_row(2, "2", "", {"answer": "y"}))
            if per_page else (answer_row(None, "", "", {"answer": "x"}),))),
        request_outcome(request_number=2, pages=(3, 5), status="network_failure", error="refused"),
        request_outcome(request_number=3, pages=(11,), status="network_failure", error="refused"),
        request_outcome(request_number=4, pages=(), status="not_attempted"),
    ]
    store_requests(c, 1, outcomes)
    record_configuration(c, True, output_mode=mode)
    c.close()

    _, rows = csv_rows(export_results_csv(db, tmp_path / "csv"))
    expected = ([("1", "00:00:00.00"), ("2", "00:00:00.50")] if per_page else [("1-2", "")])
    assert [(r["pages"], r[COLUMN]) for r in rows] == [*expected, ("3, 5", ""), ("11", "00:00:05.00"), ("", "")]


@pytest.mark.parametrize("ai", ["off", "table_per_page", "table_per_file"])
def test_the_header_is_exact_on_an_empty_run_in_both_formats(tmp_path, ai):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    if ai == "off":
        record_configuration(c, False)
        expected = ["file_id", "file_path", "file_result", "notes"]
    else:
        record_configuration(c, True, declared=("a", "b"), output_mode=ai)
        expected = ["file_id", "file_path", "pages", COLUMN, "file_result", "llm_a", "llm_b", "notes"]
    c.close()
    headers, rows = csv_rows(export_results_csv(db, tmp_path / "csv"))
    assert (headers, rows) == (expected, [])
    assert workbook_results(export_workbook(db, tmp_path / "xlsx"))[0] == expected


def test_a_declared_key_named_like_the_fixed_column_keeps_its_own_cell(tmp_path):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    add_file(c, 1, "clip.mp4", video_pages([7.0]))
    store_requests(c, 1, [request_outcome(request_number=1, pages=(1,), answer_rows=(
        answer_row(1, "1", "", {COLUMN: "model says 99:00:00.00", "pages": "model pages"}),))])
    record_configuration(c, True, declared=(COLUMN, "pages"), output_mode="table_per_page")
    c.close()
    _, rows = csv_rows(export_results_csv(db, tmp_path / "csv"))
    assert [(r["pages"], r[COLUMN], r[f"llm_{COLUMN}"], r["llm_pages"]) for r in rows] == [
        ("1", "00:00:07.00", "model says 99:00:00.00", "model pages")]


def test_the_workbook_filter_and_width_cover_the_new_column(tmp_path):
    db = tmp_path / "state.db"
    c = SQLiteDatabaseController(db)
    add_file(c, 1, "clip.mp4", video_pages([3723.45]))
    store_requests(c, 1, [request_outcome(request_number=1, pages=(1,), answer_rows=(
        answer_row(None, "", "", {"answer": "a"}),))])
    record_configuration(c, True)
    c.close()
    sheet = export_workbook(db, tmp_path / "xlsx")["Results"]
    headers = [cell.value for cell in sheet[1]]
    letter = get_column_letter(headers.index(COLUMN) + 1)
    assert sheet[f"{letter}2"].value == "01:02:03.45"
    last = get_column_letter(len(headers))
    assert sheet.auto_filter.ref == f"A1:{last}{sheet.max_row}"
    assert sheet.column_dimensions[letter].width >= len(COLUMN)
