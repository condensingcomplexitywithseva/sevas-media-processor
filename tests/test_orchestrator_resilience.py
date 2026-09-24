# Copyright 2026 Vsevolod Belonogov
# SPDX-License-Identifier: Apache-2.0

import logging
import sqlite3
import sys
import threading
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from sqlalchemy.exc import SQLAlchemyError

SRC = Path(__file__).resolve().parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from batch_orchestrator import BatchOrchestrator
from db_controller import SQLiteDatabaseController
from fs_utils import source_sha256, source_state
from media_classifier import MediaClassifier
from schemas import (
    RequestOutcome,
    ConfigurationError,
    FileSummary,
    OverallResult,
    PageResult,
    SourceFingerprint,
    Status,
)

LLM_FAIL = [RequestOutcome(1, (1,), "network_failure", "", "boom")]
LLM_OK = [RequestOutcome(1, (1,), "ok", "a fine answer", "")]


def make_settings(input_dir, **overrides) -> SimpleNamespace:
    values: dict[str, Any] = {
        "INPUT_FOLDER_PATH": input_dir,
        "NO_RETRY_STATUSES": [Status.OK],
        "JPEG_QUALITY": 90,
        "MAX_DIMENSION": 4096,
        "ENABLE_LLM_INFERENCE": False,
        "MAX_CONSECUTIVE_LLM_FAILURES": 3,
        "MAX_JPEGS_PER_INFERENCE": 10,
    }
    values.update(overrides)
    return SimpleNamespace(**values)


class RecordingDb:

    def __init__(self, frames_available=((1, "frame.jpg"),)):
        self.frames_available = [(n, Path(p), "") for n, p in frames_available]
        self.started = []
        self.frames = []
        self.completed = []
        self.llm_requests = []

    def finalize_file(self, file_id, expected_requests=0):
        pass

    def record_file_interruption(self, file_id, reason):
        self.interruption = (file_id, reason)

    def get_highest_file_id(self):
        return 0

    def get_no_retry_sources(self, results):
        return {}

    def handle_file_started(self, file_id, rel_path, ext, pipeline_name, **kwargs):
        self.started.append((file_id, rel_path))

    def handle_frame_saved(self, file_id, page_result):
        self.frames.append((file_id, page_result))

    def handle_file_completed(self, file_id, summary):
        self.completed.append((file_id, summary))

    def handle_llm_requests(self, file_id, outcomes):
        if self.llm_requests and self.llm_requests[-1][0] == file_id:
            self.llm_requests[-1][1].extend(outcomes)
        else:
            self.llm_requests.append((file_id, list(outcomes)))

    def get_file_statuses(self, file_id):
        return {}

    def get_successful_frames(self, file_id, output_folder):
        return list(self.frames_available)


class StubLogger:
    def __init__(self):
        self.app_logger = logging.getLogger("test-orchestrator-resilience")

    def log_file_started(self, *args):
        pass

    def log_frame_saved(self, *args):
        pass

    def log_file_completed(self, *args):
        pass

    def log_llm_completed(self, *args):
        pass

    def log_critical_error(self, *args):
        pass


class StubRouter:

    output_folder = Path(".")
    relative_or_orphan = staticmethod(MediaClassifier.relative_or_orphan)

    def __init__(self):
        self.routed = []

    def evaluate_and_route(self, file_id, path, root):
        self.routed.append(path.name)
        rel, orphaned = self.relative_or_orphan(path, root)
        return rel, path.suffix, "Stub", self.make_generator(path), orphaned

    def make_generator(self, path):
        def gen():
            yield PageResult(1, f"{path.stem}.jpg", Status.OK.value, "")
            return FileSummary(1, "1", Status.OK.value, "done")

        return gen()


class CountingLLM:

    def __init__(self, results=()):
        self.results = list(results)
        self.calls = 0

    def execute_network_inference(self, frames, abort_flag=None, *, on_outcome=None):
        assert on_outcome is not None
        self.calls += 1
        outcomes = self.results.pop(0)
        for outcome in outcomes:
            on_outcome(outcome)
        return outcomes


class ForbiddenLLM:

    def __init__(self):
        self.calls = 0

    def execute_network_inference(self, frames, abort_flag=None, *, on_outcome=None):
        assert on_outcome is not None
        self.calls += 1
        for outcome in LLM_OK:
            on_outcome(outcome)
        return LLM_OK


def fingerprint(path):
    state, digest = source_state(path), source_sha256(path)
    assert state is not None and digest
    return SourceFingerprint(state[0], state[1], digest)


def make_input(tmp_path, *names):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    for name in names:
        (input_dir / name).write_bytes(b"x")
    return input_dir



def test_resume_skips_completed_files_before_any_pipeline_is_built(tmp_path):
    input_dir = make_input(tmp_path, "done.png", "new.png")

    db = SQLiteDatabaseController(tmp_path / "state.db")
    db.handle_file_started(7, "done.png", ".png", "Stub", source=fingerprint(input_dir / "done.png"))
    db.handle_frame_saved(7, PageResult(1, "done.jpg", Status.OK.value, ""))
    db.handle_file_completed(7, FileSummary(1, "1", "ok", "from a previous run"))
    db.finalize_file(7)

    router = StubRouter()
    orchestrator = BatchOrchestrator(
        make_settings(input_dir), db, router, StubLogger(), None
    )
    orchestrator.execute_batch_processing_loop()
    db.close()

    assert router.routed == ["new.png"]

    connection = sqlite3.connect(tmp_path / "state.db")
    try:
        rows = connection.execute(
            "SELECT file_id, file_path FROM file_registry ORDER BY file_id"
        ).fetchall()
    finally:
        connection.close()
    assert rows == [(7, "done.png"), (8, "new.png")]



def test_empty_input_folder_ends_gracefully_at_100(tmp_path, caplog):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    db = RecordingDb()
    events = []

    with caplog.at_level(logging.WARNING):
        BatchOrchestrator(
            make_settings(input_dir), db, StubRouter(), StubLogger(), None
        ).execute_batch_processing_loop(on_progress=lambda e: events.append(e))

    assert events == [{"type": "progress", "value": 100}]
    assert db.started == [] and db.completed == []
    assert "empty" in caplog.text.lower()



def test_circuit_breaker_trips_after_consecutive_llm_failures(tmp_path):
    input_dir = make_input(tmp_path, "a.png", "b.png", "c.png", "d.png")
    llm = CountingLLM([LLM_FAIL, LLM_FAIL, LLM_FAIL, LLM_OK])
    settings = make_settings(
        input_dir, ENABLE_LLM_INFERENCE=True, MAX_CONSECUTIVE_LLM_FAILURES=3
    )

    with pytest.raises(ConfigurationError, match="CIRCUIT BREAKER"):
        BatchOrchestrator(
            settings, RecordingDb(), StubRouter(), StubLogger(), llm
        ).execute_batch_processing_loop()

    assert llm.calls == 3


def test_one_llm_success_resets_the_failure_counter(tmp_path):
    input_dir = make_input(tmp_path, "a.png", "b.png", "c.png", "d.png", "e.png")
    llm = CountingLLM([LLM_FAIL, LLM_FAIL, LLM_OK, LLM_FAIL, LLM_FAIL])
    settings = make_settings(
        input_dir, ENABLE_LLM_INFERENCE=True, MAX_CONSECUTIVE_LLM_FAILURES=3
    )

    BatchOrchestrator(
        settings, RecordingDb(), StubRouter(), StubLogger(), llm
    ).execute_batch_processing_loop()

    assert llm.calls == 5


def test_content_failures_never_trip_the_network_breaker(tmp_path):
    input_dir = make_input(tmp_path, "a.png", "b.png", "c.png", "d.png")
    content_fail = [RequestOutcome(1, (1,), "invalid_json_answer", "not json", "")]
    llm = CountingLLM([content_fail] * 4)
    settings = make_settings(
        input_dir, ENABLE_LLM_INFERENCE=True, MAX_CONSECUTIVE_LLM_FAILURES=3
    )

    with pytest.raises(ConfigurationError) as excinfo:
        BatchOrchestrator(
            settings, RecordingDb(), StubRouter(), StubLogger(), llm
        ).execute_batch_processing_loop()

    assert "valid JSON" in str(excinfo.value)
    assert "no successful reply in between" not in str(excinfo.value)
    assert llm.calls == 3


@pytest.mark.parametrize("neutral_status", [
    "aborted_by_user", "token_limit_exceeded", "encoding_failure",
    "not_attempted",
])
def test_neutral_llm_outcomes_hold_the_breaker_counter(tmp_path, neutral_status):
    input_dir = make_input(tmp_path, "a.png", "b.png", "c.png", "d.png", "e.png")
    neutral = [RequestOutcome(1, (1,), neutral_status, "", "detail")]
    llm = CountingLLM([LLM_FAIL, LLM_FAIL, neutral, LLM_FAIL, LLM_OK])
    db = RecordingDb()
    settings = make_settings(
        input_dir, ENABLE_LLM_INFERENCE=True, MAX_CONSECUTIVE_LLM_FAILURES=3
    )

    with pytest.raises(ConfigurationError, match="CIRCUIT BREAKER"):
        BatchOrchestrator(
            settings, db, StubRouter(), StubLogger(), llm
        ).execute_batch_processing_loop()

    assert llm.calls == 4
    assert [file_id for file_id, _ in db.llm_requests] == [1, 2, 3, 4]


def test_a_clean_request_amid_network_failures_resets_the_breaker(tmp_path):
    input_dir = make_input(tmp_path, "a.png", "b.png", "c.png", "d.png")
    mixed = [
        RequestOutcome(1, tuple(range(1, 6)), "ok", "a fine answer", ""),
        RequestOutcome(2, (6, 7, 8, 9), "network_failure", "", "boom"),
    ]
    llm = CountingLLM([mixed] * 4)
    settings = make_settings(
        input_dir, ENABLE_LLM_INFERENCE=True, MAX_CONSECUTIVE_LLM_FAILURES=3
    )

    BatchOrchestrator(
        settings, RecordingDb(), StubRouter(), StubLogger(), llm
    ).execute_batch_processing_loop()

    assert llm.calls == 4



def test_file_with_no_valid_frames_bypasses_the_ai_without_a_network_call(tmp_path):
    input_dir = make_input(tmp_path, "broken.png")
    db = RecordingDb(frames_available=())
    settings = make_settings(input_dir, ENABLE_LLM_INFERENCE=True)
    llm = ForbiddenLLM()

    BatchOrchestrator(
        settings, db, StubRouter(), StubLogger(), llm
    ).execute_batch_processing_loop()

    assert llm.calls == 0
    assert db.llm_requests == []


def test_orphaned_file_gets_the_orphan_note_exactly_once(tmp_path):
    input_dir = make_input(tmp_path, "stray.png")

    class OrphanRouter(StubRouter):
        def evaluate_and_route(self, file_id, path, root):
            rel, ext, name, generator, _ = super().evaluate_and_route(file_id, path, root)
            return rel, ext, name, generator, True

    db = RecordingDb()
    BatchOrchestrator(
        make_settings(input_dir), db, OrphanRouter(), StubLogger(), None
    ).execute_batch_processing_loop()

    assert len(db.completed) == 1
    comment = db.completed[0][1].file_to_jpegs_comment
    assert comment.count("Orphaned path fallback") == 1
    assert "stray.png" in comment



def test_crash_after_completion_never_overwrites_the_record(tmp_path):
    input_dir = make_input(tmp_path, "a.png", "b.png")

    class ExplodingPrepDb(RecordingDb):
        def get_successful_frames(self, file_id, output_folder):
            raise RuntimeError("AI prep exploded")

    db = ExplodingPrepDb()
    settings = make_settings(input_dir, ENABLE_LLM_INFERENCE=True)

    BatchOrchestrator(
        settings, db, StubRouter(), StubLogger(), ForbiddenLLM()
    ).execute_batch_processing_loop()


    assert [file_id for file_id, _ in db.completed] == [1, 2]
    assert all(summary.file_to_jpegs_comment == "done" for _, summary in db.completed)


def test_crash_before_completion_writes_a_fatal_crash_summary(tmp_path):
    input_dir = make_input(tmp_path, "torn.png")

    class TornRouter(StubRouter):
        def make_generator(self, path):
            def gen():
                yield PageResult(1, "torn.jpg", Status.OK.value, "")
                raise ValueError("decoder blew up mid-file")

            return gen()

    db = RecordingDb()
    BatchOrchestrator(
        make_settings(input_dir), db, TornRouter(), StubLogger(), None
    ).execute_batch_processing_loop()

    assert len(db.completed) == 1
    _, summary = db.completed[0]
    assert summary.range_status == Status.FAILURE.value
    assert summary.file_to_jpegs_comment.startswith("Fatal Orchestration Exception")
    assert "decoder blew up" in summary.file_to_jpegs_comment


def test_database_errors_escalate_to_runtime_error(tmp_path):
    input_dir = make_input(tmp_path, "a.png")

    class BrokenDb(RecordingDb):
        def handle_file_started(self, *args, **kwargs):
            raise SQLAlchemyError("disk gone")

    with pytest.raises(RuntimeError, match="Fatal DB Transaction Error"):
        BatchOrchestrator(
            make_settings(input_dir), BrokenDb(), StubRouter(), StubLogger(), None
        ).execute_batch_processing_loop()



def test_stop_during_extraction_prevents_the_llm_call(tmp_path):
    input_dir = make_input(tmp_path, "a.png")
    abort_flag = threading.Event()

    class AbortingRouter(StubRouter):
        def make_generator(self, path):
            def gen():
                yield PageResult(1, "a.jpg", Status.OK.value, "")
                abort_flag.set()
                yield PageResult(2, "b.jpg", Status.OK.value, "")
                return FileSummary(2, "1-2", Status.OK.value, "done")

            return gen()

    db = RecordingDb()
    settings = make_settings(input_dir, ENABLE_LLM_INFERENCE=True)
    llm = ForbiddenLLM()

    BatchOrchestrator(
        settings, db, AbortingRouter(), StubLogger(), llm
    ).execute_batch_processing_loop(abort_flag=abort_flag)

    assert llm.calls == 0
    assert len(db.completed) == 1
    assert "Aborted mid-extraction" in db.completed[0][1].file_to_jpegs_comment
    assert db.llm_requests == []
    assert db.interruption[0] == 1 and "Stopped" in db.interruption[1]


def test_halt_mid_llm_leg_still_records_the_carried_outcomes(tmp_path):
    input_dir = make_input(tmp_path, "a.png")
    carried = [
        RequestOutcome(1, (1,), "ok", "paid answer", ""),
        RequestOutcome(2, (2,), "network_failure", "", "FATAL AUTHENTICATION ERROR"),
        RequestOutcome(3, (3,), "not_attempted", "", "Not attempted: the batch halted."),
    ]

    class HaltingLLM:
        def recovery_identity(self):
            return "scripted-client"

        def execute_network_inference(self, frames, abort_flag=None, *, on_outcome=None):
            assert on_outcome is not None
            halt = ConfigurationError("FATAL AUTHENTICATION ERROR")
            halt.request_outcomes = list(carried)
            for outcome in carried:
                on_outcome(outcome)
            raise halt

    db = RecordingDb()
    settings = make_settings(input_dir, ENABLE_LLM_INFERENCE=True)

    with pytest.raises(ConfigurationError, match="AUTHENTICATION"):
        BatchOrchestrator(
            settings, db, StubRouter(), StubLogger(), HaltingLLM()
        ).execute_batch_processing_loop()

    assert db.llm_requests == [(1, carried)]


def test_crash_in_the_llm_leg_marks_the_ai_work_as_not_attempted(tmp_path):
    input_dir = make_input(tmp_path, "a.png")

    class CrashingLLM:
        def recovery_identity(self):
            return "scripted-client"

        def execute_network_inference(self, frames, abort_flag=None, *, on_outcome=None):
            assert on_outcome is not None
            raise RuntimeError("client blew up mid-request")

    db = RecordingDb()
    settings = make_settings(input_dir, ENABLE_LLM_INFERENCE=True)

    BatchOrchestrator(
        settings, db, StubRouter(), StubLogger(), CrashingLLM()
    ).execute_batch_processing_loop()

    assert [fid for fid, _ in db.completed] == [1]
    assert db.completed[0][1].file_to_jpegs_comment == "done"
    assert db.llm_requests == []
    assert db.interruption == (1, "Processing interrupted: client blew up mid-request")


def _file_restart_child(folder, endpoint, mode, cut, options):
    import json
    import os
    from PIL import Image
    from sqlalchemy import event
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from app_context import processing_identity, run_configuration_of
    from config_validator import Settings
    from test_llm_client_network import make_wire_client

    folder = Path(folder)
    options = json.loads(options)
    fixture = options.pop("fixture", {})
    flag = threading.Event()
    settings = Settings(INPUT_FOLDER_PATH=folder / "input", OUTPUT_FOLDER_PATH=folder / "output",
                        START_OVER=False, ENABLE_LLM_INFERENCE=mode != "off",
                        **{"MAX_JPEGS_PER_INFERENCE": 1, "LLM_OUTPUT_MODE": "table_per_file" if mode == "off" else mode,
                           "HALT_ON_LLM_PARSE_ERROR": False, "LLM_JSON_MAX_ATTEMPTS": 1, **options})
    client = make_wire_client(SimpleNamespace(url=endpoint), provider_overrides=fixture.get("provider", {}),
                             **{name: getattr(settings, name) for name in (
                                 "MAX_JPEGS_PER_INFERENCE", "LLM_OUTPUT_MODE", "LLM_JSON_MAX_ATTEMPTS",
                                 "LLM_USER_PROMPT", "LLM_USER_PROMPT_MODE", "LLM_SYSTEM_PROMPT",
                                 "LLM_SYSTEM_PROMPT_MODE", "LLM_TIMEOUT_SECONDS", "LLM_MAX_RETRIES",
                                 "LLM_RETRY_SLEEP_SECONDS", "HALT_ON_LLM_PARSE_ERROR", "LLM_ABORT_ON_MALFORMED_JSON")})
    db = SQLiteDatabaseController(folder / "state.db")
    recorded = db.get_run_configuration()
    config = run_configuration_of(settings)
    if recorded is not None and recorded != config:
        raise ConfigurationError("changed recorded configuration")
    if recorded is None:
        db.record_run_configuration(*config)
    db.ensure_run_identity(*processing_identity(settings, client if mode != "off" else None))

    class Router(StubRouter):
        output_folder = folder / "output"

        def evaluate_and_route(self, file_id, path, root):
            self.file_id = file_id
            with (folder / "routed.txt").open("a") as stream:
                stream.write(path.name + "\n")
            return super().evaluate_and_route(file_id, path, root)

        def make_generator(self, path):
            pages = fixture.get("pages", [1, 2, 3])
            for index, page in enumerate(pages):
                name = f"{self.file_id}-page-{page}.jpg"
                Image.new("RGB", (8, 8), (page % 255, 20, 30)).save(self.output_folder / name)
                failed = self.file_id > 1 and page == fixture.get("failed_page")
                yield PageResult(page, name, "failure" if failed else "ok", "page conversion failed" if failed else "")
                if self.file_id > 1 and index == 0:
                    if cut == "during_conversion":
                        os._exit(71)
                    if cut == "stop_conversion":
                        flag.set()
            return FileSummary(max(pages), "1-3", "ok", "done")

    def wrap(name, before=False):
        original = getattr(db, name)

        def call(file_id, *args, **kwargs):
            target = file_id > 1
            if before and target:
                os._exit(71)
            result = original(file_id, *args, **kwargs)
            if target:
                if cut.startswith("stop_"):
                    flag.set()
                else:
                    os._exit(71)
            return result
        setattr(db, name, call)

    if cut == "registered":
        wrap("handle_file_started")
    elif cut in ("converted", "stop_converted"):
        wrap("handle_file_completed")
    elif cut in ("first_saved", "stop_saved"):
        wrap("handle_llm_requests")
    elif cut in ("before_complete", "after_complete"):
        wrap("finalize_file", before=cut == "before_complete")
    elif cut in ("inside_answers", "inside_complete", "write_answers", "write_complete"):
        def interrupt_write(conn, cursor, statement, parameters, context, many):
            if cut.endswith("answers"):
                matched = statement.startswith("INSERT INTO llm_answers") and conn.exec_driver_sql(
                    "SELECT 1 FROM llm_answers WHERE file_id>1 LIMIT 1").first()
            else:
                matched = statement.startswith("UPDATE file_registry SET processing_complete") and conn.exec_driver_sql(
                    "SELECT 1 FROM file_registry WHERE file_id>1 AND processing_complete=1 LIMIT 1").first()
            if matched:
                if cut.startswith("inside"):
                    os._exit(71)
                raise SQLAlchemyError("injected storage failure")
        event.listen(db.sql_engine, "after_cursor_execute", interrupt_write)
    elif cut == "received":
        original = client._process_request
        calls = 0

        def received(*args, **kwargs):
            nonlocal calls
            result = original(*args, **kwargs)
            calls += 1
            if calls == 4:
                os._exit(71)
            return result
        client._process_request = received
    try:
        BatchOrchestrator(settings, db, Router(), StubLogger(), client if mode != "off" else None
                          ).execute_batch_processing_loop(abort_flag=flag)
    finally:
        db.close()


def _restart_child(folder, server, mode="table_per_file", cut="", check=True, **options):
    import json
    import subprocess
    child = "import runpy,sys; runpy.run_path(sys.argv[1])['_file_restart_child'](*sys.argv[2:])"
    command = [sys.executable, "-c", child,
               str(Path(__file__).resolve()), str(folder), server.url, mode, cut, json.dumps(options)]
    result = subprocess.run(command, capture_output=True, text=True, timeout=45,
                            cwd=Path(__file__).resolve().parents[1])
    if check:
        expected = 1 if cut.startswith("write_") else (71 if cut and not cut.startswith("stop_") else 0)
        assert result.returncode == expected, result.stdout + result.stderr
        if cut.startswith("write_"):
            assert "injected storage failure" in result.stderr
    return result


@pytest.fixture
def restart_case(tmp_path):
    import json
    import re
    from fake_llm.generic import GenericServer, openai_reply

    class Server(GenericServer):
        mode = "table_per_file"

        def handle_unknown_route(self):
            pages = [int(n) for n in re.findall(r"- page (\d+)", json.dumps(self.requests[-1].json))]
            number = len(self.requests)
            answer = {"answer": f"call-{number}-page-{pages[0]}"}
            if self.mode == "table_per_page":
                answer = [{"page": page, "answer": f"call-{number}-page-{page}"} for page in pages]
            return self.json_response(openai_reply(json.dumps(answer)))

    (tmp_path / "input").mkdir()
    (tmp_path / "output").mkdir()
    for name in ("A.tiff", "B.tiff"):
        (tmp_path / "input" / name).write_bytes(b"generated router input")
    with Server() as server:
        yield tmp_path, server


def _restart_rows(folder, query):
    from contextlib import closing
    with closing(sqlite3.connect(folder / "state.db")) as connection:
        return connection.execute(query).fetchall()


@pytest.mark.parametrize("mode", ["table_per_file", "table_per_page", "off"])
@pytest.mark.parametrize("cut", ["registered", "during_conversion", "converted", "before_complete",
                                  "inside_complete", "after_complete"])
def test_whole_file_restart_keeps_completed_files_and_separate_attempts(restart_case, mode, cut):
    folder, server = restart_case
    server.mode = mode
    _restart_child(folder, server, mode, cut)
    before = _restart_rows(folder, "SELECT * FROM llm_answers WHERE file_id=1 ORDER BY llm_answer_id")
    assert _restart_rows(folder, "SELECT processing_complete FROM file_registry ORDER BY file_id") == [
        (1,), (1 if cut == "after_complete" else 0,)]
    uploads = len(server.requests)
    _restart_child(folder, server, mode)
    repeat = cut != "after_complete"
    assert len(server.requests) == uploads + (3 if repeat and mode != "off" else 0)
    assert (folder / "routed.txt").read_text().splitlines() == ["A.tiff", "B.tiff"] + (["B.tiff"] if repeat else [])
    assert _restart_rows(folder, "SELECT * FROM llm_answers WHERE file_id=1 ORDER BY llm_answer_id") == before
    records = _restart_rows(folder, "SELECT file_id,file_path,processing_complete FROM file_registry ORDER BY file_id")
    assert records == (
        [(1, "A.tiff", 1), (2, "B.tiff", 0), (3, "B.tiff", 1)] if repeat else
        [(1, "A.tiff", 1), (2, "B.tiff", 1)])
    assert _restart_rows(folder, "PRAGMA integrity_check") == [("ok",)]


@pytest.mark.parametrize("mode", ["table_per_file", "table_per_page"])
@pytest.mark.parametrize("cut", ["first_saved", "received", "inside_answers", "write_answers", "write_complete"])
def test_interrupted_ai_repeats_all_requests_without_merging_saved_answers(restart_case, mode, cut):
    folder, server = restart_case
    server.mode = mode
    _restart_child(folder, server, mode, cut)
    old = _restart_rows(folder, "SELECT * FROM llm_answers ORDER BY llm_answer_id")
    assert _restart_rows(folder, "SELECT processing_complete FROM file_registry WHERE file_id=2") == [(0,)]
    if cut in ("received", "inside_answers", "write_answers"):
        assert not _restart_rows(folder, "SELECT * FROM llm_answers WHERE file_id=2")
    uploads = len(server.requests)
    _restart_child(folder, server, mode)
    assert len(server.requests) == uploads + 3
    assert _restart_rows(folder, "SELECT * FROM llm_answers WHERE file_id<3 ORDER BY llm_answer_id") == old
    values = _restart_rows(folder, "SELECT values_json FROM llm_answers WHERE file_id=3 ORDER BY llm_answer_id")
    assert values == [(f'{{"answer": "call-{uploads + n}-page-{n}"}}',) for n in (1, 2, 3)]
    assert _restart_rows(folder, "SELECT overall_result FROM overall_result WHERE file_id=3") == [("ok",)]


@pytest.mark.parametrize("mode", ["table_per_file", "table_per_page"])
@pytest.mark.parametrize("failed_page", [None, 2])
def test_partial_conversion_and_sparse_batches_restart_as_a_whole_file(restart_case, mode, failed_page):
    folder, server = restart_case
    server.mode = mode
    options = {"MAX_JPEGS_PER_INFERENCE": 2, "fixture": {"pages": [1, 2, 5, 9], "failed_page": failed_page}}
    _restart_child(folder, server, mode, "first_saved", **options)
    uploads = len(server.requests)
    _restart_child(folder, server, mode, **options)
    assert len(server.requests) == uploads + 2
    assert _restart_rows(folder, "SELECT processing_complete FROM file_registry WHERE file_id=3") == [(1,)]
    assert _restart_rows(folder, "SELECT overall_result FROM overall_result WHERE file_id=3") == [
        ("partial_fail" if failed_page else "ok",)]
    assert _restart_rows(folder, "SELECT file_id FROM llm_requests WHERE file_id>1 ORDER BY request_id") == [
        (2,), (3,), (3,)]


@pytest.mark.parametrize("cut", ["stop_conversion", "stop_converted", "stop_saved"])
def test_stop_keeps_finished_files_and_leaves_retryable_whole_file(restart_case, cut):
    folder, server = restart_case
    _restart_child(folder, server, cut=cut)
    assert len(server.requests) == (4 if cut == "stop_saved" else 3)
    assert _restart_rows(folder, "SELECT processing_complete FROM file_registry ORDER BY file_id") == [(1,), (0,)]
    uploads = len(server.requests)
    _restart_child(folder, server)
    assert len(server.requests) == uploads + 3


@pytest.mark.parametrize("change", [
    {"JPEG_QUALITY": 85}, {"DOCUMENT_RANGE": "2-3"}, {"MAX_JPEGS_PER_INFERENCE": 2},
    {"LLM_USER_PROMPT": "another task"}, {"fixture": {"provider": {"model": "another-model"}}},
    {"fixture": {"provider": {"extra_header_key": "API-Version", "extra_header_value": "2"}}},
    {"fixture": {"provider": {"response_extraction_path": "other.answer"}}},
    {"fixture": {"provider": {"image_payload_style": "base64_dict"}}},
    {"fixture": {"provider": {"reasoning_handling": "strip_xml"}}},
])
@pytest.mark.parametrize("cut,failed_page", [("", None), ("first_saved", None), ("first_saved", 2)])
def test_run_semantics_refuse_before_historical_skips_or_uploads(restart_case, change, cut, failed_page):
    folder, server = restart_case
    _restart_child(folder, server, cut=cut, fixture={"failed_page": failed_page})
    change = {**change, "fixture": {"failed_page": failed_page, **change.get("fixture", {})}}
    before = _restart_rows(folder, "SELECT * FROM file_registry ORDER BY file_id")
    uploads = len(server.requests)
    result = _restart_child(folder, server, check=False, **change)
    assert result.returncode != 0 and "err_resume_ai_work_changed" in result.stderr
    assert len(server.requests) == uploads
    assert _restart_rows(folder, "SELECT * FROM file_registry ORDER BY file_id") == before


@pytest.mark.parametrize("change", [
    {"LLM_TIMEOUT_SECONDS": 25}, {"LLM_MAX_RETRIES": 2}, {"LLM_RETRY_SLEEP_SECONDS": 0},
    {"LLM_JSON_MAX_ATTEMPTS": 3}, {"HALT_ON_LLM_PARSE_ERROR": True}, {"LLM_ABORT_ON_MALFORMED_JSON": True},
    {"LOGGING_LEVEL": "DEBUG"}, {"VIDEO_SAMPLING_MAX_FRAMES_BUDGET": 7},
    {"fixture": {"provider": {"max_tokens": 1234}}},
    {"fixture": {"provider": {"max_tokens_field": "unused_limit"}}},
    {"fixture": {"provider": {"extra_header_key": "Unused-Header"}}},
    {"fixture": {"provider": {"auth_header_format": "Token {token}"}}},
    {"fixture": {"provider": {"require_max_tokens": True}}},
    {"fixture": {"provider": {"require_max_tokens": True, "max_tokens": 4000}}},
    {"PILLOW_MAX_PIXELS": 10**9},
])
def test_operational_and_dormant_edits_allow_whole_file_retry(restart_case, change):
    folder, server = restart_case
    _restart_child(folder, server, cut="first_saved")
    _restart_child(folder, server, **change)
    assert len(server.requests) == 7
    assert _restart_rows(folder, "SELECT file_id FROM file_registry ORDER BY file_id") == [(1,), (2,), (3,)]


@pytest.mark.parametrize("no_retry", [["partial_fail", "ok"], ["ok"], []])
def test_explicit_retry_policy_does_not_turn_incomplete_attempts_into_success(restart_case, no_retry):
    folder, server = restart_case
    _restart_child(folder, server, cut="first_saved")
    _restart_child(folder, server, NO_RETRY_STATUSES=no_retry)
    assert len(server.requests) == (4 if "partial_fail" in no_retry else (7 if "ok" in no_retry else 10))
    assert _restart_rows(folder, "SELECT processing_complete FROM file_registry WHERE file_id=2") == [(0,)]
    assert _restart_rows(folder, "SELECT overall_result FROM overall_result WHERE file_id=2") == [("partial_fail",)]


@pytest.mark.parametrize("mode", ["table_per_file", "table_per_page"])
def test_repeated_restarts_keep_diagnostic_rows_and_complete_csv_xlsx(restart_case, mode):
    import csv
    from data_exporter import SQLiteDataExporter
    from openpyxl import load_workbook
    folder, server = restart_case
    server.mode = mode
    _restart_child(folder, server, mode, "first_saved")
    _restart_child(folder, server, mode, "first_saved")
    _restart_child(folder, server, mode)
    assert len(server.requests) == 8
    old_values = _restart_rows(folder, "SELECT file_id,values_json FROM llm_answers ORDER BY llm_answer_id")
    assert [r[0] for r in old_values] == [1, 1, 1, 2, 3, 4, 4, 4]
    exporter = SQLiteDataExporter(folder / "state.db")
    result = exporter.export_all_formats(folder / "reports")
    assert result.as_dict()["status"] == "success" and not result.recovery
    paths = list((folder / "reports").glob("*"))
    results_csv = next(p for p in paths if p.suffix == ".csv" and "results" in p.name.lower())
    with results_csv.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert [int(r["file_id"]) for r in rows] == [1, 1, 1, 2, 3, 4, 4, 4]
    assert [r["llm_answer"] for r in rows] == [
        f"call-{i + 1}-page-{page}" for i, page in enumerate([1, 2, 3, 1, 1, 1, 2, 3])]
    for index, row in enumerate(rows):
        assert ("Incomplete attempt; diagnostic only; not reused" in row["notes"]) == (index in (3, 4))
        assert row["file_result"] == ("partial_fail" if index in (3, 4) else "ok")
    book = load_workbook(next(p for p in paths if p.suffix == ".xlsx"), read_only=True)
    try:
        for csv_path in (p for p in paths if p.suffix == ".csv"):
            with csv_path.open(encoding="utf-8-sig", newline="") as stream:
                expected = list(csv.reader(stream))
            candidates = [sheet for sheet in book if list(next(sheet.values, ())) == expected[0]]
            assert len(candidates) == 1
            actual = [["" if value is None else str(value) for value in row] for row in candidates[0].values]
            assert actual == expected
    finally:
        book.close()


@pytest.mark.parametrize("stop_at", [0, 1, 3])
def test_cancellation_stops_image_preparation_before_the_next_image(tmp_path, monkeypatch, stop_at):
    from test_llm_client_network import make_wire_client
    client = make_wire_client(SimpleNamespace(url="http://127.0.0.1:1/unused"))
    flag = threading.Event()
    if not stop_at:
        flag.set()
    reads = []

    def encode(path):
        reads.append(path)
        if len(reads) == stop_at:
            flag.set()
        return "ZmFrZQ=="

    monkeypatch.setattr(client, "_encode_image_to_base64", encode)
    def forbidden(*args, **kwargs):
        raise AssertionError("No request may start after Stop")
    monkeypatch.setattr("requests.Session", forbidden)
    result = client.execute_network_inference([(n, tmp_path / f"{n}.jpg", "") for n in range(1, 6)], flag)
    assert len(reads) == stop_at
    assert result[0].status == ("not_attempted" if not stop_at else "aborted_by_user")


@pytest.mark.parametrize("cut,expected_uploads", [("write_answers", 4), ("write_complete", 6)])
def test_storage_failure_stops_before_any_later_file_upload(restart_case, cut, expected_uploads):
    folder, server = restart_case
    (folder / "input" / "C.tiff").write_bytes(b"generated third file")
    _restart_child(folder, server, cut=cut)
    assert len(server.requests) == expected_uploads
    assert (folder / "routed.txt").read_text().splitlines() == ["A.tiff", "B.tiff"]
    assert _restart_rows(folder, "SELECT file_path,processing_complete FROM file_registry ORDER BY file_id") == [
        ("A.tiff", 1), ("B.tiff", 0)]


@pytest.mark.parametrize("mode", ["table_per_file", "table_per_page"])
def test_restart_reextracts_even_when_all_previous_jpegs_disappeared(restart_case, mode):
    folder, server = restart_case
    server.mode = mode
    _restart_child(folder, server, mode, "first_saved")
    for image in (folder / "output").glob("*.jpg"):
        image.unlink()
    _restart_child(folder, server, mode)
    assert len(server.requests) == 7
    assert sorted(p.name for p in (folder / "output").glob("*.jpg")) == [
        "3-page-1.jpg", "3-page-2.jpg", "3-page-3.jpg"]
    assert _restart_rows(folder, "SELECT processing_complete FROM file_registry ORDER BY file_id") == [(1,), (0,), (1,)]


def test_waiting_for_provider_holds_no_database_write_transaction(restart_case):
    from fake_llm.generic import openai_reply
    folder, server = restart_case
    controls = []

    def lock_database(record):
        with sqlite3.connect(folder / "state.db", timeout=0.1) as connection:
            connection.execute("BEGIN IMMEDIATE")
            controls.append(True)
            connection.rollback()
    server.queue(on_request=lock_database, json=openai_reply('{"answer":"saved"}'))
    _restart_child(folder, server)
    assert controls == [True]


@pytest.mark.parametrize("end", ["invalid-body", "valid-answer", "unavailable"])
def test_corrective_retry_final_outcome_never_promotes_an_old_invalid_answer(restart_case, end):
    from fake_llm.generic import openai_reply
    folder, server = restart_case
    for page in (1, 2, 3):
        server.queue(json=openai_reply(f'{{"answer":"A-{page}"}}'))
    server.queue(json=openai_reply("old malformed answer"))
    if end == "valid-answer":
        server.queue(json=openai_reply('{"answer":"new valid answer"}'))
    elif end == "invalid-body":
        server.queue(raw='{"truncated":')
    else:
        server.queue(status=503, json={"error":"unavailable"})
    _restart_child(folder, server, LLM_JSON_MAX_ATTEMPTS=2, LLM_MAX_RETRIES=1)
    status, = _restart_rows(folder, "SELECT request_status FROM llm_requests WHERE file_id=2 AND request_number=1")
    rows = _restart_rows(folder, "SELECT values_json FROM llm_answers WHERE file_id=2 ORDER BY llm_answer_id")
    if end == "valid-answer":
        assert status == ("ok",) and rows[0] == ('{"answer": "new valid answer"}',)
    else:
        assert status[0] != "ok"
        assert all("old malformed answer" not in value for value, in rows)


def test_completion_oracle_rejects_a_suppressed_final_commit(tmp_path, monkeypatch):
    db = SQLiteDatabaseController(tmp_path / "state.db")
    try:
        db.handle_file_started(1, "file.png", ".png", "Generated")
        db.handle_frame_saved(1, PageResult(1, "1.jpg", "ok", ""))
        db.handle_file_completed(1, FileSummary(1, "1", "ok", "done"))
        def check():
            assert db.get_file_statuses(1)["overall_result"] == "ok"
            assert set(db.get_no_retry_sources(["ok"])) == {"file.png"}
        with monkeypatch.context() as broken:
            broken.setattr(db, "finalize_file", lambda *args: None)
            db.finalize_file(1)
            with pytest.raises(AssertionError):
                check()
        db.finalize_file(1)
        check()
    finally:
        db.close()


def test_decoder_cleanup_failure_cannot_mask_a_fatal_database_failure(tmp_path):
    input_dir = make_input(tmp_path, "A.png", "B.png")
    class BrokenDb(RecordingDb):
        def handle_frame_saved(self, file_id, page_result):
            raise SQLAlchemyError("storage failure")
    class BrokenCleanupRouter(StubRouter):
        def make_generator(self, path):
            try:
                yield PageResult(1, "1.jpg", "ok", "")
            finally:
                raise ValueError("decoder cleanup failed")
    router = BrokenCleanupRouter()
    with pytest.raises(RuntimeError, match="storage failure"):
        BatchOrchestrator(make_settings(input_dir), BrokenDb(), router, StubLogger()).execute_batch_processing_loop()
    assert router.routed == ["A.png"]


@pytest.mark.parametrize("change", ["same", "changed", "relocated"])
def test_run_identity_uses_loaded_prompt_contents(restart_case, change):
    folder, server = restart_case
    prompt = folder / "prompt.txt"
    prompt.write_text("same task", encoding="utf-8")
    _restart_child(folder, server, cut="first_saved", LLM_USER_PROMPT_MODE="FILE", LLM_USER_PROMPT=str(prompt))
    if change == "changed":
        prompt.write_text("new task", encoding="utf-8")
    elif change == "relocated":
        prompt = folder / "moved-prompt.txt"
        prompt.write_text("same task", encoding="utf-8")
    result = _restart_child(folder, server, check=False,
                            LLM_USER_PROMPT_MODE="FILE", LLM_USER_PROMPT=str(prompt))
    if change == "changed":
        assert result.returncode != 0 and "err_resume_ai_work_changed" in result.stderr
        assert len(server.requests) == 4
    else:
        assert result.returncode == 0, result.stderr
        assert len(server.requests) == 7


def test_changed_source_root_is_refused_even_when_every_file_completed(tmp_path, monkeypatch):
    from app_context import ProcessorCore
    from config_validator import Settings
    import central_logger
    from PIL import Image
    monkeypatch.setattr(central_logger, "setup_logging", lambda *args: None)
    first, other = tmp_path / "first", tmp_path / "other"
    first.mkdir()
    other.mkdir()
    Image.new("RGB", (8, 8), "red").save(first / "same.png")
    Image.new("RGB", (8, 8), "blue").save(other / "same.png")
    settings = Settings(INPUT_FOLDER_PATH=first, OUTPUT_FOLDER_PATH=tmp_path / "output", START_OVER=False)
    core = ProcessorCore(settings, threading.Event())
    core.run()
    db_path = core.db_path
    with sqlite3.connect(db_path) as connection:
        before = connection.execute("SELECT * FROM file_registry").fetchall()
    settings.INPUT_FOLDER_PATH = other
    with pytest.raises(ConfigurationError, match="err_resume_ai_work_changed"):
        ProcessorCore(settings, threading.Event())
    with sqlite3.connect(db_path) as connection:
        assert connection.execute("SELECT * FROM file_registry").fetchall() == before
    settings.INPUT_FOLDER_PATH = first / "."
    ProcessorCore(settings, threading.Event()).run()
    with sqlite3.connect(db_path) as connection:
        assert connection.execute("SELECT * FROM file_registry").fetchall() == before



def _resuming_core(tmp_path, monkeypatch):
    from app_context import ProcessorCore
    from config_validator import Settings
    import central_logger
    monkeypatch.setattr(central_logger, "setup_logging", lambda *args: None)
    (tmp_path / "input").mkdir(exist_ok=True)
    settings = Settings(INPUT_FOLDER_PATH=tmp_path / "input", OUTPUT_FOLDER_PATH=tmp_path / "output",
                        START_OVER=False)

    def run():
        core = ProcessorCore(settings, threading.Event())
        core.run()
        return core
    return settings, run


def _attempts(core):
    return _acceptance_rows(core.db_path, "SELECT f.file_id, f.file_path, o.overall_result "
                            "FROM file_registry f JOIN overall_result o USING(file_id) ORDER BY f.file_id")


def _jpeg_colour(core, file_id):
    from contextlib import closing
    from PIL import Image
    with closing(sqlite3.connect(core.db_path)) as connection:
        (name,), = connection.execute("SELECT output_file FROM page_log WHERE file_id=?", (file_id,)).fetchall()
    with Image.open(core.router.output_folder / name) as image:
        colour = image.convert("RGB").getpixel((8, 8))
    assert isinstance(colour, tuple)
    return colour


@pytest.mark.parametrize("replacement", ["other size, same time", "same size, other time"])
def test_a_replaced_source_at_the_same_path_is_processed_again(tmp_path, monkeypatch, replacement):
    import os
    from PIL import Image
    _, run = _resuming_core(tmp_path, monkeypatch)
    receipt = tmp_path / "input" / "receipt.png"
    Image.new("RGB", (64, 64), "red").save(receipt)
    run()
    before = receipt.stat()
    Image.new("RGB", (96, 96) if replacement == "other size, same time" else (64, 64), "blue").save(receipt)
    if replacement == "other size, same time":
        os.utime(receipt, ns=(before.st_atime_ns, before.st_mtime_ns))
        assert receipt.stat().st_size != before.st_size
    else:
        os.utime(receipt, ns=(before.st_atime_ns, before.st_mtime_ns + 10**9))
        assert receipt.stat().st_size == before.st_size
    core = run()
    assert _attempts(core) == [(1, "receipt.png", "ok"), (2, "receipt.png", "ok")]
    red, blue = _jpeg_colour(core, 1), _jpeg_colour(core, 2)
    assert red[0] > 200 > red[2] and blue[2] > 200 > blue[0]
    from data_exporter import SQLiteDataExporter
    SQLiteDataExporter(core.db_path).export_all_formats(tmp_path / "report")
    table, sheet = _results_and_workbook(tmp_path / "report")
    for rows in (table, sheet):
        assert [row[-1] for row in rows[1:]] == [
            "File attempt 1 | Source file changed after this attempt",
            "File attempt 2 | Source file changed since attempt 1"]
    run()
    assert len(_attempts(core)) == 2


@pytest.mark.parametrize("change", ["unchanged", "touched", "same size and time"])
def test_an_unchanged_size_and_time_skip_without_reading_the_file(tmp_path, monkeypatch, change):
    import os
    import batch_orchestrator
    from PIL import Image
    _, run = _resuming_core(tmp_path, monkeypatch)
    receipt = tmp_path / "input" / "receipt.png"
    Image.new("RGB", (64, 64), "red").save(receipt)
    core = run()
    before = receipt.stat()
    if change == "touched":
        os.utime(receipt, ns=(before.st_atime_ns, before.st_mtime_ns + 10**9))
    elif change == "same size and time":
        Image.new("RGB", (64, 64), "blue").save(receipt)
        os.utime(receipt, ns=(before.st_atime_ns, before.st_mtime_ns))
        assert receipt.stat().st_size == before.st_size
    reads = []
    real_sha256 = batch_orchestrator.source_sha256
    monkeypatch.setattr(batch_orchestrator, "source_sha256",
                        lambda path, stop=None: reads.append(path.name) or real_sha256(path, stop))
    run()
    assert _attempts(core) == [(1, "receipt.png", "ok")]
    assert reads == (["receipt.png"] if change == "touched" else [])


def test_a_source_changed_during_conversion_stays_incomplete_and_unsent(tmp_path):
    input_dir = make_input(tmp_path, "A.png")
    source = input_dir / "A.png"

    class RewritingRouter(StubRouter):
        rewrite = True

        def make_generator(self, path):
            yield PageResult(1, f"{path.stem}.jpg", Status.OK.value, "")
            if self.rewrite:
                source.write_bytes(b"a longer, different version")
            return FileSummary(1, "1", Status.OK.value, "done")

    db = SQLiteDatabaseController(tmp_path / "state.db")
    llm = ForbiddenLLM()
    router = RewritingRouter()
    settings = make_settings(input_dir, ENABLE_LLM_INFERENCE=True)
    try:
        BatchOrchestrator(settings, db, router, StubLogger(), llm).execute_batch_processing_loop()
        assert llm.calls == 0
        assert _acceptance_rows(tmp_path / "state.db",
                                "SELECT processing_complete, processing_error, range_status FROM file_registry"
                                ) == [(0, "The source file changed while it was being converted.", "ok")]
        router.rewrite = False
        BatchOrchestrator(settings, db, router, StubLogger(), llm).execute_batch_processing_loop()
        assert llm.calls == 1
        assert set(db.get_no_retry_sources(["ok"])) == {"A.png"}
        assert db.get_no_retry_sources(["ok"])["A.png"] == [fingerprint(source)]
    finally:
        db.close()


def test_stop_while_fingerprinting_leaves_the_file_unregistered(tmp_path, monkeypatch):
    import batch_orchestrator
    import fs_utils
    input_dir = make_input(tmp_path, "A.png", "B.png")
    (input_dir / "B.png").write_bytes(b"x" * (3 * fs_utils.SOURCE_HASH_CHUNK_BYTES))
    stop = threading.Event()
    real_sha256 = batch_orchestrator.source_sha256

    def stop_during_b(path, flag=None):
        if path.name == "B.png":
            stop.set()
        return real_sha256(path, flag)

    monkeypatch.setattr(batch_orchestrator, "source_sha256", stop_during_b)
    db = RecordingDb()
    router = StubRouter()
    BatchOrchestrator(make_settings(input_dir), db, router, StubLogger()).execute_batch_processing_loop(stop)
    assert router.routed == ["A.png"]
    assert db.started == [(1, "A.png")]


def test_stop_while_confirming_a_completed_file_ends_the_run(tmp_path, monkeypatch):
    import os
    import batch_orchestrator
    input_dir = make_input(tmp_path, "A.png", "B.png")
    done = input_dir / "A.png"
    db = SQLiteDatabaseController(tmp_path / "state.db")
    try:
        BatchOrchestrator(make_settings(input_dir), db, StubRouter(), StubLogger()).execute_batch_processing_loop()
        before = done.stat()
        done.write_bytes(b"y")
        os.utime(done, ns=(before.st_atime_ns, before.st_mtime_ns + 10**9))
        assert done.stat().st_size == before.st_size
        stop = threading.Event()
        real_sha256 = batch_orchestrator.source_sha256

        def stop_while_confirming(path, flag=None):
            stop.set()
            return real_sha256(path, flag)

        monkeypatch.setattr(batch_orchestrator, "source_sha256", stop_while_confirming)
        router = StubRouter()
        BatchOrchestrator(make_settings(input_dir), db, router, StubLogger()).execute_batch_processing_loop(stop)
        assert router.routed == []
        assert db.get_highest_file_id() == 2
    finally:
        db.close()


def test_a_completed_attempt_without_a_readable_source_is_processed_again(tmp_path, monkeypatch):
    import batch_orchestrator
    input_dir = make_input(tmp_path, "A.png")
    db = SQLiteDatabaseController(tmp_path / "state.db")
    try:
        with monkeypatch.context() as unreadable:
            unreadable.setattr(batch_orchestrator, "source_state", lambda path: None)
            BatchOrchestrator(make_settings(input_dir), db, StubRouter(), StubLogger()).execute_batch_processing_loop()
        assert db.get_no_retry_sources(["ok"]) == {"A.png": []}
        router = StubRouter()
        BatchOrchestrator(make_settings(input_dir), db, router, StubLogger()).execute_batch_processing_loop()
        assert router.routed == ["A.png"]
        assert db.get_no_retry_sources(["ok"])["A.png"] == [fingerprint(input_dir / "A.png")]
    finally:
        db.close()


def test_unsupported_files_are_not_read_and_resume_by_size_and_time(tmp_path, monkeypatch):
    import os
    import batch_orchestrator
    input_dir = make_input(tmp_path, "archive.iso")
    reads = []
    real_sha256 = batch_orchestrator.source_sha256
    monkeypatch.setattr(batch_orchestrator, "source_sha256",
                        lambda path, stop=None: reads.append(path.name) or real_sha256(path, stop))
    db = SQLiteDatabaseController(tmp_path / "state.db")
    settings = make_settings(input_dir, NO_RETRY_STATUSES=[OverallResult.OK, OverallResult.FAIL])

    class RejectingRouter(StubRouter):
        def make_generator(self, path):
            return FileSummary(0, "", Status.FAILURE.value, "Unsupported file extension: .iso")
            yield

    try:
        router = RejectingRouter()
        for _ in range(2):
            BatchOrchestrator(settings, db, router, StubLogger()).execute_batch_processing_loop()
        assert router.routed == ["archive.iso"]
        iso = input_dir / "archive.iso"
        os.utime(iso, ns=(iso.stat().st_atime_ns, iso.stat().st_mtime_ns + 10**9))
        BatchOrchestrator(settings, db, router, StubLogger()).execute_batch_processing_loop()
        assert router.routed == ["archive.iso", "archive.iso"]
        assert reads == []
        assert {fp.sha256 for fp in db.get_no_retry_sources(["fail"])["archive.iso"]} == {""}
    finally:
        db.close()


@pytest.mark.parametrize("bad_rows", [
    [{"page": 99, "answer": "wrong"}],
    [{"page": 1, "answer": "old"}, {"page": 1, "answer": "duplicate"}],
    [{"page": 1, "answer": "old"}, {"page": 3, "answer": "old"}],
])
def test_bad_answer_history_cannot_poison_a_new_file_attempt(restart_case, bad_rows):
    import json
    from fake_llm.generic import openai_reply
    folder, server = restart_case
    server.mode = "table_per_page"
    server.queue(json=openai_reply(json.dumps([{"page": n, "answer": f"A-{n}"} for n in (1, 2, 3)])))
    server.queue(json=openai_reply(json.dumps(bad_rows)))
    _restart_child(folder, server, "table_per_page", MAX_JPEGS_PER_INFERENCE=3)
    old = _restart_rows(folder, "SELECT * FROM llm_answers WHERE file_id=2 ORDER BY llm_answer_id")
    assert old
    assert _restart_rows(folder, "SELECT overall_result FROM overall_result WHERE file_id=2") == [("partial_fail",)]
    _restart_child(folder, server, "table_per_page", MAX_JPEGS_PER_INFERENCE=3)
    assert len(server.requests) == 3
    assert _restart_rows(folder, "SELECT * FROM llm_answers WHERE file_id=2 ORDER BY llm_answer_id") == old
    assert _restart_rows(folder, "SELECT overall_result FROM overall_result WHERE file_id=3") == [("ok",)]
    assert _restart_rows(folder, "SELECT values_json FROM llm_answers WHERE file_id=3 ORDER BY llm_answer_id") == [
        (f'{{"answer": "call-3-page-{n}"}}',) for n in (1, 2, 3)]


ACCEPTANCE_PROMPT = "ACCEPTANCE-PROMPT-SENTINEL-7f3a"
INCOMPLETE_NOTE = "Incomplete attempt; diagnostic only; not reused"
IDENTITY_URL = "http://127.0.0.1:9/v1/chat/completions"


def _acceptance_settings(folder, base_url, mode, **overrides):
    from urllib.parse import urlparse
    from config_validator import Settings, _default_provider_configs
    path = urlparse(_default_provider_configs()["openai"].url).path
    values: dict[str, Any] = {
        "INPUT_FOLDER_PATH": str(Path(folder) / "input"), "OUTPUT_FOLDER_PATH": str(Path(folder) / "output"),
        "ENABLE_LLM_INFERENCE": True, "LLM_PROVIDER": "openai",
        "LLM_PROVIDERS": {"openai": {"url": base_url.rstrip("/") + path}},
        "LLM_MAX_RETRIES": 1, "LLM_RETRY_SLEEP_SECONDS": 0, "LLM_TIMEOUT_SECONDS": 30,
        "HALT_ON_LLM_PARSE_ERROR": False, "LLM_JSON_MAX_ATTEMPTS": 1, "MAX_JPEGS_PER_INFERENCE": 2,
        "LLM_OUTPUT_MODE": mode, "LLM_OUTPUT_COLUMNS": "answer", "LLM_USER_PROMPT": ACCEPTANCE_PROMPT,
        "START_OVER": False,
    }
    values.update(overrides)
    return Settings(**values)


def _acceptance_inputs(folder, *extra):
    from PIL import Image
    (folder / "input").mkdir()
    Image.new("RGB", (40, 40), (200, 30, 30)).save(folder / "input" / "A.png")
    pages = [Image.new("RGB", (40, 40), (60 * n, 90, 200 - 50 * n)) for n in range(3)]
    pages[0].save(folder / "input" / "B.tif", save_all=True, append_images=pages[1:])
    for name in extra:
        Image.new("RGB", (40, 40), (20, 160, 60)).save(folder / "input" / name)


def _label_model(mode):
    import json
    import re
    from fake_llm import ContentPolicy

    class LabelModel(ContentPolicy):
        def __init__(self):
            super().__init__()
            self.calls = 0
            self.kill_at = 0
            self.process: Any = None
            self.after_call: dict[int, Any] = {}
            self.guard = threading.Lock()

        def answer(self, body=None):
            with self.guard:
                self.calls += 1
                call = self.calls
            pages = [int(page) for page in re.findall(r"Image \d+ of \d+ - page (\d+)", json.dumps(body))]
            if call == self.kill_at and self.process is not None:
                self.process.kill()
                self.process.wait(30)
            if call in self.after_call:
                self.after_call[call]()
            if mode == "table_per_page":
                return json.dumps([{"page": page, "answer": f"call-{call}-page-{page}"} for page in pages])
            return json.dumps({"answer": f"call-{call}-pages-{pages[0]}-{pages[-1]}"})
    return LabelModel()


def _real_core_child(folder, base_url, mode, token):
    import json
    import os
    import central_logger
    import llm_client
    from app_context import ProcessorCore
    for key in [key for key in os.environ if key.endswith("_TOKEN")]:
        del os.environ[key]
    central_logger.setup_logging = lambda *args, **kwargs: None
    llm_client.get_env_tokens = lambda: {"OPENAI_TOKEN": token}
    events: list[Any] = []
    ProcessorCore(_acceptance_settings(folder, base_url, mode), threading.Event(), on_progress=events.append).run()
    (Path(folder) / "events.json").write_text(json.dumps(events, default=str), encoding="utf-8")


def _run_real_core_child(folder, server, model, mode, kill_at=0):
    import subprocess
    from fake_llm.harness import KEY_SHAPED_TOKENS
    (folder / "events.json").unlink(missing_ok=True)
    child = "import runpy,sys; runpy.run_path(sys.argv[1])['_real_core_child'](*sys.argv[2:])"
    command = [sys.executable, "-c", child, str(Path(__file__).resolve()), str(folder), server.base_url, mode,
               KEY_SHAPED_TOKENS["openai"]]
    process = subprocess.Popen(command, cwd=Path(__file__).resolve().parents[1], text=True,
                               stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    model.kill_at, model.process = kill_at, process
    try:
        output, errors = process.communicate(timeout=120)
    finally:
        if process.poll() is None:
            process.kill()
            process.wait(30)
    return process.returncode, output + errors


def _acceptance_rows(database, query):
    from contextlib import closing
    with closing(sqlite3.connect(database)) as connection:
        return connection.execute(query).fetchall()


def _results_and_workbook(tech_folder):
    import csv
    from openpyxl import load_workbook
    with sorted(tech_folder.glob("results_*.csv"))[-1].open(encoding="utf-8-sig", newline="") as stream:
        table = list(csv.reader(stream))
    book = load_workbook(sorted(tech_folder.glob("database_export_*.xlsx"))[-1], read_only=True)
    try:
        sheet = [["" if value is None else str(value) for value in row] for row in book["Results"].values]
    finally:
        book.close()
    return table, sheet


def _expected_answers(mode, call, pages):
    if mode == "table_per_page":
        return [f"call-{call}-page-{page}" for page in pages]
    return [f"call-{call}-pages-{pages[0]}-{pages[-1]}"]


@pytest.mark.parametrize("mode", ["table_per_file", "table_per_page"])
@pytest.mark.parametrize("kill_at", [2, 3], ids=["before-any-B-outcome", "after-first-B-request"])
def test_real_core_killed_during_a_request_restarts_only_the_unfinished_file(tmp_path, mode, kill_at):
    import base64
    import io
    import json
    import re
    from PIL import Image
    from fake_llm import make_server
    from fake_llm.harness import KEY_SHAPED_TOKENS
    _acceptance_inputs(tmp_path)
    model = _label_model(mode)
    server = make_server("openai").start()
    server.content = model
    try:
        code, output = _run_real_core_child(tmp_path, server, model, mode, kill_at)
        assert code != 0 and not (tmp_path / "events.json").exists(), output
        first_uploads = len(server.requests)
        assert first_uploads == kill_at
        database = next((tmp_path / "output").rglob("application_state.db"))

        def rows(query):
            return _acceptance_rows(database, query)
        registry = "SELECT file_id,file_path,processing_complete FROM file_registry ORDER BY file_id"
        assert rows(registry) == [(1, "A.png", 1), (2, "B.tif", 0)]
        assert rows("SELECT overall_result FROM overall_result ORDER BY file_id") == [("ok",), ("partial_fail",)]
        assert rows("SELECT request_number,request_status FROM llm_requests WHERE file_id=2") == (
            [] if kill_at == 2 else [(1, "ok")])
        old_answers = rows("SELECT * FROM llm_answers ORDER BY llm_answer_id")
        stored = b"".join(path.read_bytes() for path in database.parent.glob("application_state.db*"))
        token = KEY_SHAPED_TOKENS["openai"]
        assert token and token.encode() not in stored
        assert ACCEPTANCE_PROMPT.encode() not in stored
        earlier_b = [path for path in (tmp_path / "output").rglob("*.jpg") if path.name.startswith("2_")]
        assert len(earlier_b) == 3
        for path in earlier_b:
            path.write_bytes(b"tampered earlier attempt, not an image")

        code, output = _run_real_core_child(tmp_path, server, model, mode)
        assert code == 0, output
        events = json.loads((tmp_path / "events.json").read_text(encoding="utf-8"))
        assert events[-1] == {"type": "done"}
        resent = server.requests[first_uploads:]
        assert [re.findall(r"Image (\d+) of (\d+) - page (\d+)", json.dumps(request.json)) for request in resent] == [
            [("1", "3", "1"), ("2", "3", "2")], [("3", "3", "3")]]
        for request in resent:
            for data in re.findall(r"data:image/jpeg;base64,([A-Za-z0-9+/=]+)", json.dumps(request.json)):
                with Image.open(io.BytesIO(base64.b64decode(data))) as sent:
                    assert sent.format == "JPEG"
        assert rows(registry) == [(1, "A.png", 1), (2, "B.tif", 0), (3, "B.tif", 1)]
        assert rows("SELECT overall_result FROM overall_result ORDER BY file_id") == [
            ("ok",), ("partial_fail",), ("ok",)]
        assert rows("SELECT * FROM llm_answers WHERE file_id<3 ORDER BY llm_answer_id") == old_answers
        renewed = [*_expected_answers(mode, first_uploads + 1, [1, 2]),
                   *_expected_answers(mode, first_uploads + 2, [3])]
        assert rows("SELECT values_json FROM llm_answers WHERE file_id=3 ORDER BY llm_answer_id") == [
            (json.dumps({"answer": value}),) for value in renewed]
        assert sorted(path.name.split("_")[0] for path in (tmp_path / "output").rglob("*.jpg")) == [
            "1", "2", "2", "2", "3", "3", "3"]

        table, sheet = _results_and_workbook(database.parent)
        assert sheet == table
        column = {name: index for index, name in enumerate(table[0])}
        earlier = [] if kill_at == 2 else _expected_answers(mode, 2, [1, 2])
        assert [(row[column["file_id"]], row[column["llm_answer"]]) for row in table[1:]] == [
            ("1", *_expected_answers(mode, 1, [1])), *(("2", value) for value in earlier or [""]),
            *(("3", value) for value in renewed)]
        for row in table[1:]:
            file_id, notes = row[column["file_id"]], row[column["notes"]]
            assert (INCOMPLETE_NOTE in notes) == (file_id == "2"), row
            assert row[column["file_result"]] == ("partial_fail" if file_id == "2" else "ok"), row
            if file_id != "1":
                assert f"File attempt {int(file_id) - 1}" in notes, row
    finally:
        server.stop()


@pytest.fixture
def isolated_core(monkeypatch):
    import os
    import central_logger
    import llm_client
    from fake_llm.harness import KEY_SHAPED_TOKENS
    for key in [key for key in os.environ if key.endswith("_TOKEN")]:
        monkeypatch.delenv(key)
    monkeypatch.setattr(central_logger, "setup_logging", lambda *args, **kwargs: None)
    tokens = {"OPENAI_TOKEN": KEY_SHAPED_TOKENS["openai"]}
    monkeypatch.setattr(llm_client, "get_env_tokens", lambda: dict(tokens))
    return tokens


def test_authentication_halt_exports_diagnostics_and_retry_restarts_only_the_halted_file(tmp_path, isolated_core):
    import json
    from app_context import ProcessorCore
    from fake_llm import make_server
    _acceptance_inputs(tmp_path, "C.png")
    model = _label_model("table_per_file")
    server = make_server("openai").start()
    server.content = model
    model.after_call[2] = lambda: server.queue(status=401, json={"error": {
        "message": "Incorrect API key provided", "type": "invalid_request_error", "code": "invalid_api_key"}})
    try:
        settings = _acceptance_settings(tmp_path, server.base_url, "table_per_file")
        events: list[Any] = []
        core = ProcessorCore(settings, threading.Event(), on_progress=events.append)
        core.run()
        database = core.db_path
        assert events[-1]["type"] == "failed"
        assert len(server.requests) == 3

        def rows(query):
            return _acceptance_rows(database, query)
        registry = "SELECT file_id,file_path,processing_complete FROM file_registry ORDER BY file_id"
        assert rows(registry) == [(1, "A.png", 1), (2, "B.tif", 0)]
        assert rows("SELECT request_number,request_status FROM llm_requests WHERE file_id=2 ORDER BY request_id") == [
            (1, "ok"), (2, "network_failure")]
        error, = rows("SELECT processing_error FROM file_registry WHERE file_id=2")[0]
        assert "FATAL AUTHENTICATION ERROR" in error
        table, sheet = _results_and_workbook(database.parent)
        assert sheet == table
        column = {name: index for index, name in enumerate(table[0])}
        assert [row[column["file_id"]] for row in table[1:]] == ["1", "2", "2"]
        assert all(INCOMPLETE_NOTE in row[column["notes"]] for row in table[2:])
        assert "Request failed" in table[3][column["notes"]]

        isolated_core["OPENAI_TOKEN"] = isolated_core["OPENAI_TOKEN"][:-1] + (
            "A" if isolated_core["OPENAI_TOKEN"][-1] != "A" else "B")
        repaired = _acceptance_settings(tmp_path, server.base_url, "table_per_file", LLM_TIMEOUT_SECONDS=45,
                                        INPUT_FOLDER_PATH=str(tmp_path / "input").replace("\\", "/") + "/")
        events = []
        ProcessorCore(repaired, threading.Event(), on_progress=events.append).run()
        assert events[-1] == {"type": "done"}
        resent = server.requests[3:]
        assert [len(json.dumps(request.json).split("data:image/jpeg")) - 1 for request in resent] == [2, 1, 1]
        assert rows(registry) == [(1, "A.png", 1), (2, "B.tif", 0), (3, "B.tif", 1), (4, "C.png", 1)]
        assert rows("SELECT overall_result FROM overall_result WHERE file_id>2 ORDER BY file_id") == [("ok",), ("ok",)]

        with sqlite3.connect(database) as connection:
            before = list(connection.iterdump())
        changed = _acceptance_settings(tmp_path, server.base_url, "table_per_file", LLM_USER_PROMPT="another task")
        with pytest.raises(ConfigurationError, match="err_resume_ai_work_changed"):
            ProcessorCore(changed, threading.Event())
        with sqlite3.connect(database) as connection:
            assert list(connection.iterdump()) == before
        assert len(server.requests) == 6
    finally:
        server.stop()


IDENTITY_TOKEN = "identity-token-a"


def _identity(tmp_path, token=IDENTITY_TOKEN, **overrides):
    from app_context import processing_identity
    from config_validator import Settings
    from llm_client import LLMClient
    for name in ("input", "output"):
        (tmp_path / name).mkdir(exist_ok=True)
    values: dict[str, Any] = {
        "INPUT_FOLDER_PATH": str(tmp_path / "input"), "OUTPUT_FOLDER_PATH": str(tmp_path / "output"),
        "ENABLE_LLM_INFERENCE": True, "LLM_PROVIDER": "openai", "LLM_PROVIDERS": {"openai": {"url": IDENTITY_URL}},
    }
    values.update(overrides)
    settings = Settings(**values)
    return processing_identity(settings, LLMClient(settings, token=token) if settings.ENABLE_LLM_INFERENCE else None)


def _openai_preset():
    from config_validator import _default_provider_configs
    return _default_provider_configs()["openai"]


ALLOWED_IDENTITY_EDITS = {
    "credential": {"token": "identity-token-b"},
    "logging": {"LOGGING_LEVEL": "DEBUG"},
    "language": {"GUI_LANGUAGE": "ru"},
    "retry policy": {"NO_RETRY_STATUSES": ["ok", "partial_fail"]},
    "timeout": {"LLM_TIMEOUT_SECONDS": 7},
    "network retries": {"LLM_MAX_RETRIES": 9},
    "retry pause": {"LLM_RETRY_SLEEP_SECONDS": 11},
    "answer attempts": {"LLM_JSON_MAX_ATTEMPTS": 6},
    "parse halt": {"HALT_ON_LLM_PARSE_ERROR": False},
    "malformed halt": {"LLM_ABORT_ON_MALFORMED_JSON": True},
    "breaker threshold": {"MAX_CONSECUTIVE_LLM_FAILURES": 2},
    "inactive provider": {"LLM_PROVIDERS": {"openai": {"url": IDENTITY_URL}, "claude": {"model": "other-model"}}},
    "explicit preset model": {"LLM_PROVIDERS": {"openai": {"url": IDENTITY_URL, "model": _openai_preset().model}}},
    "inactive video mode": {"VIDEO_SAMPLING_CAPTURE_RATE_FPS": 3.0, "VIDEO_SAMPLING_MAX_FRAMES_BUDGET": 7},
    "column spacing": {"LLM_OUTPUT_COLUMNS": " answer "},
    "range spacing": {"DOCUMENT_RANGE": " 1 - 3 ", "_base": {"DOCUMENT_RANGE": "1-3"}},
}
if not _openai_preset().require_max_tokens:
    ALLOWED_IDENTITY_EDITS["dormant token limit"] = {"LLM_PROVIDERS": {"openai": {"url": IDENTITY_URL,
                                                                               "max_tokens": 1234}}}


@pytest.mark.parametrize("edit", list(ALLOWED_IDENTITY_EDITS.values()), ids=list(ALLOWED_IDENTITY_EDITS))
def test_run_identity_accepts_credential_operational_dormant_and_equivalent_edits(tmp_path, edit):
    edit = dict(edit)
    base = edit.pop("_base", {})
    token = edit.pop("token", IDENTITY_TOKEN)
    assert _identity(tmp_path, token, **{**base, **edit}) == _identity(tmp_path, **base)


REFUSED_IDENTITY_EDITS = {
    "endpoint": {"LLM_PROVIDERS": {"openai": {"url": IDENTITY_URL.replace(":9/", ":10/")}}},
    "model": {"LLM_PROVIDERS": {"openai": {"url": IDENTITY_URL, "model": "another-model"}}},
    "active extra header": {"LLM_PROVIDERS": {"openai": {"url": IDENTITY_URL, "extra_header_key": "API-Version",
                                                         "extra_header_value": "2"}}},
    "user prompt": {"LLM_USER_PROMPT": "another task"},
    "system prompt": {"LLM_SYSTEM_PROMPT": "another role"},
    "column meanings": {"LLM_OUTPUT_COLUMN_INSTRUCTIONS": "answer: the colour"},
    "columns": {"LLM_OUTPUT_COLUMNS": "answer,colour"},
    "answer mode": {"LLM_OUTPUT_MODE": "table_per_page"},
    "request grouping": {"MAX_JPEGS_PER_INFERENCE": 5},
    "jpeg quality": {"JPEG_QUALITY": 80},
    "resolution": {"MAX_DIMENSION": 1000},
    "document range": {"DOCUMENT_RANGE": "2-3"},
    "video mode": {"VIDEO_MODE": "SAMPLING"},
    "active video setting": {"VIDEO_SUMMARY_TARGET_TOTAL_FRAMES": 12},
    "ai switched off": {"ENABLE_LLM_INFERENCE": False},
}


@pytest.mark.parametrize("edit", list(REFUSED_IDENTITY_EDITS.values()), ids=list(REFUSED_IDENTITY_EDITS))
def test_run_identity_changes_with_effective_processing_semantics(tmp_path, edit):
    assert _identity(tmp_path, **edit) != _identity(tmp_path)


TOKEN_LIMIT_EDITS = {
    "higher limit": ({"require_max_tokens": True, "max_tokens": 1000},
                     {"require_max_tokens": True, "max_tokens": 4000}),
    "limit switched on": ({}, {"require_max_tokens": True, "max_tokens": 4000}),
    "limit field renamed": ({"require_max_tokens": True, "max_tokens_field": "max_tokens"},
                            {"require_max_tokens": True, "max_tokens_field": "max_completion_tokens"}),
}


@pytest.mark.parametrize("before,after", list(TOKEN_LIMIT_EDITS.values()), ids=list(TOKEN_LIMIT_EDITS))
def test_token_limit_edits_are_repairs_that_keep_the_run_identity(tmp_path, before, after):
    def identity(fields):
        return _identity(tmp_path, LLM_PROVIDERS={"openai": {"url": IDENTITY_URL, **fields}})
    assert identity(after) == identity(before)


def test_raising_the_token_limit_after_a_truncation_halt_resends_only_the_halted_file(tmp_path, isolated_core):
    import json
    from app_context import ProcessorCore
    from fake_llm import make_server
    from fake_llm.generic import openai_reply
    _acceptance_inputs(tmp_path)
    model = _label_model("table_per_file")
    server = make_server("openai").start()
    server.content = model
    model.after_call[1] = lambda: server.queue(json=openai_reply('{"answer": "cut o', finish_reason="length"))
    try:
        def settings(limit):
            return _acceptance_settings(
                tmp_path, server.base_url, "table_per_file", HALT_ON_LLM_PARSE_ERROR=True,
                LLM_PROVIDERS={"openai": {"url": server.base_url.rstrip("/") + "/v1/chat/completions",
                                          "require_max_tokens": True, "max_tokens": limit}})
        events: list[Any] = []
        core = ProcessorCore(settings(1000), threading.Event(), on_progress=events.append)
        core.run()
        database = core.db_path
        assert events[-1]["type"] == "failed"
        assert len(server.requests) == 2
        assert _acceptance_rows(
            database, "SELECT request_number,request_status FROM llm_requests WHERE file_id=2 ORDER BY request_id") == [
            (1, "token_limit_exceeded"), (2, "not_attempted")]
        accepted = _acceptance_rows(database, "SELECT * FROM llm_answers WHERE file_id=1")

        events = []
        ProcessorCore(settings(4000), threading.Event(), on_progress=events.append).run()
        assert events[-1] == {"type": "done"}
        resent = server.requests[2:]
        assert [len(json.dumps(request.json).split("data:image/jpeg")) - 1 for request in resent] == [2, 1]
        assert {request.json["max_completion_tokens"] for request in resent} == {4000}
        assert _acceptance_rows(database, "SELECT * FROM llm_answers WHERE file_id=1") == accepted
        assert _acceptance_rows(
            database, "SELECT file_id,file_path,processing_complete FROM file_registry ORDER BY file_id") == [
            (1, "A.png", 1), (2, "B.tif", 0), (3, "B.tif", 1)]
    finally:
        server.stop()


def test_raising_the_image_pixel_limit_after_a_refusal_retries_only_the_refused_image(
        tmp_path, isolated_core, monkeypatch):
    import hashlib
    from PIL import Image
    from app_context import ProcessorCore
    from config_validator import Settings
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", Image.MAX_IMAGE_PIXELS)
    (tmp_path / "input").mkdir()
    Image.new("RGB", (40, 40), (200, 30, 30)).save(tmp_path / "input" / "A.png")
    Image.new("RGB", (60, 60), (30, 90, 200)).save(tmp_path / "input" / "B.png")

    def run(limit, output="output"):
        events: list[Any] = []
        ProcessorCore(Settings(INPUT_FOLDER_PATH=tmp_path / "input",
                               OUTPUT_FOLDER_PATH=tmp_path / output, START_OVER=False,
                               ENABLE_LLM_INFERENCE=False, PILLOW_MAX_PIXELS=limit),
                      threading.Event(), on_progress=events.append).run()
        assert events[-1] == {"type": "done"}
        return tmp_path / output / "current_run"

    def jpeg(run_folder, stem):
        [path] = run_folder.rglob(f"*_{stem}_page_1.jpg")
        return path, hashlib.sha256(path.read_bytes()).hexdigest()

    results = ("SELECT r.file_id, r.file_path, o.overall_result FROM file_registry r "
               "JOIN overall_result o USING (file_id) ORDER BY r.file_id")
    with pytest.warns(Image.DecompressionBombWarning):
        run_folder = run(1000)
    database = run_folder / "TECH" / "application_state.db"
    a_path, a_hash = jpeg(run_folder, "A")
    a_written = a_path.stat().st_mtime_ns
    assert _acceptance_rows(database, results) == [(1, "A.png", "ok"), (2, "B.png", "fail")]
    assert "decompression bomb" in _acceptance_rows(
        database, "SELECT group_concat(page_to_jpeg_comment) FROM page_log WHERE file_id=2")[0][0]

    run(10_000)
    assert _acceptance_rows(database, results) == [(1, "A.png", "ok"), (2, "B.png", "fail"), (3, "B.png", "ok")]
    assert jpeg(run_folder, "A") == (a_path, a_hash) and a_path.stat().st_mtime_ns == a_written
    assert jpeg(run(10**9, "reference"), "A")[1] == a_hash


def test_run_identity_uses_prompt_text_not_its_storage_mode(tmp_path):
    prompt = tmp_path / "prompt.txt"
    prompt.write_text("Describe the receipt", encoding="utf-8")
    as_text = _identity(tmp_path, LLM_USER_PROMPT="Describe the receipt")
    assert _identity(tmp_path, LLM_USER_PROMPT_MODE="FILE", LLM_USER_PROMPT=str(prompt)) == as_text
    prompt.write_text("Describe the invoice", encoding="utf-8")
    assert _identity(tmp_path, LLM_USER_PROMPT_MODE="FILE", LLM_USER_PROMPT=str(prompt)) != as_text


def test_dormant_ai_settings_never_change_an_ai_off_identity(tmp_path):
    off = _identity(tmp_path, ENABLE_LLM_INFERENCE=False)
    assert _identity(tmp_path, ENABLE_LLM_INFERENCE=False, LLM_USER_PROMPT="another task",
                     LLM_OUTPUT_COLUMNS="answer,colour", MAX_JPEGS_PER_INFERENCE=5,
                     LLM_PROVIDERS={"openai": {"url": IDENTITY_URL, "model": "another-model"}}) == off
    assert _identity(tmp_path, ENABLE_LLM_INFERENCE=False, JPEG_QUALITY=80) != off


ROOT_SPELLINGS = ["trailing separator", "forward slashes", "relative"] + (
    ["letter case"] if sys.platform == "win32" else [])


@pytest.mark.parametrize("spelling", ROOT_SPELLINGS)
def test_run_identity_normalizes_equivalent_source_root_spellings(tmp_path, monkeypatch, spelling):
    import os
    root, digest = _identity(tmp_path)
    folder = str(tmp_path / "input")
    spelled = {"trailing separator": folder + os.sep, "forward slashes": folder.replace("\\", "/"),
               "relative": "input", "letter case": folder.upper()}[spelling]
    monkeypatch.chdir(tmp_path)
    assert _identity(tmp_path, INPUT_FOLDER_PATH=spelled) == (root, digest)
    (tmp_path / "other").mkdir()
    assert _identity(tmp_path, INPUT_FOLDER_PATH=str(tmp_path / "other"))[0] != root


@pytest.mark.skipif(sys.platform != "win32", reason="NTFS junctions and 8.3 short names")
@pytest.mark.parametrize("alias", ["junction", "8.3 short name"])
def test_run_identity_resolves_windows_aliases_of_the_source_root(tmp_path, alias):
    import ctypes
    import os
    import subprocess
    folder = tmp_path / "Long Source Folder"
    folder.mkdir()
    identity = _identity(tmp_path, INPUT_FOLDER_PATH=str(folder))
    if alias == "junction":
        spelled = str(tmp_path / "linked")
        shell = os.path.join(os.environ["SYSTEMROOT"], "System32", "cmd.exe")
        subprocess.run([shell, "/c", "mklink", "/J", spelled, str(folder)], check=True, capture_output=True)
    else:
        buffer = ctypes.create_unicode_buffer(1024)
        assert ctypes.windll.kernel32.GetShortPathNameW(str(folder), buffer, len(buffer))
        spelled = buffer.value
        if spelled.lower() == str(folder).lower():
            pytest.skip("8.3 name generation is disabled on this volume")
    assert _identity(tmp_path, INPUT_FOLDER_PATH=spelled) == identity


def test_process_exit_inside_a_multi_answer_request_leaves_neither_request_nor_answers(tmp_path):
    import json
    import re
    from fake_llm.generic import GenericServer, openai_reply

    class Server(GenericServer):
        def handle_unknown_route(self):
            pages = [int(n) for n in re.findall(r"- page (\d+)", json.dumps(self.requests[-1].json))]
            number = len(self.requests)
            return self.json_response(openai_reply(json.dumps(
                [{"page": page, "answer": f"call-{number}-page-{page}"} for page in pages])))

    (tmp_path / "input").mkdir()
    (tmp_path / "output").mkdir()
    for name in ("A.tiff", "B.tiff"):
        (tmp_path / "input" / name).write_bytes(b"generated router input")
    with Server() as server:
        _restart_child(tmp_path, server, "table_per_page", "inside_answers", MAX_JPEGS_PER_INFERENCE=3)
        assert _restart_rows(tmp_path, "SELECT COUNT(*) FROM llm_requests WHERE file_id=2") == [(0,)]
        assert _restart_rows(tmp_path, "SELECT COUNT(*) FROM llm_answers WHERE file_id=2") == [(0,)]
        completed = _restart_rows(tmp_path, "SELECT * FROM llm_answers WHERE file_id=1 ORDER BY llm_answer_id")
        assert len(completed) == 3
        _restart_child(tmp_path, server, "table_per_page", MAX_JPEGS_PER_INFERENCE=3)
        assert len(server.requests) == 3
        assert _restart_rows(tmp_path, "SELECT * FROM llm_answers WHERE file_id=1 ORDER BY llm_answer_id"
                             ) == completed
        assert _restart_rows(tmp_path, "SELECT values_json FROM llm_answers WHERE file_id=3 ORDER BY llm_answer_id"
                             ) == [
            (json.dumps({"answer": f"call-3-page-{page}"}),) for page in (1, 2, 3)]
